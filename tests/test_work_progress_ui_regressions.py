"""Regression tests for Work Progress media, notes, and report export UI contracts."""
import re
import subprocess
from pathlib import Path

from openpyxl.styles import Alignment

from app.services.report import build_daily_report


ROOT = Path(__file__).resolve().parents[1]
WORK_PROGRESS_JS = ROOT / "static/js/render/work-progress.js"
WORK_PROGRESS_CSS = ROOT / "static/css/style.work-progress.css"
GALLERY_LIFECYCLE_TEST = ROOT / "tests/work_progress_gallery_lifecycle.test.js"
DETAIL_TARGET_LIFECYCLE_TEST = ROOT / "tests/work_progress_detail_target_lifecycle.test.js"
CALENDAR_JS = ROOT / "static/js/render/calendar.js"
CALENDAR_CSS = ROOT / "static/css/style.calendar.css"
APP_JS = ROOT / "static/js/app.js"
CORE_CSS = ROOT / "static/css/style.core.css"
STOCKOUT_CSS = ROOT / "static/css/style.stockout.css"


def _read(path: Path) -> str:
    """Read one UTF-8 source asset for a contract assertion."""
    return path.read_text(encoding="utf-8")


def test_work_progress_gallery_is_a_viewport_overlay_and_is_closed_on_tab_change():
    """Regression: photo preview must not render as ordinary content at page bottom."""
    css = _read(WORK_PROGRESS_CSS)
    js = _read(WORK_PROGRESS_JS)
    app = _read(APP_JS)
    overlay_rules = re.findall(r"(?m)^[ \t]*\.wpr-gallery-overlay\s*\{([^}]*)\}", css)
    assert any("position: fixed" in rule and "inset: 0" in rule and "z-index:" in rule for rule in overlay_rules)
    assert "#content.wpr-content .wpr-gallery-overlay" not in css
    assert "document.body.appendChild(overlay)" in js
    assert "wprCloseGallery()" in js
    leave_block = app.split("if (previousTab === 'work-progress'", 1)[1].split("currentTab = tab", 1)[0]
    assert "wprCloseGallery()" in leave_block


def test_gallery_async_lifecycle_runtime_contract():
    """Regression: tab leave, stale errors, and rapid A/B opens must be safe."""
    result = subprocess.run(
        ["node", str(GALLERY_LIFECYCLE_TEST)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=120,
    )
    assert result.returncode == 0, f"gallery lifecycle runtime failed:\n{result.stdout}\n{result.stderr}"
    assert "14 passed" in result.stdout


def test_detail_target_lifecycle_runtime_contract():
    """Regression: selected detail actions must preserve target and token scope."""
    result = subprocess.run(
        ["node", str(DETAIL_TARGET_LIFECYCLE_TEST)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=120,
    )
    assert result.returncode == 0, f"detail target lifecycle runtime failed:\n{result.stdout}\n{result.stderr}"
    assert "4 passed" in result.stdout


def test_work_progress_and_calendar_notes_preserve_multiline_text():
    """Regression: escaped multiline notes must remain multiline at every affected surface."""
    wpr_css = _read(WORK_PROGRESS_CSS)
    cal_css = _read(CALENDAR_CSS)
    core_css = _read(CORE_CSS)
    stockout_css = _read(STOCKOUT_CSS)
    cal_js = _read(CALENDAR_JS)
    assert ".wpr-selected-summary span" in wpr_css
    assert ".wpr-edit-readonly-row strong" in wpr_css
    assert "white-space: pre-wrap" in wpr_css
    assert ".cal-note" in cal_css and "white-space: pre-wrap" in cal_css
    assert ".item-note-text" in core_css and "white-space: pre-wrap" in core_css
    assert ".stockout-note" in stockout_css and "white-space: pre-wrap" in stockout_css
    assert "${e.note ?" in cal_js
    assert "e.note || '無備註'" not in cal_js


def test_empty_work_progress_notes_do_not_render_placeholder_copy():
    """Regression: empty optional notes must hide their labels instead of saying no note."""
    js = _read(WORK_PROGRESS_JS)
    assert "無備註" not in js
    assert "尚未填寫備註" not in js
    assert "工作進度備註" not in js


def test_daily_report_note_cells_enable_wrapping():
    """Regression: multiline calendar notes must be visible in generated XLSX output."""
    buf, _ = build_daily_report(
        "2026-09-21",
        [{
            "client_name": "王先生",
            "address": "板橋區",
            "service_type_id": 1,
            "service_name": "保養",
            "start_time": "09:00",
            "end_time": "12:00",
            "note": "第一行\n第二行",
        }],
    )
    from openpyxl import load_workbook

    ws = load_workbook(buf).active
    assert ws["B6"].value == "第一行\n第二行"
    assert ws["B6"].alignment == Alignment(wrap_text=True, vertical="top")
