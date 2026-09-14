import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import unquote
sys.path.insert(0, str(Path(__file__).parents[1]))

import pytest
import app.database as app_db
import main as app_main
from app.services.auth import SESSION_COOKIE, create_session, init_admin_if_missing
from app.services.engineering_petty_cash import engineering_filename, engineering_safe_filename, engineering_sheet_title
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def eng_client(tmp_path, monkeypatch):
    monkeypatch.setattr(app_db, 'DB_PATH', str(tmp_path / 'engineering.db'))
    app_db.init_db()
    conn = app_db.get_db()
    init_admin_if_missing(conn)
    uid = conn.execute("SELECT id FROM users WHERE username='admin'").fetchone()['id']
    token = create_session(conn, uid)
    conn.close()
    c = TestClient(app_main.app)
    c.cookies.set(SESSION_COOKIE, token)
    return c


def payload():
    return {
        'report_type': 'engineering', 'start_date': '2026-09-01', 'end_date': '2026-09-04',
        'upload_person': '藍先生', 'prepared_by': '王小明', 'filename_text': '發票',
        'status': 'completed', 'categories': [{
            'name': '交通費', 'groups': [{'name': '油資', 'receipts': [
                {'tax_id_mark': 'V', 'receipt_number': 'DH-98939443', 'amount': 100, 'details': ['九二無鉛']},
                {'tax_id_mark': '12345678', 'receipt_number': 'DH-98728718', 'amount': 725, 'details': ['九二無鉛']},
            ]}]
        }, {
            'name': '工程材料費', 'groups': [{'name': '五金 / 工具', 'receipts': [{
                'tax_id_mark': 'V', 'receipt_number': 'CB 20011829', 'amount': 1004,
                'details': ['ABS塑鋼管 ×2', 'ABS等徑三通 ×2', 'ABS閥接頭 ×2', 'ABS90度彎頭 ×16', 'ABS等徑接頭 ×4', 'ABS專用接著劑 ×1']
            }]}]
        }]
    }


def test_engineering_create_calculates_hierarchy_and_preserves_owner_separation(eng_client):
    r = eng_client.post('/api/petty-cash-reports', json=payload())
    assert r.status_code == 201, r.text
    data = r.json()
    assert data['report_type'] == 'engineering'
    assert data['upload_person'] != data['prepared_by']
    assert data['total_amount'] == 1829
    assert data['categories'][0]['subtotal'] == 825
    assert data['categories'][1]['groups'][0]['receipts'][0]['amount'] == 1004


def test_engineering_filter_and_update_are_shared_shell(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload()).json()
    listed = eng_client.get('/api/petty-cash-reports?report_type=engineering')
    assert listed.status_code == 200
    assert listed.json()['items'][0]['report_type'] == 'engineering'
    body = payload(); body['prepared_by'] = '李主任'; body['categories'][0]['groups'][0]['receipts'][0]['amount'] = 200
    updated = eng_client.put(f"/api/petty-cash-reports/{created['id']}", json=body)
    assert updated.status_code == 200, updated.text
    assert updated.json()['prepared_by'] == '李主任'
    assert updated.json()['total_amount'] == 1929


def general_payload():
    return {
        'report_type': 'general', 'start_date': '2026-09-01', 'end_date': '2026-09-04',
        'upload_person': '藍先生', 'prepared_by': '另一位製表人', 'filename_text': '發票',
        'opening_balance': 100, 'status': 'completed', 'entries': [],
    }


def test_general_and_engineering_duplicate_keys_are_isolated(eng_client):
    engineering = eng_client.post('/api/petty-cash-reports', json=payload())
    assert engineering.status_code == 201, engineering.text
    general = eng_client.post('/api/petty-cash-reports', json=general_payload())
    assert general.status_code == 201, general.text
    assert general.json()['report_type'] == 'general'


def test_update_rejects_cross_report_type_payload(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload())
    assert created.status_code == 201, created.text
    rejected = eng_client.put(
        f"/api/petty-cash-reports/{created.json()['id']}", json=general_payload()
    )
    assert rejected.status_code == 409, rejected.text
    unchanged = eng_client.get(f"/api/petty-cash-reports/{created.json()['id']}")
    assert unchanged.status_code == 200
    assert unchanged.json()['report_type'] == 'engineering'


def test_previous_balance_ignores_engineering_reports(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload())
    assert created.status_code == 201, created.text
    previous = eng_client.get(
        '/api/petty-cash-reports/previous-balance',
        params={'upload_person': '藍先生', 'before': '2026-09-05'},
    )
    assert previous.status_code == 200
    assert previous.json()['found'] is False


def test_engineering_duplicate_detection(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload())
    assert created.status_code == 201, created.text
    duplicate = eng_client.post('/api/petty-cash-reports', json=payload())
    assert duplicate.status_code == 409
    assert '工程零用金月報' in duplicate.json()['detail']
    updated = eng_client.put(f"/api/petty-cash-reports/{created.json()['id']}", json=payload())
    assert updated.status_code == 200, updated.text


def test_engineering_concurrent_duplicate_create_is_serialized(eng_client):
    token = eng_client.cookies.get(SESSION_COOKIE)
    clients = []
    for _ in range(2):
        client = TestClient(app_main.app)
        client.cookies.set(SESSION_COOKIE, token)
        clients.append(client)
    barrier = threading.Barrier(2)

    def create(client):
        barrier.wait(timeout=5)
        return client.post('/api/petty-cash-reports', json=payload())

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(create, clients))
    assert sorted(response.status_code for response in responses) == [201, 409]


def test_engineering_export_tax_values_merges_and_filename(eng_client, tmp_path):
    created = eng_client.post('/api/petty-cash-reports', json=payload()).json()
    response = eng_client.get(f"/api/petty-cash-reports/{created['id']}/export.xlsx")
    assert response.status_code == 200
    assert '(0901-0904 發票)藍先生 工程零用金.xlsx' in unquote(response.headers['content-disposition'])
    path = tmp_path / 'out.xlsx'; path.write_bytes(response.content)
    ws = load_workbook(path).active
    assert ws.title == '0901-0904'
    assert ws['C3'].value == 'V'; assert ws['C4'].value == '12345678'
    assert ws['C3'].number_format == '@'; assert ws['C4'].font.name == ws['C3'].font.name
    assert 'C6:C11' in [str(x) for x in ws.merged_cells.ranges]
    assert ws['F14'].value == '=SUM(F5,F12)'
    assert ws['F5'].value == '=SUM(F3:F4)'


def test_engineering_frontend_contract():
    render = Path(__file__).parents[1] / 'static/js/render/petty-cash.js'
    modal = Path(__file__).parents[1] / 'static/js/modals/engineering-petty-cash.js'
    index = Path(__file__).parents[1] / 'static/index.html'
    render_text = render.read_text(encoding='utf-8')
    modal_text = modal.read_text(encoding='utf-8')
    index_text = index.read_text(encoding='utf-8')
    assert "report_type" in render_text and "engineering" in render_text
    assert "pcChooseReportType" in render_text and "engRenderDetail" in render_text
    assert "engAddCategory" in modal_text and "tax_id_mark" in modal_text
    assert "JSON.parse(JSON.stringify(d))" in modal_text
    assert 'engToggleEditorReceipt' in modal_text
    assert 'eng-editor-receipt-toggle' in modal_text
    assert 'engEditorExpandedReceipts' in modal_text
    assert 'id="eng-owner"' in modal_text and 'oninput="engFilenamePreview()"' in modal_text
    assert "engineering-petty-cash.js" in index_text




def test_petty_cash_options_are_separate_and_configured(eng_client):
    engineering = {'report_type': 'engineering', 'option_type': 'category', 'name': '交通費', 'sort_order': 0}
    general = {'report_type': 'general', 'option_type': 'category', 'name': '交通費', 'sort_order': 0}
    assert eng_client.post('/api/petty-cash-options', json=engineering).status_code == 201
    assert eng_client.post('/api/petty-cash-options', json=general).status_code == 201
    assert eng_client.post('/api/petty-cash-options', json=engineering).status_code == 409
    eng = eng_client.get('/api/petty-cash-options?report_type=engineering&option_type=category')
    gen = eng_client.get('/api/petty-cash-options?report_type=general&option_type=category')
    assert [x['name'] for x in eng.json()['items']] == ['交通費']
    assert [x['name'] for x in gen.json()['items']] == ['交通費']
    option_id = eng.json()['items'][0]['id']
    assert eng_client.put(f'/api/petty-cash-options/{option_id}', json={'name': '燃料費'}).status_code == 200
    assert eng_client.delete(f'/api/petty-cash-options/{option_id}').status_code == 200


def test_petty_cash_option_update_rejects_blank_name(eng_client):
    created = eng_client.post('/api/petty-cash-options', json={
        'report_type': 'engineering', 'option_type': 'group', 'name': '油資', 'sort_order': 0,
    })
    assert created.status_code == 201, created.text
    option_id = created.json()['id']
    rejected = eng_client.put(
        f'/api/petty-cash-options/{option_id}', json={'name': '   '}
    )
    assert rejected.status_code == 422, rejected.text


def test_petty_cash_options_reject_invalid_scope(eng_client):
    assert eng_client.get('/api/petty-cash-options?report_type=other&option_type=category').status_code == 400
    assert eng_client.get('/api/petty-cash-options?report_type=general&option_type=other').status_code == 400


def test_engineering_transaction_rolls_back_and_leaves_no_orphans(eng_client):
    body = payload(); body['categories'][0]['groups'][0]['receipts'][0]['details'] = ['']
    response = eng_client.post('/api/petty-cash-reports', json=body)
    assert response.status_code == 422
    conn = app_db.get_db()
    try:
        assert conn.execute("SELECT COUNT(*) FROM petty_cash_reports").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM engineering_expense_details").fetchone()[0] == 0
    finally:
        conn.close()


def test_engineering_validation_and_filename_cases(eng_client):
    assert engineering_sheet_title('2026-09-04', '2026-09-04') == '0904'
    assert engineering_sheet_title('2026-09-01', '2026-09-04') == '0901-0904'
    assert engineering_sheet_title('2026-12-30', '2027-01-03') == '20261230-20270103'
    body = payload()
    body['start_date'] = '2026-09-05'
    body['end_date'] = '2026-09-04'
    assert eng_client.post('/api/petty-cash-reports', json=body).status_code == 422
    body = payload(); body['upload_person'] = '   '
    assert eng_client.post('/api/petty-cash-reports', json=body).status_code == 422
    assert engineering_filename('2026-09-04', '2026-09-04', '藍先生') == '(0904)藍先生 工程零用金.xlsx'
    assert engineering_filename('2026-09-04', '2026-09-04', '藍先生', '發票') == '(0904 發票)藍先生 工程零用金.xlsx'
    assert engineering_filename('2026-12-30', '2027-01-03', '藍先生') == '(20261230-20270103)藍先生 工程零用金.xlsx'


def test_engineering_delete_cascades_all_children(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload()).json()
    report_id = created['id']
    assert eng_client.delete(f'/api/petty-cash-reports/{report_id}').status_code == 200
    conn = app_db.get_db()
    try:
        for table in ('petty_cash_reports', 'engineering_expense_categories', 'engineering_expense_groups', 'engineering_expense_receipts', 'engineering_expense_details'):
            assert conn.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0] == 0
    finally:
        conn.close()


def test_engineering_export_formula_values_are_safe(eng_client, tmp_path):
    body = payload()
    body['categories'][0]['groups'][0]['receipts'][0]['tax_id_mark'] = '=1+1'
    body['categories'][0]['groups'][0]['receipts'][0]['receipt_number'] = '+CMD'
    created = eng_client.post('/api/petty-cash-reports', json=body).json()
    response = eng_client.get(f"/api/petty-cash-reports/{created['id']}/export.xlsx")
    path = tmp_path / 'formula-safe.xlsx'; path.write_bytes(response.content)
    ws = load_workbook(path).active
    assert ws['C3'].value == "'=1+1"
    assert ws['D3'].value == "'+CMD"


def test_engineering_viewer_can_read_but_cannot_write(eng_client):
    created = eng_client.post('/api/petty-cash-reports', json=payload()).json()
    conn = app_db.get_db()
    try:
        conn.execute("INSERT INTO users (username,password_hash,display_name,role) VALUES ('eng-viewer','x','工程檢視者','viewer')")
        conn.commit()
        viewer_id = conn.execute("SELECT id FROM users WHERE username='eng-viewer'").fetchone()['id']
        token = create_session(conn, viewer_id)
    finally:
        conn.close()
    viewer = TestClient(app_main.app)
    viewer.cookies.set(SESSION_COOKIE, token)
    assert viewer.get(f"/api/petty-cash-reports/{created['id']}").status_code == 200
    assert viewer.post('/api/petty-cash-reports', json=payload()).status_code == 403
    assert viewer.delete(f"/api/petty-cash-reports/{created['id']}").status_code == 403


def test_mixed_general_engineering_list_keeps_summary_shapes(eng_client):
    engineering = eng_client.post('/api/petty-cash-reports', json=payload()).json()
    general = {
        'report_type': 'general', 'start_date': '2026-08-26', 'end_date': '2026-09-25',
        'filename_text': '一般', 'upload_person': '王小明', 'prepared_by': '王小明',
        'opening_balance': 0, 'opening_balance_source': 'manual', 'status': 'completed', 'entries': [],
    }
    assert eng_client.post('/api/petty-cash-reports', json=general).status_code == 201
    items = eng_client.get('/api/petty-cash-reports').json()['items']
    by_type = {item['report_type']: item for item in items}
    assert by_type['engineering']['total_amount'] == engineering['total_amount']
    assert 'closing_balance' not in by_type['engineering']
    assert 'closing_balance' in by_type['general']



def test_engineering_safe_filename_keeps_extension_with_max_length():
    """Regression: truncation must never remove the .xlsx extension."""
    raw = engineering_filename(
        '2026-12-30', '2027-01-03', 'O' * 50, 'N' * 50,
    )
    safe = engineering_safe_filename(raw)
    assert len(safe) <= 120
    assert safe.endswith('.xlsx')
    assert '/' not in safe and chr(92) not in safe


def test_engineering_report_dates_are_stored_as_iso_canonical(eng_client):
    """Regression: basic ISO input is canonicalized before SQLite storage."""
    body = payload()
    body['start_date'] = '20260901'
    body['end_date'] = '20260904'
    created = eng_client.post('/api/petty-cash-reports', json=body)
    assert created.status_code == 201, created.text
    report = eng_client.get(f"/api/petty-cash-reports/{created.json()['id']}")
    assert report.status_code == 200, report.text
    assert report.json()['start_date'] == '2026-09-01'
    assert report.json()['end_date'] == '2026-09-04'
