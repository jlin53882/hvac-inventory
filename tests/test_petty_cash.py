# -*- coding: utf-8 -*-
"""零用金月報 API 防回歸測試（2026-09-12）。

涵蓋規格 §43：建立/修改/多上傳人/篩選/計算/上一期帶入/多項目/
交易/重複/驗證/檔名/標題/合併/製表人/渲染/原日報表 intact。
所有資料庫均使用 pytest tmp_path，絕不碰正式 inventory.db。
"""
import app.database as app_db
import main as app_main
import pytest
import threading
from concurrent.futures import ThreadPoolExecutor
from app.services.auth import SESSION_COOKIE, create_session, init_admin_if_missing
from fastapi.testclient import TestClient


@pytest.fixture()
def pc_env(tmp_path, monkeypatch):
    """隔離 DB，回傳可建立已登入 client 的工廠。"""
    monkeypatch.setattr(app_db, "DB_PATH", str(tmp_path / "petty_cash.db"))
    app_db.init_db()
    conn = app_db.get_db()
    try:
        init_admin_if_missing(conn)
    finally:
        conn.close()

    def make_client(username="admin"):
        conn = app_db.get_db()
        try:
            row = conn.execute(
                "SELECT id, role FROM users WHERE username=?", (username,)
            ).fetchone()
            if row is None:
                role = "viewer" if username.startswith("viewer") else "user"
                conn.execute(
                    "INSERT INTO users (username, password_hash, display_name, role)"
                    " VALUES (?, 'x', ?, ?)",
                    (username, username, role),
                )
                conn.commit()
                user_id = conn.execute(
                    "SELECT id FROM users WHERE username=?", (username,)
                ).fetchone()["id"]
            else:
                user_id = row["id"]
            token = create_session(conn, user_id)
        finally:
            conn.close()
        client = TestClient(app_main.app)
        client.cookies.set(SESSION_COOKIE, token)
        return client

    yield make_client


def _entry(date, typ, desc, amount, category="", items=None, sort=0):
    return {
        "entry_date": date, "entry_type": typ, "description": desc,
        "amount": amount, "category": category, "sort_order": sort,
        "items": items or [],
    }


def _item(name, qty, unit, amount, sort=0):
    return {"item_name": name, "qty": qty, "unit": unit, "amount": amount, "sort_order": sort}


def _scenario_a(upload_person="王小明", prepared_by="王小明", status="completed"):
    """規格 §44 情境 A：收入 7334 / 支出 3647 / 餘額 3911。"""
    return {
        "start_date": "2026-08-26", "end_date": "2026-09-25",
        "filename_text": "資材", "upload_person": upload_person,
        "prepared_by": prepared_by, "opening_balance": 224,
        "opening_balance_source": "manual", "status": status,
        "entries": [
            _entry("2026-09-01", "income", "零用金", 3000, sort=0),
            _entry("2026-09-10", "income", "付款申請", 1334, sort=1),
            _entry("2026-09-10", "income", "撥補零用金", 3000, sort=2),
            _entry("2026-09-01", "expense", "畚箕 ×1", 75, "五金", sort=3),
            _entry("2026-09-03", "expense", "寬型美工刀 ×4", 139, "文具", sort=4),
            _entry("2026-09-03", "expense", "卡順噴燈瓦斯 ×30", 1150, "五金", sort=5),
            _entry("2026-09-04", "expense", "瓦斯火口500號 ×5", 810, "五金", sort=6),
            _entry("2026-09-10", "expense", "圓型點火器 ×1", 139, "五金", sort=7),
            _entry("2026-09-10", "expense", "員工福利", 1334, "員工福利", [
                _item("竹炭水", 24, "瓶", 200, 0),
                _item("寶礦力", 24, "瓶", 457, 1),
                _item("咖啡廣場", 24, "瓶", 449, 2),
                _item("四季春無糖", 24, "瓶", 439, 3),
            ], sort=8),
        ],
    }


def _create(client, body):
    r = client.post("/api/petty-cash-reports", json=body)
    assert r.status_code == 201, f"建立失敗 {r.status_code}: {r.text}"
    return r.json()


def _set_user_permissions(username, values):
    conn = app_db.get_db()
    try:
        user = conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()
        for key, value in values.items():
            permission = conn.execute("SELECT id FROM permissions WHERE key=?", (key,)).fetchone()
            conn.execute(
                "INSERT OR REPLACE INTO user_permissions (user_id, permission_id, value, updated_at) "
                "VALUES (?, ?, ?, datetime('now'))",
                (user["id"], permission["id"], int(value)),
            )
        conn.commit()
    finally:
        conn.close()


def _listed_report(client, report_id):
    return next(
        item for item in client.get("/api/petty-cash-reports").json()["items"]
        if item["id"] == report_id
    )


# ---------- 建立 / 計算 ----------

def test_create_computes_totals(pc_env):
    c = pc_env()
    d = _create(c, _scenario_a())
    assert d["totals"] == {
        "opening_balance": 224, "income": 7334, "expense": 3647, "closing_balance": 3911,
    }
    assert len(d["entries"]) == 9
    welfare = [e for e in d["entries"] if e["description"] == "員工福利"][0]
    assert len(welfare["items"]) == 4
    assert welfare["amount_warning"] is not None  # 1545 ≠ 1334，詳見 warning 測試
    assert welfare["detail_total"] == 1545
    assert welfare["difference"] == 211
    assert d["status"] == "completed"


def test_amount_mismatch_warning_saved(pc_env):
    """§25：明細合計 1545 ≠ 支出 1334 → warning 但仍可儲存。"""
    c = pc_env()
    d = _create(c, _scenario_a())
    welfare = [e for e in d["entries"] if e["description"] == "員工福利"][0]
    item_sum = sum(i["amount"] for i in welfare["items"])
    assert item_sum == 1545
    assert welfare["amount_warning"] is not None
    assert "不一致" in welfare["amount_warning"]


def test_update_report_replaces_entries(pc_env):
    """§43-2：修改月報（PUT 全量替換，舊明細不殘留）。"""
    c = pc_env()
    d = _create(c, _scenario_a())
    rid = d["id"]
    body = _scenario_a()
    body["prepared_by"] = "陳主任"
    body["entries"] = [_entry("2026-09-05", "income", "撥補", 5000, sort=0)]
    r = c.put(f"/api/petty-cash-reports/{rid}", json=body)
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["prepared_by"] == "陳主任"
    assert len(got["entries"]) == 1
    assert got["totals"] == {
        "opening_balance": 224, "income": 5000, "expense": 0, "closing_balance": 5224,
    }
    # 舊 4 明細項目已清除（無 orphan）
    conn = app_db.get_db()
    try:
        n = conn.execute("SELECT COUNT(*) FROM petty_cash_entry_items").fetchone()[0]
        assert n == 0
    finally:
        conn.close()


# ---------- 多上傳人 / 篩選 ----------

def test_same_period_different_upload_persons(pc_env):
    """§43-3 / §45：同期間不同上傳人可各自擁有報表。"""
    c = pc_env()
    a = _create(c, _scenario_a("王小明"))
    b = _create(c, _scenario_a("老闆", "老闆"))
    assert a["id"] != b["id"]
    assert b["prepared_by"] == "老闆"


def test_filter_by_upload_person(pc_env):
    c = pc_env()
    _create(c, _scenario_a("王小明"))
    _create(c, _scenario_a("老闆", "老闆", status="draft"))
    r = c.get("/api/petty-cash-reports", params={"upload_person": "老闆"})
    assert r.status_code == 200
    data = r.json()
    assert data["total"] == 1
    assert data["items"][0]["upload_person"] == "老闆"


def test_filter_by_date_overlap(pc_env):
    c = pc_env()
    _create(c, _scenario_a())  # 08/26~09/25
    body = _scenario_a("王小明")
    body.update({"start_date": "2026-09-26", "end_date": "2026-10-25", "filename_text": "十月"})
    body["entries"] = [_entry("2026-10-01", "income", "撥補", 1000, sort=0)]
    _create(c, body)
    r = c.get("/api/petty-cash-reports",
              params={"start_date": "2026-08-01", "end_date": "2026-08-31"})
    assert r.json()["total"] == 1
    r = c.get("/api/petty-cash-reports",
              params={"start_date": "2026-10-01", "end_date": "2026-10-31"})
    assert r.json()["total"] == 1
    r = c.get("/api/petty-cash-reports",
              params={"start_date": "2026-11-01", "end_date": "2026-11-30"})
    assert r.json()["total"] == 0


def test_filter_by_status_and_search(pc_env):
    c = pc_env()
    _create(c, _scenario_a("王小明", status="completed"))
    _create(c, _scenario_a("老闆", "老闆", status="draft"))
    assert c.get("/api/petty-cash-reports", params={"status": "draft"}).json()["total"] == 1
    assert c.get("/api/petty-cash-reports", params={"status": "completed"}).json()["total"] == 1
    assert c.get("/api/petty-cash-reports", params={"search": "資材"}).json()["total"] == 2
    assert c.get("/api/petty-cash-reports", params={"search": "老闆"}).json()["total"] == 1


def test_pagination_total(pc_env):
    c = pc_env()
    for i, person in enumerate(["王小明", "老闆", "陳主任"]):
        _create(c, _scenario_a(person, person))
    r = c.get("/api/petty-cash-reports", params={"page": 1, "page_size": 2})
    data = r.json()
    assert data["total"] == 3 and len(data["items"]) == 2
    r = c.get("/api/petty-cash-reports", params={"page": 2, "page_size": 2})
    assert len(r.json()["items"]) == 1


def test_kpi_counts_full_scope(pc_env):
    """§32：KPI 取自篩選全量，不受分頁影響。"""
    c = pc_env()
    _create(c, _scenario_a("王小明", status="completed"))
    _create(c, _scenario_a("老闆", "老闆", status="draft"))
    kpi = c.get("/api/petty-cash/kpi").json()
    assert kpi == {"total": 2, "completed": 1, "draft": 1}
    kpi = c.get("/api/petty-cash/kpi", params={"upload_person": "老闆"}).json()
    assert kpi == {"total": 1, "completed": 0, "draft": 1}


def test_persons_endpoint(pc_env):
    c = pc_env()
    _create(c, _scenario_a("王小明"))
    pc_env("sarah")  # 啟用中使用者顯示名也應列入候選
    persons = c.get("/api/petty-cash-persons").json()["persons"]
    assert "王小明" in persons
    assert "sarah" in persons


# ---------- 上一期帶入 ----------

def _prev_month_body(person, opening, closing_extra=0):
    body = _scenario_a(person, person, status="completed")
    body.update({"start_date": "2026-07-26", "end_date": "2026-08-25",
                 "filename_text": "七月", "opening_balance": opening})
    body["entries"] = [_entry("2026-08-01", "income", "撥補", 1000 + closing_extra, sort=0)]
    return body


def test_previous_balance_autofill(pc_env):
    c = pc_env()
    prev = _create(c, _prev_month_body("王小明", 250))  # 餘額 250+1000=1250
    assert prev["totals"]["closing_balance"] == 1250
    r = c.get("/api/petty-cash-reports/previous-balance",
              params={"upload_person": "王小明", "before": "2026-08-26"})
    data = r.json()
    assert data["found"] is True
    assert data["opening_balance"] == 1250
    assert data["previous_id"] == prev["id"]


def test_previous_balance_same_person_only(pc_env):
    """§7.2：不同上傳人餘額不相通。"""
    c = pc_env()
    _create(c, _prev_month_body("王小明", 2000))   # 餘額 3000
    _create(c, _prev_month_body("老闆", 7000))     # 餘額 8000
    r = c.get("/api/petty-cash-reports/previous-balance",
              params={"upload_person": "老闆", "before": "2026-08-26"})
    assert r.json()["opening_balance"] == 8000


def test_previous_balance_ignores_draft(pc_env):
    """§27：只找 completed 上一期。"""
    c = pc_env()
    body = _prev_month_body("王小明", 100)
    body["status"] = "draft"
    _create(c, body)
    r = c.get("/api/petty-cash-reports/previous-balance",
              params={"upload_person": "王小明", "before": "2026-08-26"})
    assert r.json()["found"] is False


def test_previous_balance_not_found_message(pc_env):
    c = pc_env()
    r = c.get("/api/petty-cash-reports/previous-balance",
              params={"upload_person": "新人", "before": "2026-08-26"})
    data = r.json()
    assert data["found"] is False
    assert "手動" in data["message"]


def test_opening_manual_preserved_on_reload(pc_env):
    """§27：人工修改後 reload 不被自動帶入覆蓋。"""
    c = pc_env()
    _create(c, _prev_month_body("王小明", 250))  # 上一期餘額 1250
    d = _create(c, _scenario_a())
    rid = d["id"]
    body = _scenario_a()
    body["opening_balance"] = 999
    body["opening_balance_source"] = "manual"
    r = c.put(f"/api/petty-cash-reports/{rid}", json=body)
    assert r.json()["opening_balance"] == 999
    got = c.get(f"/api/petty-cash-reports/{rid}").json()
    assert got["opening_balance"] == 999
    assert got["opening_balance_source"] == "manual"
    assert got["totals"]["closing_balance"] == 999 + 7334 - 3647


# ---------- 交易 / 重複 / 驗證 ----------

def test_transaction_rollback_on_invalid(pc_env):
    """§36：多表寫入任一步失敗 → rollback，不留半筆。"""
    c = pc_env()
    body = _scenario_a()
    body["entries"].append(_entry("2026-10-01", "expense", "超期", 100, sort=9))
    r = c.post("/api/petty-cash-reports", json=body)
    assert r.status_code == 400
    assert c.get("/api/petty-cash-reports").json()["total"] == 0
    conn = app_db.get_db()
    try:
        assert conn.execute("SELECT COUNT(*) FROM petty_cash_entries").fetchone()[0] == 0
    finally:
        conn.close()


def test_duplicate_detection(pc_env):
    c = pc_env()
    d = _create(c, _scenario_a())
    r = c.post("/api/petty-cash-reports", json=_scenario_a())
    assert r.status_code == 409
    assert "已存在" in r.json()["detail"]
    # PUT 自身更新不觸發重複
    body = _scenario_a()
    body["filename_text"] = "資材"
    r = c.put(f"/api/petty-cash-reports/{d['id']}", json=body)
    assert r.status_code == 200


def test_concurrent_duplicate_create_is_serialized(pc_env):
    primary = pc_env()
    token = primary.cookies.get(SESSION_COOKIE)
    clients = []
    for _ in range(2):
        client = TestClient(app_main.app)
        client.cookies.set(SESSION_COOKIE, token)
        clients.append(client)
    barrier = threading.Barrier(2)

    def create(client):
        barrier.wait(timeout=5)
        return client.post('/api/petty-cash-reports', json=_scenario_a())

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(create, clients))
    assert sorted(response.status_code for response in responses) == [201, 409]


def test_entry_date_outside_period_rejected(pc_env):
    c = pc_env()
    body = _scenario_a()
    body["entries"] = [_entry("2026-07-01", "income", "太早", 100, sort=0)]
    assert c.post("/api/petty-cash-reports", json=body).status_code == 400


def test_income_with_items_rejected(pc_env):
    c = pc_env()
    body = _scenario_a()
    body["entries"] = [_entry("2026-09-01", "income", "怪收入", 100,
                              items=[_item("X", 1, "個", 100)], sort=0)]
    assert c.post("/api/petty-cash-reports", json=body).status_code == 400


def test_invalid_inputs_rejected(pc_env):
    c = pc_env()
    base = _scenario_a()
    bad_period = dict(base, start_date="2026-09-25", end_date="2026-08-26")
    assert c.post("/api/petty-cash-reports", json=bad_period).status_code == 400
    bad_date = dict(base, start_date="2026/08/26")
    assert c.post("/api/petty-cash-reports", json=bad_date).status_code == 400
    zero_amt = dict(base)
    zero_amt["entries"] = [_entry("2026-09-01", "income", "零", 0, sort=0)]
    assert c.post("/api/petty-cash-reports", json=zero_amt).status_code == 422
    bad_status = dict(base, status="archived")
    assert c.post("/api/petty-cash-reports", json=bad_status).status_code == 422
    blank_name = dict(base, filename_text="  ")
    assert c.post("/api/petty-cash-reports", json=blank_name).status_code == 422


# ---------- 匯出 ----------

def test_export_filename_title_merges(pc_env, tmp_path):
    c = pc_env()
    d = _create(c, _scenario_a())
    r = c.get(f"/api/petty-cash-reports/{d['id']}/export.xlsx")
    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert "0826-0925.xlsx" in cd  # 安全檔名無斜線
    assert "/" not in cd.split("filename*=")[1].split(";")[0].replace("%2F", "")
    p = tmp_path / "pc.xlsx"
    p.write_bytes(r.content)
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb.active
    assert ws["A1"].value == "08/26~09/25零用金收支明細表"  # Excel 內標題維持斜線
    assert ws["B2"].value == "115年"
    assert ws["E2"].value == 224
    assert ws["E16"].value == "=E2+SUM(D4:D15)-SUM(E4:E15)"
    assert ws["E16"].number_format == "#,##0"
    merges = [str(m) for m in ws.merged_cells.ranges]
    assert "A1:F1" in merges
    # 員工福利 4 明細佔 12~15 列（前 8 筆單列佔 4~11 列），日期/支出/科目跨列合併
    assert "B12:B15" in merges and "E12:E15" in merges and "F12:F15" in merges
    assert "D12:D15" in merges
    assert ws["C12"].value == "竹炭水 24瓶 $200"
    assert ws["C15"].value == "四季春無糖 24瓶 $439"


def test_export_prepared_by_without_uploader(pc_env, tmp_path):
    """§26.4 / §5.1：輸出製表人，不輸出上傳人。"""
    c = pc_env()
    d = _create(c, _scenario_a("老闆", "王小明"))
    r = c.get(f"/api/petty-cash-reports/{d['id']}/export.xlsx")
    assert r.status_code == 200
    p = tmp_path / "pc2.xlsx"
    p.write_bytes(r.content)
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert ws["D18"].value == "製表人:"
    assert ws["E18"].value == "王小明"
    texts = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
    assert not any("老闆" in t for t in texts)
    assert any("王小明" in t for t in texts)


def test_export_formula_injection_safe(pc_env, tmp_path):
    c = pc_env()
    body = _scenario_a()
    body["entries"] = [_entry("2026-09-01", "expense", "=1+1", 50, "+cat", sort=0)]
    d = _create(c, body)
    r = c.get(f"/api/petty-cash-reports/{d['id']}/export.xlsx")
    p = tmp_path / "pc3.xlsx"
    p.write_bytes(r.content)
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert ws["C4"].value == "'=1+1"
    assert ws["F4"].value == "'+cat"


def test_export_sets_last_exported_at(pc_env):
    c = pc_env()
    d = _create(c, _scenario_a())
    assert d["last_exported_at"] == ""
    c.get(f"/api/petty-cash-reports/{d['id']}/export.xlsx")
    got = c.get(f"/api/petty-cash-reports/{d['id']}").json()
    assert got["last_exported_at"] != ""


# ---------- 權限 / 刪除 ----------

def test_unauthorized_401(pc_env):
    c = TestClient(app_main.app)
    assert c.get("/api/petty-cash-reports").status_code == 401
    assert c.get("/api/petty-cash/kpi").status_code == 401
    assert c.post("/api/petty-cash-reports", json=_scenario_a()).status_code == 401
    assert c.get("/api/petty-cash-reports/1").status_code == 401
    assert c.get("/api/petty-cash-reports/1/export.xlsx").status_code == 401


def test_owner_only_edit_delete(pc_env):
    """§30：authorization 看登入者；上傳人只是報表主體。"""
    admin = pc_env()
    sarah = pc_env("sarah")
    boss = pc_env("boss")
    d = _create(sarah, _scenario_a("老闆", "王小明"))  # sarah 幫老闆建
    rid = d["id"]
    assert boss.put(f"/api/petty-cash-reports/{rid}", json=_scenario_a()).status_code == 403
    assert boss.delete(f"/api/petty-cash-reports/{rid}").status_code == 403
    viewer = pc_env("viewer1")
    assert viewer.get("/api/petty-cash-reports").status_code == 200  # 登入可查
    assert admin.delete(f"/api/petty-cash-reports/{rid}").status_code == 200  # 全域刪除


def test_delete_cascade_no_orphans(pc_env):
    c = pc_env()
    d = _create(c, _scenario_a())
    rid = d["id"]
    assert c.delete(f"/api/petty-cash-reports/{rid}").status_code == 200
    assert c.get(f"/api/petty-cash-reports/{rid}").status_code == 404
    conn = app_db.get_db()
    try:
        assert conn.execute("SELECT COUNT(*) FROM petty_cash_entries").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM petty_cash_entry_items").fetchone()[0] == 0
    finally:
        conn.close()


# ---------- 原功能 intact ----------

def test_daily_signed_reports_still_work(pc_env, tmp_path, monkeypatch):
    """§43-27：每日簽名日報表未被破壞。"""
    import app.routes.signed_reports as signed_reports
    monkeypatch.setattr(signed_reports, "STATIC_DIR", str(tmp_path / "static"))
    c = pc_env()
    r = c.post(
        "/api/signed-reports",
        data={"report_date": "2026-09-07", "uploader_name": "王小明", "note": "已簽回"},
        files={"file": ("daily.pdf", b"%PDF-signed", "application/pdf")},
    )
    assert r.status_code == 200, r.text
    assert c.get("/api/signed-reports").json()["total"] == 1


def test_capabilities_match_permissions_and_scope(pc_env):
    client = pc_env("user")
    report = _create(client, _scenario_a())
    report_id = report["id"]

    summary = _listed_report(client, report_id)
    assert summary["can_edit"] is True
    assert summary["can_delete"] is True
    detail = client.get(f"/api/petty-cash-reports/{report_id}").json()
    assert detail["can_edit"] is True
    assert detail["can_delete"] is True

    _set_user_permissions("user", {"petty-cash-edit": 0})
    after_edit_off = _listed_report(client, report_id)
    assert after_edit_off["can_edit"] is False
    assert after_edit_off["can_delete"] is True
    after_edit_off_detail = client.get(f"/api/petty-cash-reports/{report_id}").json()
    assert after_edit_off_detail["can_edit"] is False
    assert after_edit_off_detail["can_delete"] is True
    assert client.put(f"/api/petty-cash-reports/{report_id}", json=_scenario_a()).status_code == 403

    _set_user_permissions("user", {"petty-cash-delete": 0, "petty-cash-delete-all": 0})
    after_delete_off = _listed_report(client, report_id)
    assert after_delete_off["can_edit"] is False
    assert after_delete_off["can_delete"] is False
    after_delete_off_detail = client.get(f"/api/petty-cash-reports/{report_id}").json()
    assert after_delete_off_detail["can_edit"] is False
    assert after_delete_off_detail["can_delete"] is False
    assert client.delete(f"/api/petty-cash-reports/{report_id}").status_code == 403


def test_non_owner_edit_requires_capability_and_global_scope(pc_env):
    owner = pc_env("owner-a")
    non_owner = pc_env("other-b")
    report_id = _create(owner, _scenario_a())["id"]

    assert _listed_report(non_owner, report_id)["can_edit"] is False
    assert non_owner.get(f"/api/petty-cash-reports/{report_id}").json()["can_edit"] is False
    assert non_owner.put(f"/api/petty-cash-reports/{report_id}", json=_scenario_a()).status_code == 403

    _set_user_permissions("other-b", {"petty-cash-delete-all": 1})
    assert _listed_report(non_owner, report_id)["can_edit"] is True
    assert non_owner.get(f"/api/petty-cash-reports/{report_id}").json()["can_edit"] is True
    assert non_owner.put(f"/api/petty-cash-reports/{report_id}", json=_scenario_a()).status_code == 200

    _set_user_permissions("other-b", {"petty-cash-edit": 0})
    assert _listed_report(non_owner, report_id)["can_edit"] is False
    assert non_owner.put(f"/api/petty-cash-reports/{report_id}", json=_scenario_a()).status_code == 403


def test_non_owner_delete_requires_capability_and_global_scope(pc_env):
    owner = pc_env("owner-a")
    non_owner = pc_env("other-b")
    first_report_id = _create(owner, _scenario_a())["id"]

    assert _listed_report(non_owner, first_report_id)["can_delete"] is False
    assert non_owner.delete(f"/api/petty-cash-reports/{first_report_id}").status_code == 403

    _set_user_permissions("other-b", {"petty-cash-delete-all": 1})
    assert _listed_report(non_owner, first_report_id)["can_delete"] is True
    assert non_owner.delete(f"/api/petty-cash-reports/{first_report_id}").status_code == 200

    second_report_id = _create(owner, _scenario_a())["id"]
    _set_user_permissions("other-b", {"petty-cash-delete": 0})
    assert _listed_report(non_owner, second_report_id)["can_delete"] is False
    assert non_owner.get(f"/api/petty-cash-reports/{second_report_id}").json()["can_delete"] is False
    assert non_owner.delete(f"/api/petty-cash-reports/{second_report_id}").status_code == 403
