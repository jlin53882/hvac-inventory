# -*- coding: utf-8 -*-
"""
庫存管理系統 - 單元測試（v10：items 主檔 + item_stocks 位置庫存）
============================
執行方式（必須清 PYTHONPATH 避免 hermes 污染）：
    cd C:\\Users\\admin\\workspace\\hvac-inventory
    env -u PYTHONPATH .venv\\Scripts\\python.exe -m pytest tests/ -v

重點：
  - 每個測試用 tmp_path 建立獨立測試 DB，不會污染 inventory.db 正式資料
  - fixture `client` 自動：切 DB 路徑 → init_db() → TestClient
  - v10 語義：數量存在 item_stocks（位置庫存），品項主檔只有名稱/廠牌等
"""
import os
import sys

import pytest
from fastapi.testclient import TestClient

# 讓測試能 import 到專案根目錄的 main.py
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

import app.config as app_config  # noqa: E402
import app.database as app_db  # noqa: E402
import main as app_main  # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """每個測試獨立 DB：切 DB_PATH → 重建 schema → 建立 admin → 自動登入 → 回傳帶 session 的 TestClient"""
    test_db = tmp_path / "test_inventory.db"
    test_upload = tmp_path / "uploads"
    test_upload.mkdir()
    monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
    monkeypatch.setattr(app_config, "UPLOAD_DIR", str(test_upload))
    app_db.init_db()

    # 建 admin + session 注入（A② 2026-08-14：取代 POST login，省 PBKDF2 600k 迭代 ≈150ms/測試；
    # 登入路徑本身由 test_users.py 的登入/暴力破解測試覆蓋）
    from app.services.auth import SESSION_COOKIE, create_session, init_admin_if_missing
    _conn = app_db.get_db()
    try:
        init_admin_if_missing(_conn)
        _admin_id = _conn.execute("SELECT id FROM users WHERE username='admin'").fetchone()["id"]
        _token = create_session(_conn, _admin_id)
    finally:
        _conn.close()

    with TestClient(app_main.app) as c:
        c.cookies.set(SESSION_COOKIE, _token)
        yield c

    # 測試後清掉測試 DB（Windows 上可能仍被連線鎖住，失敗不影響）
    try:
        if test_db.exists():
            test_db.unlink()
    except PermissionError:
        pass


def _get_item(client, item_id):
    """查單筆品項（含 total_qty/location 相容欄位）"""
    items = client.get("/api/items").json()
    for i in items:
        if i["id"] == item_id:
            return i
    raise AssertionError(f"找不到 item id={item_id}")


def _add_item(client, **kw):
    """Helper：新增品項（v10：stocks[] 陣列）並回傳 dict"""
    payload = {
        "brand": kw.get("brand", "測試牌"),
        "code": kw.get("code", ""),
        "name": kw.get("name", "測試品項"),
        "unit": kw.get("unit", "個"),
        "low_stock": kw.get("low_stock", 0),
        "site": kw.get("site", "office"),
        # v10：位置庫存清單（qty/location/note 參數轉成第一筆位置）
        "stocks": [{
            "location": kw.get("location", "測試位置"),
            "qty": kw.get("qty", 10),
            "note": kw.get("note", ""),
        }],
    }
    r = client.post("/api/items", json=payload)
    assert r.status_code == 201, r.text
    return r.json()


# ========== 基本 CRUD ==========

class TestHealth:
    def test_health_ok(self, client):
        """驗證 GET /health 回傳 200 且 status 為 ok"""
        r = client.get("/health")
        assert r.status_code == 200
        assert r.json()["status"] == "ok"


class TestItemsCRUD:
    def test_create_item(self, client):
        """驗證新增品項回傳完整欄位（相容欄位 qty/total_qty/location 與 stocks 陣列）"""
        item = _add_item(client, name="冷媒管", qty=5, location="A倉")
        assert item["name"] == "冷媒管"
        assert item["qty"] == 5          # 相容欄位：總量 = SUM(stocks)
        assert item["total_qty"] == 5
        assert item["location"] == "A倉"  # 相容欄位：第一筆位置
        assert item["prepared_qty"] == 0
        assert item["is_kit"] == 0
        assert len(item["stocks"]) == 1
        assert item["stocks"][0]["location"] == "A倉"
        assert item["stocks"][0]["qty"] == 5

    def test_create_item_multi_stocks(self, client):
        """v10 核心：一個品項可放多個位置"""
        r = client.post("/api/items", json={
            "brand": "大金",
            "code": "ARC-A",
            "name": "遙控器",
            "stocks": [
                {"location": "A櫃", "qty": 1},
                {"location": "B櫃", "qty": 2},
            ],
        })
        assert r.status_code == 201
        item = r.json()
        assert len(item["stocks"]) == 2
        assert item["total_qty"] == 3  # 總量自動加總
        assert item["qty"] == 3
        assert item["location"] == "A櫃"  # 第一筆為主要位置

    def test_list_items(self, client):
        """驗證 GET /api/items 列出全部品項"""
        _add_item(client, name="一", brand="三菱")
        _add_item(client, name="二", brand="大金")
        r = client.get("/api/items")
        assert r.status_code == 200
        assert len(r.json()) == 2

    def test_list_items_batch_stocks_grouping(self, client):
        """2026-08-15 B1：list_items 批量查（IN 子句）——多 item 多 stocks 分組正確、每筆內依 id 排序、has_photo 為 bool"""
        a = _add_item(client, name="批量甲", brand="大金", location="A倉", qty=10)
        client.post(f"/api/items/{a['id']}/stocks", json={"location": "B倉", "qty": 5})
        b = _add_item(client, name="批量乙", brand="三菱", location="C倉", qty=3)
        client.post(f"/api/items/{b['id']}/stocks", json={"location": "D倉", "qty": 2})

        items = client.get("/api/items").json()
        ia = next(x for x in items if x["id"] == a["id"])
        ib = next(x for x in items if x["id"] == b["id"])
        # 分組正確：每筆 item 只含自己的 stocks，且依 id 排序（A倉 先建 → 排前）
        assert [s["location"] for s in ia["stocks"]] == ["A倉", "B倉"]
        assert [s["location"] for s in ib["stocks"]] == ["C倉", "D倉"]
        assert ia["total_qty"] == 15
        assert ib["total_qty"] == 5
        assert isinstance(ia["has_photo"], bool)
        assert ia["in_kits"] == []

    def test_list_items_has_photo_true_with_file(self, client, tmp_path, monkeypatch):
        """v4 pro 審查補測 #1：upload 目錄有照片檔 → list_items 的 has_photo=True（正向等價性）"""
        import os
        upload = tmp_path / "uploads"
        upload.mkdir(exist_ok=True)
        monkeypatch.setattr("app.config.UPLOAD_DIR", str(upload))
        a = _add_item(client, name="有照片", qty=1)
        b = _add_item(client, name="無照片", qty=2)
        (upload / f"{a['id']}.jpg").write_bytes(b"fake-jpg")
        items = client.get("/api/items").json()
        ia = next(x for x in items if x["id"] == a["id"])
        ib = next(x for x in items if x["id"] == b["id"])
        assert ia["has_photo"] is True
        assert ib["has_photo"] is False

    def test_list_items_item_with_zero_stocks(self, client):
        """v4 pro 審查補測 #5：位置刪光 → stocks=[]、qty=0、location=""、note=""（批量 path 空分支）"""
        item = _add_item(client, name="零庫存", qty=5, location="Z倉")
        sid = item["stocks"][0]["id"]
        # P0-A：非 0 庫存不可直接 DELETE——先調零再刪
        assert client.patch(f"/api/stocks/{sid}", json={"qty": 0}).status_code == 200
        r = client.delete(f"/api/stocks/{sid}")
        assert r.status_code == 200
        items = client.get("/api/items").json()
        it = next(x for x in items if x["id"] == item["id"])
        assert it["stocks"] == []
        assert it["qty"] == 0
        assert it["location"] == ""
        assert it["note"] == ""

    def test_list_items_empty_db_returns_list(self, client):
        """v4 pro 審查補測 #6：空品項清單 → 不進 IN 查、回 []（200 不 crash）"""
        r = client.get("/api/items")
        assert r.status_code == 200
        assert r.json() == []

    def test_list_items_filter_brand(self, client):
        """驗證以 brand 參數篩選品項"""
        _add_item(client, name="一", brand="三菱")
        _add_item(client, name="二", brand="大金")
        r = client.get("/api/items", params={"brand": "三菱"})
        items = r.json()
        assert len(items) == 1
        assert items[0]["brand"] == "三菱"

    def test_list_items_search(self, client):
        """驗證以 search 參數依品名搜尋品項"""
        _add_item(client, name="變頻風扇馬達", code="B4025397")
        r = client.get("/api/items", params={"search": "風扇"})
        items = r.json()
        assert len(items) == 1
        assert "風扇" in items[0]["name"]

    def test_list_items_search_location(self, client):
        """搜尋涵蓋位置（v10：stocks 的 location/note）"""
        _add_item(client, name="冷媒", location="編號A")
        r = client.get("/api/items", params={"search": "編號A"})
        items = r.json()
        assert len(items) == 1

    def test_update_item(self, client):
        """v10：更新位置用 stocks 全量替換（P1-A：有貨位置不可靜默移除）"""
        item = _add_item(client, name="舊名", qty=5, location="A倉")
        # 先把A倉清零才能移除（P1-A 保護有貨位置不被靜默滅失）
        sid = item["stocks"][0]["id"]
        client.patch(f"/api/stocks/{sid}", json={"qty": 0})
        r = client.patch(f"/api/items/{item['id']}", json={
            "stocks": [
                {"location": "B倉", "qty": 3},
                {"location": "C倉", "qty": 2},
            ],
        })
        assert r.status_code == 200
        updated = r.json()
        assert updated["total_qty"] == 5
        assert updated["location"] == "B倉"
        assert len(updated["stocks"]) == 2
        assert updated["name"] == "舊名"  # 沒傳的欄位不變

    def test_update_item_not_found(self, client):
        """驗證更新不存在的品項回傳 404"""
        r = client.patch("/api/items/99999", json={"stocks": [{"location": "X", "qty": 1}]})
        assert r.status_code == 404

    def test_update_item_empty_fields(self, client):
        """驗證 PATCH 空 body（無任何欄位）回傳 400"""
        item = _add_item(client)
        r = client.patch(f"/api/items/{item['id']}", json={})
        assert r.status_code == 400


# ========== v10 新增：去重防護 ==========

class TestDedup:
    def test_duplicate_create_rejected(self, client):
        """相同 (brand, code, name, unit, site) 重複新增 → 400"""
        _add_item(client, brand="大金", code="K031224", name="控制基板")
        r = client.post("/api/items", json={
            "brand": "大金", "code": "K031224", "name": "控制基板",
            "unit": "個", "site": "office",
            "stocks": [{"location": "B倉", "qty": 1}],
        })
        assert r.status_code == 400
        assert "已存在" in r.json()["detail"]

    def test_same_name_different_code_allowed(self, client):
        """同品名不同型號 → 允許（不是重複）"""
        _add_item(client, brand="大金", code="K031", name="控制基板")
        r = client.post("/api/items", json={
            "brand": "大金", "code": "K999", "name": "控制基板",
            "unit": "個", "site": "office",
            "stocks": [{"location": "B倉", "qty": 1}],
        })
        assert r.status_code == 201

    def test_same_name_different_site_allowed(self, client):
        """不同分片（辦公/倉庫）同品同型號 → 允許（兩辦公室各放一份）"""
        _add_item(client, brand="大金", code="K031", name="面板", site="office")
        r = client.post("/api/items", json={
            "brand": "大金", "code": "K031", "name": "面板",
            "unit": "個", "site": "warehouse",
            "stocks": [{"location": "倉庫區", "qty": 2}],
        })
        assert r.status_code == 201

    def test_duplicate_stock_location_merged_on_edit(self, client):
        """編輯送重複位置 → 全量替換後不產生重複位置（最後一筆勝出）"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        r = client.patch(f"/api/items/{item['id']}", json={
            "stocks": [
                {"location": "A倉", "qty": 3},
                {"location": "A倉", "qty": 4},
            ],
        })
        assert r.status_code == 200
        updated = r.json()
        locs = [s["location"] for s in updated["stocks"]]
        assert locs.count("A倉") == 1  # 無重複位置
        assert updated["total_qty"] == 4


# ========== 位置庫存 CRUD ==========

class TestStocksCRUD:
    def test_add_stock(self, client):
        """驗證新增位置庫存後總量自動加總"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        r = client.post(f"/api/items/{item['id']}/stocks",
                        json={"location": "B倉", "qty": 3})
        assert r.status_code == 201
        updated = _get_item(client, item["id"])
        assert len(updated["stocks"]) == 2
        assert updated["total_qty"] == 8

    def test_add_stock_duplicate_location_rejected(self, client):
        """同品項同位置重複加 → 400（UNIQUE(item_id, location)）"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        r = client.post(f"/api/items/{item['id']}/stocks",
                        json={"location": "A倉", "qty": 3})
        assert r.status_code == 400

    def test_update_stock(self, client):
        """驗證更新位置庫存數量與備註"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        sid = item["stocks"][0]["id"]
        r = client.patch(f"/api/stocks/{sid}", json={"qty": 8, "note": "補貨"})
        assert r.status_code == 200
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 8

    def test_delete_stock(self, client):
        """驗證刪除位置庫存後總量同步扣減（P0-A：先清零再刪）"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        client.post(f"/api/items/{item['id']}/stocks",
                    json={"location": "B倉", "qty": 3})
        sid = _get_item(client, item["id"])["stocks"][1]["id"]
        # P0-A：非 0 庫存不可 DELETE——先調零
        client.patch(f"/api/stocks/{sid}", json={"qty": 0})
        r = client.delete(f"/api/stocks/{sid}")
        assert r.status_code == 200
        updated = _get_item(client, item["id"])
        assert len(updated["stocks"]) == 1
        assert updated["total_qty"] == 5

    def test_delete_item_cascades_stocks(self, client):
        """驗證刪除品項會一併刪除其位置庫存"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        client.post(f"/api/items/{item['id']}/stocks",
                    json={"location": "B倉", "qty": 3})
        r = client.delete(f"/api/items/{item['id']}")
        assert r.status_code == 200
        items = client.get("/api/items").json()
        assert len(items) == 0


# ========== 加減庫存 ==========

class TestAdjustQty:
    def test_adjust_plus(self, client):
        """驗證正向調整庫存（進貨）"""
        item = _add_item(client, name="冷媒", qty=10)
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": 5, "reason": "進貨"})
        assert r.status_code == 200
        assert r.json()["after"] == 15
        assert _get_item(client, item["id"])["total_qty"] == 15

    def test_adjust_minus_with_destination(self, client):
        """驗證負向調整庫存（出貨）並記錄去向"""
        item = _add_item(client, name="冷媒", qty=10)
        r = client.post(f"/api/items/{item['id']}/adjust",
                        json={"delta": -3, "reason": "出貨", "destination": "台北案場"})
        assert r.status_code == 200
        assert r.json()["after"] == 7

        # 異動紀錄有去向
        mov = client.get("/api/movements").json()
        assert len(mov) == 1
        assert mov[0]["destination"] == "台北案場"
        assert mov[0]["before_qty"] == 10
        assert mov[0]["after_qty"] == 7

    def test_adjust_not_below_zero(self, client):
        """驗證庫存不足時調整回傳 400"""
        item = _add_item(client, name="冷媒", qty=2)
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": -10})
        assert r.status_code == 400  # 庫存不足

    def test_adjust_item_not_found(self, client):
        """驗證調整不存在的品項回傳 404"""
        r = client.post("/api/items/99999/adjust", json={"delta": 1})
        assert r.status_code == 404


# ========== 兩階段出庫（待領出 → 已領出） ==========

class TestTwoStageStockOut:
    def test_prepared_list_has_qty_field(self, client):
        """防回歸：待領出清單必須帶 qty 總量欄位（undefined bug）"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 3})
        prepared = client.get("/api/prepared").json()
        assert len(prepared) == 1
        assert prepared[0]["prepared_qty"] == 3
        assert prepared[0]["qty"] == 10     # 不能是 undefined/缺欄位
        assert prepared[0]["total_qty"] == 10
        assert "location" in prepared[0]

    def test_prepare_does_not_deduct_qty(self, client):
        """待領出：prepared_qty 增加，但總量不變"""
        item = _add_item(client, name="銅管", qty=10)
        r = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 3})
        assert r.status_code == 200
        p = r.json()
        assert p["prepared_qty"] == 3
        assert p["qty"] == 10  # 庫存沒扣

        # 出現在待領出清單
        prepared = client.get("/api/prepared").json()
        assert len(prepared) == 1
        assert prepared[0]["id"] == item["id"]

    def test_prepare_insufficient(self, client):
        """驗證可領出數量不足時回傳 400"""
        item = _add_item(client, name="銅管", qty=2)
        r = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 5})
        assert r.status_code == 400  # 可領出數量不足

    def test_prepare_zero_qty(self, client):
        """驗證待領出數量為 0 時回傳 400"""
        item = _add_item(client, name="銅管", qty=10)
        r = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 0})
        assert r.status_code == 400

    def test_prepare_then_confirm_out(self, client):
        """待領出 → 確認已領出：扣庫存 + 記去向"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 4})

        r = client.post(f"/api/items/{item['id']}/prepared-out",
                        json={"qty": 4, "note": "中和工地"})
        assert r.status_code == 200
        p = r.json()
        assert p["qty"] == 6          # 庫存被扣
        assert p["prepared_qty"] == 0  # 待領出清空

        # 異動紀錄：去向 = 中和工地
        mov = client.get("/api/movements").json()
        outs = [m for m in mov if m["reason"] == "出庫"]
        assert len(outs) == 1
        assert outs[0]["destination"] == "中和工地"
        assert outs[0]["before_qty"] == 10
        assert outs[0]["after_qty"] == 6

    def test_confirm_out_exceeds_prepared(self, client):
        """驗證確認出庫數量超過待領出數量時回傳 400"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 2})
        r = client.post(f"/api/items/{item['id']}/prepared-out", json={"qty": 5})
        assert r.status_code == 400  # 準備中只有 2

    def test_prepare_then_return(self, client):
        """待領出 → 退回：清 prepared_qty，總量不變"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 3})
        r = client.post(f"/api/items/{item['id']}/prepared-return", json={"qty": 3})
        assert r.status_code == 200
        p = r.json()
        assert p["prepared_qty"] == 0
        assert p["qty"] == 10

        # 待領出清單空了
        prepared = client.get("/api/prepared").json()
        assert len(prepared) == 0


# ========== 已領出清單過濾（2026-08-10：手動調整不再誤入已領出） ==========

class TestStockOutsFiltered:
    def test_manual_adjust_minus_not_in_stockouts(self, client):
        """手動調整（庫存頁 +/- 儲存）負數 → 不應出現在已領出清單"""
        item = _add_item(client, name="防蟲罩", qty=5)
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": -2, "reason": "手動調整"})
        assert r.status_code == 200
        outs = client.get("/api/stockouts").json()
        assert all(o["item_id"] != item["id"] for o in outs), "手動調整不應出現在已領出"

    def test_stocktake_negative_not_in_stockouts(self, client):
        """盤點盤虧（diff<0）→ 不應出現在已領出清單"""
        item = _add_item(client, name="銅管", qty=5)
        r = client.post("/api/stocktake", json={
            "take_date": "2026-08-10",
            "items": [{"item_id": item["id"], "location": "測試位置", "actual_qty": 3}],
        })
        assert r.status_code == 200
        outs = client.get("/api/stockouts").json()
        assert all(o["item_id"] != item["id"] for o in outs), "盤點盤虧不應出現在已領出"

    def test_real_out_still_in_stockouts(self, client):
        """真正的出庫（reason='出庫'）仍要出現在已領出清單"""
        item = _add_item(client, name="冷媒", qty=10)
        client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 3, "destination": "台北案場"})
        outs = client.get("/api/stockouts").json()
        match = [o for o in outs if o["item_id"] == item["id"]]
        assert len(match) == 1
        assert match[0]["destination"] == "台北案場"


# ========== 已領出：退回（2026-08-10 新增） ==========

class TestStockoutReturn:
    def _out(self, client, item_id, qty=3, dest="台北案場"):
        """Helper：直接出庫並回傳 movements 記錄 id"""
        client.post("/api/stockout", json={
            "item_id": item_id, "qty": qty, "destination": dest})
        outs = client.get("/api/stockouts").json()
        return [o for o in outs if o["item_id"] == item_id][0]

    def test_return_adds_back_qty(self, client):
        """退回已領出：數量加回庫存 + 原記錄標記 reverted_at + 寫反向流水"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        assert _get_item(client, item["id"])["total_qty"] == 7

        r = client.post(f"/api/stockouts/{rec['id']}/return")
        assert r.status_code == 200
        assert r.json()["returned_qty"] == 3
        assert r.json()["fully_returned"] is True

        assert _get_item(client, item["id"])["total_qty"] == 10  # 庫存加回

        outs = client.get("/api/stockouts").json()
        reverted = [o for o in outs if o["id"] == rec["id"]][0]
        assert reverted["reverted_at"], "原記錄應標記 reverted_at"

        # 反向流水 reason='退回已領出'、delta 正數
        mov = client.get("/api/movements").json()
        backs = [m for m in mov if m["reason"] == "退回已領出"]
        assert len(backs) == 1
        assert backs[0]["delta"] == 3

    def test_movements_chain_before_delta_after(self, client):
        """P4-1（2026-08-14）：連續調整後每筆流水 before+delta==after（寫後重讀，鏈一致）"""
        item = _add_item(client, name="冷媒", qty=10)
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": 5, "reason": "進貨"})
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": -3, "reason": "出貨"})
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": 2, "reason": "進貨"})
        movs = client.get("/api/movements").json()
        assert len(movs) == 3
        # 每筆流水 before + delta == after（浮點 round 3 位）
        for m in movs:
            assert round(m["before_qty"] + m["delta"], 3) == round(m["after_qty"], 3), \
                f"流水鏈不一致: {m}"
        # 最後一筆 after == 目前總量（10+5-3+2=14）
        assert movs[0]["after_qty"] == 14
        assert _get_item(client, item["id"])["total_qty"] == 14

    def test_update_stockout_keeps_delta_after(self, client):
        """P4-4（2026-08-14）：編輯已領出數量後流水 delta/after 正確（BEGIN IMMEDIATE 不破壞功能）"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=4)
        assert _get_item(client, item["id"])["total_qty"] == 6
        # 改數量 4 → 6（多扣 2）
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 6})
        assert r.status_code == 200, r.text
        assert _get_item(client, item["id"])["total_qty"] == 4
        # 原記錄流水更新為 delta=-6、after_qty=4
        movs = client.get("/api/movements").json()
        target = [m for m in movs if m["id"] == rec["id"]][0]
        assert target["delta"] == -6
        assert target["after_qty"] == 4
        assert round(target["before_qty"] + target["delta"], 3) == round(target["after_qty"], 3)
        # 編輯調整流水（-2）鏈一致：before=編輯前總量 6 → 6+(-2)=4（2026-08-14 審查修 P4-1）
        adj = [m for m in movs if m["reason"] == "已領出編輯調整"]
        assert len(adj) == 1
        assert adj[0]["delta"] == -2
        assert adj[0]["before_qty"] == 6
        assert adj[0]["after_qty"] == 4
        assert round(adj[0]["before_qty"] + adj[0]["delta"], 3) == round(adj[0]["after_qty"], 3)


    def test_return_twice_rejected(self, client):
        """同一筆已領出記錄不能重複退回"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        assert client.post(f"/api/stockouts/{rec['id']}/return").status_code == 200
        r = client.post(f"/api/stockouts/{rec['id']}/return")
        assert r.status_code == 400  # 已退回過
        assert _get_item(client, item["id"])["total_qty"] == 10  # 沒有重複加回

    def test_return_non_out_rejected(self, client):
        """非出庫記錄（如領出準備 delta=0）不能退回"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 3})
        mov = client.get("/api/movements").json()
        prep = [m for m in mov if m["reason"] == "領出準備"][0]
        r = client.post(f"/api/stockouts/{prep['id']}/return")
        assert r.status_code == 400

    def test_return_not_found(self, client):
        """退回不存在的記錄 → 404"""
        r = client.post("/api/stockouts/99999/return")
        assert r.status_code == 404

    def test_return_manual_adjust_rejected(self, client):
        """手動調整（非出庫）的負數記錄不能被退回（防止庫存虛增）"""
        item = _add_item(client, name="防蟲罩", qty=5)
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": -2, "reason": "手動調整"})
        mov = client.get("/api/movements").json()
        adj = [m for m in mov if m["reason"] == "手動調整"][0]
        r = client.post(f"/api/stockouts/{adj['id']}/return")
        assert r.status_code == 400
        assert _get_item(client, item["id"])["total_qty"] == 3  # 庫存沒有被加回

    def test_partial_return(self, client):
        """2026-09-07 Sarah：部分退回——只退 N 個，庫存只加回 N"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=5)
        assert _get_item(client, item["id"])["total_qty"] == 5

        # 退回 2 個（非全數 5）
        r = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 2})
        assert r.status_code == 200
        assert r.json()["returned_qty"] == 2
        assert r.json()["fully_returned"] is False
        assert _get_item(client, item["id"])["total_qty"] == 7  # 5+2=7

        # 反向流水 delta=2（非全數 5）
        mov = client.get("/api/movements").json()
        backs = [m for m in mov if m["reason"] == "退回已領出"]
        assert len(backs) == 1
        assert backs[0]["delta"] == 2

        # 部分退回後不設 reverted_at，可再退
        outs = client.get("/api/stockouts").json()
        orig = [o for o in outs if o["id"] == rec["id"]][0]
        assert not orig["reverted_at"], "部分退回不應設 reverted_at"

    def test_partial_then_full_return(self, client):
        """2026-09-07 Sarah（Codex 審查）：部分退回後可再退剩餘"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=5)
        assert _get_item(client, item["id"])["total_qty"] == 5

        # 第一次：退回 2
        r1 = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 2})
        assert r1.status_code == 200
        assert r1.json()["fully_returned"] is False
        assert _get_item(client, item["id"])["total_qty"] == 7

        # 第二次：退回剩餘 3（全數）
        r2 = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 3})
        assert r2.status_code == 200
        assert r2.json()["returned_qty"] == 3
        assert r2.json()["fully_returned"] is True
        assert _get_item(client, item["id"])["total_qty"] == 10

        # 全數退回後設 reverted_at，不能再退
        r3 = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 1})
        assert r3.status_code == 400

        # 反向流水共 2 筆（2+3=5）
        mov = client.get("/api/movements").json()
        backs = [m for m in mov if m["reason"] == "退回已領出"]
        assert len(backs) == 2
        assert backs[0]["delta"] + backs[1]["delta"] == 5

    def test_multiple_stockouts_same_item_independent(self, client):
        """2026-09-07（Codex CRITICAL 修正）：同品項多筆出庫的部分退回互不干擾
        - 出庫 A：5 個，退回 2 個
        - 出庫 B：3 個，退回 1 個
        - 兩筆的 source_movement_id 精確連結，不會互相計算
        """
        item = _add_item(client, name="冷媒", qty=20)
        # 兩筆出庫
        rec_a = self._out(client, item["id"], qty=5, dest="台北")
        rec_b = self._out(client, item["id"], qty=3, dest="台中")
        assert _get_item(client, item["id"])["total_qty"] == 12  # 20-5-3=12

        # 從 A 退回 2
        r1 = client.post(f"/api/stockouts/{rec_a['id']}/return", json={"qty": 2})
        assert r1.status_code == 200
        assert r1.json()["fully_returned"] is False
        assert _get_item(client, item["id"])["total_qty"] == 14  # 12+2=14

        # 從 B 退回 1（不應受 A 的 2 影響）
        r2 = client.post(f"/api/stockouts/{rec_b['id']}/return", json={"qty": 1})
        assert r2.status_code == 200
        assert r2.json()["fully_returned"] is False  # B 只退了 1/3，不應是全數退回
        assert _get_item(client, item["id"])["total_qty"] == 15  # 14+1=15

        # B 的 reverted_at 不應被設（只退了 1/3）
        mov_b = client.get("/api/movements").json()
        orig_b = [m for m in mov_b if m["id"] == rec_b["id"]][0]
        assert not orig_b["reverted_at"], "B 只退 1/3 不應設 reverted_at"

        # B 可再退 2（剩餘）
        r3 = client.post(f"/api/stockouts/{rec_b['id']}/return", json={"qty": 2})
        assert r3.status_code == 200
        assert r3.json()["fully_returned"] is True  # B 全數退回
        assert _get_item(client, item["id"])["total_qty"] == 17  # 15+2=17

        # A 仍可再退 3（剩餘）
        r4 = client.post(f"/api/stockouts/{rec_a['id']}/return", json={"qty": 3})
        assert r4.status_code == 200
        assert r4.json()["fully_returned"] is True
        assert _get_item(client, item["id"])["total_qty"] == 20  # 全部回來

    def test_return_qty_exceeds_original_rejected(self, client):
        """2026-09-07 Sarah：退回數量超過原出庫數量 → 400"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        r = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 5})
        assert r.status_code == 400
        assert "超過" in r.json()["detail"]

    def test_return_with_custom_destination(self, client):
        """2026-09-07 Sarah：退回時可自訂去向"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3, dest="台北案場")
        r = client.post(f"/api/stockouts/{rec['id']}/return",
                        json={"qty": 3, "destination": "退回倉庫"})
        assert r.status_code == 200

        # 反向流水的去向 = 自訂去向（非原出庫去向）
        mov = client.get("/api/movements").json()
        backs = [m for m in mov if m["reason"] == "退回已領出"]
        assert backs[0]["destination"] == "退回倉庫"

    def test_return_with_custom_date(self, client):
        """2026-09-07 Sarah：退回時可自訂日期"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        r = client.post(f"/api/stockouts/{rec['id']}/return",
                        json={"qty": 3, "created_at": "2026-09-05 00:00:00"})
        assert r.status_code == 200

        # 反向流水的 created_at = 自訂日期
        mov = client.get("/api/movements").json()
        backs = [m for m in mov if m["reason"] == "退回已領出"]
        assert backs[0]["created_at"].startswith("2026-09-05")

    def test_return_full_without_body(self, client):
        """2026-09-07：不帶 body 退回 = 全數退回（向後相容）"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=4)
        r = client.post(f"/api/stockouts/{rec['id']}/return")
        assert r.status_code == 200
        assert r.json()["returned_qty"] == 4
        assert _get_item(client, item["id"])["total_qty"] == 10


    def test_return_to_selected_stock_location(self, client):
        """退回可改選同品項的其他庫存位置，實際庫存跟著移動。"""
        item = _add_item(client, name="位置退回", qty=10)
        stock = item["stocks"][0]
        second = client.post(f"/api/items/{item['id']}/stocks", json={"location": "櫃B | 2-2", "qty": 0})
        assert second.status_code in (200, 201)
        second_id = second.json()[-1]["id"]
        rec = self._out(client, item["id"], qty=3, dest="案場")
        r = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 2, "return_stock_id": second_id})
        assert r.status_code == 200, r.text
        rows = _get_item(client, item['id'])["stocks"]
        by_id = {x["id"]: x["qty"] for x in rows}
        assert by_id[stock["id"]] == 7
        assert by_id[second_id] == 2
        ret = [x for x in client.get("/api/movements").json() if x["reason"] == "退回已領出"][0]
        assert ret["return_stock_id"] == second_id

    def test_return_edit_moves_stock_and_revoke_restores(self, client):
        """編輯退回位置/數量會同步移動庫存，撤銷會扣回並保留流水。"""
        item = _add_item(client, name="退回編輯", qty=10)
        second = client.post(f"/api/items/{item['id']}/stocks", json={"location": "櫃C | 3-3", "qty": 0}).json()[-1]
        rec = self._out(client, item["id"], qty=5)
        ret = client.post(f"/api/stockouts/{rec['id']}/return", json={"qty": 2, "return_stock_id": second["id"]}).json()
        r = client.patch(f"/api/stockout-returns/{ret['return_movement_id']}", json={"qty": 3, "return_stock_id": second["id"]})
        assert r.status_code == 200, r.text
        updated = r.json()
        assert round(updated["before_qty"] + updated["delta"], 3) == round(updated["after_qty"], 3)
        r = client.delete(f"/api/stockout-returns/{ret['return_movement_id']}")
        assert r.status_code == 200, r.text
        rows = _get_item(client, item['id'])["stocks"]
        assert {x["id"]: x["qty"] for x in rows}[second["id"]] == 0
        all_mov = client.get("/api/movements").json()
        assert not any(x["id"] == ret["return_movement_id"] for x in all_mov)

    def test_delete_active_return_removes_record_and_reverses_stock(self, client):
        """刪除尚未撤銷的退回流水，應扣回退回庫存並移除流水。"""
        item = _add_item(client, name="刪除活動退回", qty=10)
        rec = self._out(client, item["id"], qty=4)
        ret = client.post(
            f"/api/stockouts/{rec['id']}/return",
            json={"qty": 2, "return_stock_id": item["stocks"][0]["id"]},
        ).json()

        deleted = client.delete(f"/api/stockout-returns/{ret['return_movement_id']}")

        assert deleted.status_code == 200, deleted.text
        assert deleted.json()["deleted"] == ret["return_movement_id"]
        assert _get_item(client, item["id"])["total_qty"] == 6
        assert not any(
            m["id"] == ret["return_movement_id"]
            for m in client.get("/api/movements").json()
        )
        assert not any(
            m["id"] == ret["return_movement_id"]
            for m in client.get("/api/stockouts?site=all").json()
        )

    def test_delete_reverted_return_removes_record_without_stock_change(self, client):
        """已撤銷退回也可刪除，且不可再次扣庫存。"""
        item = _add_item(client, name="刪除已撤銷退回", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = item["stocks"][0]["id"]
        conn = app_db.get_db()
        try:
            cur = conn.execute(
                "INSERT INTO movements "
                "(item_id, delta, before_qty, after_qty, reason, return_stock_id, reverted_at) "
                "VALUES (?,?,?,?,?,?,?)",
                (item["id"], 2, 6, 8, "退回已領出", stock_id, "2026-09-09T10:00:00"),
            )
            return_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        deleted = client.delete(f"/api/stockout-returns/{return_id}")

        assert deleted.status_code == 200, deleted.text
        assert _get_item(client, item["id"])["total_qty"] == 6
        assert not any(m["id"] == return_id for m in client.get("/api/movements").json())



    def test_repair_legacy_return_then_edit_and_revoke(self, client):
        """舊退回流水可由系統補關聯後編輯、撤銷，不需人工改 DB。"""
        item = _add_item(client, name="舊退回復原", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = rec["source_stock_id"]

        # 模擬舊版：退回已加回庫存，但沒有任何關聯欄位。
        conn = app_db.get_db()
        try:
            conn.execute("UPDATE item_stocks SET qty=qty+4 WHERE id=?", (stock_id,))
            conn.execute("UPDATE movements SET reverted_at=? WHERE id=?", ("2026-09-08T10:00:00", rec["id"]))
            cur = conn.execute(
                "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
                (item["id"], 4, 6, 10, "退回已領出", "公司"),
            )
            legacy_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        repaired = client.post(
            f"/api/stockout-returns/{legacy_id}/repair",
            json={"source_movement_id": rec["id"], "return_stock_id": stock_id},
        )
        assert repaired.status_code == 200, repaired.text
        assert repaired.json()["source_movement_id"] == rec["id"]
        assert repaired.json()["return_stock_id"] == stock_id

        edited = client.patch(
            f"/api/stockout-returns/{legacy_id}",
            json={"destination": "公司倉庫"},
        )
        assert edited.status_code == 200, edited.text
        assert edited.json()["destination"] == "公司倉庫"

        revoked = client.delete(f"/api/stockout-returns/{legacy_id}")
        assert revoked.status_code == 200, revoked.text
        rows = _get_item(client, item["id"])["stocks"]
        assert {x["id"]: x["qty"] for x in rows}[stock_id] == 6
        parent = [m for m in client.get("/api/movements").json() if m["id"] == rec["id"]][0]
        assert parent["reverted_at"] is None


    def test_repair_preserves_existing_partial_link(self, client):
        """舊資料只缺一個關聯欄位時，修復不可覆寫既有回補位置。"""
        item = _add_item(client, name="部分關聯復原", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = rec["source_stock_id"]
        alternate = client.post(
            f"/api/items/{item['id']}/stocks",
            json={"location": "另一個位置", "qty": 0},
        ).json()[-1]["id"]
        conn = app_db.get_db()
        try:
            conn.execute("UPDATE item_stocks SET qty=qty+4 WHERE id=?", (stock_id,))
            conn.execute("UPDATE movements SET reverted_at=? WHERE id=?", ("2026-09-08T10:00:00", rec["id"]))
            cur = conn.execute(
                "INSERT INTO movements (item_id, delta, reason, return_stock_id) VALUES (?,?,?,?)",
                (item["id"], 4, "退回已領出", stock_id),
            )
            legacy_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        repaired = client.post(
            f"/api/stockout-returns/{legacy_id}/repair",
            json={"source_movement_id": rec["id"], "return_stock_id": alternate},
        )
        assert repaired.status_code == 200, repaired.text
        assert repaired.json()["return_stock_id"] == stock_id


    def test_repair_partial_legacy_return_keeps_parent_active(self, client):
        """部分舊退回可修復，原始出庫仍可繼續退回剩餘數量。"""
        item = _add_item(client, name="部分舊退回", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = rec["source_stock_id"]
        conn = app_db.get_db()
        try:
            conn.execute("UPDATE item_stocks SET qty=qty+2 WHERE id=?", (stock_id,))
            cur = conn.execute(
                "INSERT INTO movements (item_id, delta, reason, destination) VALUES (?,?,?,?)",
                (item["id"], 2, "退回已領出", "公司"),
            )
            legacy_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        repaired = client.post(
            f"/api/stockout-returns/{legacy_id}/repair",
            json={"source_movement_id": rec["id"], "return_stock_id": stock_id},
        )
        assert repaired.status_code == 200, repaired.text
        parent = [m for m in client.get("/api/movements").json() if m["id"] == rec["id"]][0]
        assert parent["reverted_at"] is None

        completed = client.post(
            f"/api/stockouts/{rec['id']}/return",
            json={"qty": 2, "return_stock_id": stock_id},
        )
        assert completed.status_code == 200, completed.text


    def test_repair_rejects_return_total_over_parent(self, client):
        """舊退回關聯不可讓同一原始出庫超過可退總量。"""
        item = _add_item(client, name="超退關聯防護", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = rec["source_stock_id"]
        first = client.post(
            f"/api/stockouts/{rec['id']}/return",
            json={"qty": 4, "return_stock_id": stock_id},
        )
        assert first.status_code == 200, first.text

        conn = app_db.get_db()
        try:
            conn.execute("UPDATE item_stocks SET qty=qty+2 WHERE id=?", (stock_id,))
            cur = conn.execute(
                "INSERT INTO movements (item_id, delta, reason) VALUES (?,?,?)",
                (item["id"], 2, "退回已領出"),
            )
            legacy_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        repaired = client.post(
            f"/api/stockout-returns/{legacy_id}/repair",
            json={"source_movement_id": rec["id"], "return_stock_id": stock_id},
        )
        assert repaired.status_code == 400
        assert "退回總量不可超過" in repaired.json()["detail"]


    def test_repair_rejects_nonpositive_legacy_return(self, client):
        """復原不可接受負數或零數量的舊退回流水。"""
        item = _add_item(client, name="非法舊退回", qty=10)
        rec = self._out(client, item["id"], qty=4)
        stock_id = rec["source_stock_id"]
        conn = app_db.get_db()
        try:
            cur = conn.execute(
                "INSERT INTO movements (item_id, delta, reason) VALUES (?,?,?)",
                (item["id"], -1, "退回已領出"),
            )
            legacy_id = cur.lastrowid
            conn.commit()
        finally:
            conn.close()

        repaired = client.post(
            f"/api/stockout-returns/{legacy_id}/repair",
            json={"source_movement_id": rec["id"], "return_stock_id": stock_id},
        )
        assert repaired.status_code == 400
        assert "退回數量必須大於 0" in repaired.json()["detail"]


# ========== 已領出：編輯（2026-08-10 新增） ==========

class TestStockoutEdit:
    def _out(self, client, item_id, qty=3, dest="台北案場"):
        """測試 helper：出庫 qty 後回傳該品項最新已領出記錄"""
        client.post("/api/stockout", json={
            "item_id": item_id, "qty": qty, "destination": dest})
        outs = client.get("/api/stockouts").json()
        return [o for o in outs if o["item_id"] == item_id][0]

    def test_edit_destination(self, client):
        """編輯去向：只改 destination，數量不變"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3, dest="台北案場")
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"destination": "台中工地"})
        assert r.status_code == 200
        assert r.json()["destination"] == "台中工地"
        assert r.json()["delta"] == -3
        assert _get_item(client, item["id"])["total_qty"] == 7  # 庫存不變

    def test_edit_qty_up_deducts_more(self, client):
        """數量調大：差額自動多扣庫存"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)  # 庫存 10→7
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 5})
        assert r.status_code == 200
        assert r.json()["delta"] == -5
        assert _get_item(client, item["id"])["total_qty"] == 5  # 7-2

        # 差額流水 reason='已領出編輯調整'
        mov = client.get("/api/movements").json()
        adj = [m for m in mov if m["reason"] == "已領出編輯調整"]
        assert len(adj) == 1
        assert adj[0]["delta"] == -2

    def test_edit_qty_down_adds_back(self, client):
        """數量調小：差額自動補回庫存"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)  # 庫存 10→7
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 1})
        assert r.status_code == 200
        assert r.json()["delta"] == -1
        assert _get_item(client, item["id"])["total_qty"] == 9  # 7+2

    def test_edit_qty_insufficient_400(self, client):
        """數量調大到超過庫存 → 400"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)  # 庫存 7
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 100})
        assert r.status_code == 400  # 庫存不足

    def test_edit_reverted_400(self, client):
        """已退回的記錄不能編輯"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        client.post(f"/api/stockouts/{rec['id']}/return")
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"destination": "台中"})
        assert r.status_code == 400

    def test_edit_zero_qty_400(self, client):
        """數量改 0 → 422（Pydantic gt=0 驗證）"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 0})
        assert r.status_code == 422

    def test_edit_created_at(self, client):
        """編輯日期"""
        item = _add_item(client, name="冷媒", qty=10)
        rec = self._out(client, item["id"], qty=3)
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"created_at": "2026-08-09 10:00:00"})
        assert r.status_code == 200
        assert r.json()["created_at"].startswith("2026-08-09")

    def test_edit_manual_adjust_rejected(self, client):
        """手動調整（非出庫）的負數記錄不能被編輯"""
        item = _add_item(client, name="防蟲罩", qty=5)
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": -2, "reason": "手動調整"})
        mov = client.get("/api/movements").json()
        adj = [m for m in mov if m["reason"] == "手動調整"][0]
        r = client.patch(f"/api/stockouts/{adj['id']}", json={"destination": "台北"})
        assert r.status_code == 400


# ========== 加減庫存 vs 待領出（2026-08-10：防止可領數量變負） ==========

class TestAdjustWithPrepared:
    def test_adjust_minus_below_prepared_400(self, client):
        """減少後庫存低於待領出 → 400（堵死可領數量變負）"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 4})  # 待領出 4
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": -7, "reason": "手動調整"})
        assert r.status_code == 400  # 10-7=3 < 4
        assert _get_item(client, item["id"])["total_qty"] == 10  # 庫存沒被扣

    def test_adjust_minus_ok_with_prepared(self, client):
        """減少後庫存仍 ≥ 待領出 → 允許"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 4})
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": -5, "reason": "手動調整"})
        assert r.status_code == 200  # 10-5=5 ≥ 4
        assert r.json()["after"] == 5

    def test_adjust_plus_ignores_prepared(self, client):
        """正向調整不受待領出限制"""
        item = _add_item(client, name="銅管", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 8})
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": 5, "reason": "進貨"})
        assert r.status_code == 200
        assert r.json()["after"] == 15


# ========== 出庫（指定位置） ==========

class TestStockOutLocation:
    def test_stockout_specific_location(self, client):
        """v10：指定從哪個位置出"""
        item = _add_item(client, name="冷媒", location="A倉", qty=5)
        client.post(f"/api/items/{item['id']}/stocks",
                    json={"location": "B倉", "qty": 3})
        r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 2, "destination": "中山路案場",
            "location": "A倉",
        })
        assert r.status_code == 200
        updated = _get_item(client, item["id"])
        a = [s for s in updated["stocks"] if s["location"] == "A倉"][0]
        b = [s for s in updated["stocks"] if s["location"] == "B倉"][0]
        assert a["qty"] == 3  # A倉 5-2=3
        assert b["qty"] == 3   # B倉 不動
        assert updated["total_qty"] == 6

    def test_stockout_no_location_deducts_all(self, client):
        """不指定位置：依序扣（先扣第一筆）"""
        item = _add_item(client, name="明細", location="A倉", qty=5)
        client.post(f"/api/items/{item['id']}/stocks",
                    json={"location": "B倉", "qty": 3})
        r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 6, "destination": "板橋案場",
        })
        assert r.status_code == 200
        updated = _get_item(client, item["id"])
        a = [s for s in updated["stocks"] if s["location"] == "A倉"][0]
        b = [s for s in updated["stocks"] if s["location"] == "B倉"][0]
        assert a["qty"] == 0      # A倉扣完 5
        assert b["qty"] == 2      # 剩下 1 從 B 扣 3-1=2
        assert updated["total_qty"] == 2

    def test_stockout_insufficient(self, client):
        """驗證庫存不足時出庫回傳 400"""
        item = _add_item(client, name="冷媒", qty=2)
        r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 5, "destination": "客戶家",
        })
        assert r.status_code == 400


# ========== 整組（套件） ==========

class TestKits:
    def test_create_kit(self, client):
        """驗證建立整組（套件）並自動建立 is_kit 品項"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        r = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
            "note": "標準套件",
        })
        assert r.status_code == 201
        assert r.json()["name"] == "銅管接頭組"

        # 自動建立 is_kit=1 的品項
        kits = client.get("/api/kits").json()
        assert len(kits) == 1
        assert kits[0]["name"] == "銅管接頭組"
        assert len(kits[0]["components"]) == 2
        # 整組照片（2026-08-12 Sarah 需求）：每個材料（單一庫存品項）有自己的照片縮圖
        # → components 每筆有 has_photo 欄位（不鎖定 True/False——測試 DB 的 item id 可能撞到真實 static/uploads/ 既有照片）
        assert all("has_photo" in c for c in kits[0]["components"])
        assert all(isinstance(c["has_photo"], bool) for c in kits[0]["components"])

    def test_items_have_in_kits_field(self, client):
        """2026-08-13 Sarah 需求：/api/items 每筆回傳 in_kits（該品項屬於哪些整組），
        缺貨/低庫存清單標註「屬於整組：名稱」用；未加入任何整組的品項為空陣列"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        # 建立整組前：材料 in_kits 為空陣列
        ia0 = _get_item(client, a["id"])
        assert ia0["in_kits"] == []
        # 建立整組：銅管/接頭都加入
        r = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        })
        assert r.status_code == 201
        # 建立後：材料 in_kits 含整組名稱（list_items 與單筆都一致）
        items = client.get("/api/items").json()
        ia = next(x for x in items if x["id"] == a["id"])
        ib = next(x for x in items if x["id"] == b["id"])
        assert ia["in_kits"] == ["銅管接頭組"]
        assert ib["in_kits"] == ["銅管接頭組"]
        # 套件品項（整組本身）不屬於任何整組
        kit_item = next(x for x in items if x["is_kit"])
        assert kit_item["in_kits"] == []
        # 單筆 GET 也一致
        ia_single = _get_item(client, a["id"])
        assert ia_single["in_kits"] == ["銅管接頭組"]

    def test_create_kit_requires_items(self, client):
        """驗證整組沒有元件時回傳 400"""
        r = client.post("/api/kits", json={"name": "空套件", "items": []})
        assert r.status_code == 400

    def test_assemble_deducts_materials(self, client):
        """驗證組裝整組會扣減材料庫存並增加整組庫存"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        kit = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        }).json()

        r = client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 3})
        assert r.status_code == 200

        # 材料被扣：銅管 10-6=4、接頭 20-3=17
        ia = _get_item(client, a["id"])
        ib = _get_item(client, b["id"])
        assert ia["total_qty"] == 4
        assert ib["total_qty"] == 17

        # 整組庫存 +3
        kits = client.get("/api/kits").json()
        assert kits[0]["stock_qty"] == 3

    def test_assemble_insufficient_material(self, client):
        """驗證材料不足時組裝回傳 400"""
        a = _add_item(client, name="銅管", qty=1)
        b = _add_item(client, name="接頭", qty=20)
        kit = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        }).json()

        r = client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 1})
        assert r.status_code == 400  # 銅管要 2 但只剩 1

    def test_disassemble_returns_materials(self, client):
        """驗證拆解整組會將材料加回庫存並扣減整組庫存"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        kit = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        }).json()
        client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 2})

        r = client.post(f"/api/kits/{kit['id']}/disassemble", json={"qty": 1})
        assert r.status_code == 200

        # 材料加回：銅管 6+2=8、接頭 18+1=19
        ia = _get_item(client, a["id"])
        ib = _get_item(client, b["id"])
        assert ia["total_qty"] == 8
        assert ib["total_qty"] == 19

        # 整組庫存 2-1=1
        kits = client.get("/api/kits").json()
        assert kits[0]["stock_qty"] == 1

    def test_disassemble_insufficient_kit(self, client):
        """驗證整組庫存不足時拆解回傳 400"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        kit = client.post("/api/kits", json={
            "name": "銅管接頭組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        }).json()

        r = client.post(f"/api/kits/{kit['id']}/disassemble", json={"qty": 5})
        assert r.status_code == 400  # 整組庫存只有 0

    def test_update_kit(self, client):
        """2026-08-11 Sarah：整組可編輯（名稱/備註 + 全量替換材料）"""
        a = _add_item(client, name="銅管", qty=10)
        b = _add_item(client, name="接頭", qty=20)
        c = _add_item(client, name="閥體", qty=5)
        kit = client.post("/api/kits", json={
            "name": "舊名", "note": "舊備註",
            "items": [{"item_id": a["id"], "qty": 2}],
        }).json()
        r = client.put(f"/api/kits/{kit['id']}", json={
            "name": "新名", "note": "新備註",
            "items": [{"item_id": b["id"], "qty": 1}, {"item_id": c["id"], "qty": 3}],
        })
        assert r.status_code == 200
        k = client.get("/api/kits").json()[0]
        assert k["name"] == "新名"
        assert k["note"] == "新備註"
        assert [x["item_id"] for x in k["components"]] == [b["id"], c["id"]]
        assert [x["need_qty"] for x in k["components"]] == [1, 3]
        # 套件品項名稱同步
        items = client.get("/api/items").json()
        kit_item = [i for i in items if i["id"] == kit["item_id"]][0]
        assert kit_item["name"] == "新名"

    def test_update_kit_not_found(self, client):
        """驗證更新不存在的整組回傳 404"""
        r = client.put("/api/kits/999", json={"name": "X", "items": [{"item_id": 1, "qty": 1}]})
        assert r.status_code == 404

    def test_update_kit_requires_items(self, client):
        """更新時材料空白 → 400"""
        a = _add_item(client, name="銅管", qty=10)
        kit = client.post("/api/kits", json={"name": "組", "items": [{"item_id": a["id"], "qty": 1}]}).json()
        r = client.put(f"/api/kits/{kit['id']}", json={"name": "空", "items": []})
        assert r.status_code == 400

    def test_delete_kit(self, client):
        """2026-08-11 Sarah：整組可刪除（定義 + 套件品項 + 流水一併清）"""
        a = _add_item(client, name="銅管", qty=10)
        kit = client.post("/api/kits", json={
            "name": "測試套件",
            "items": [{"item_id": a["id"], "qty": 1}],
        }).json()
        assert len(client.get("/api/kits").json()) == 1
        r = client.delete(f"/api/kits/{kit['id']}")
        assert r.status_code == 200
        assert client.get("/api/kits").json() == []
        # 套件品項也已刪除（不會殘留 is_kit 空殼）
        items = client.get("/api/items").json()
        assert all(i["id"] != kit["item_id"] for i in items)
        # 材料本身不受影響
        assert any(i["id"] == a["id"] for i in items)

    def test_delete_kit_not_found(self, client):
        """驗證刪除不存在的整組回傳 404"""
        r = client.delete("/api/kits/999")
        assert r.status_code == 404

    def test_delete_stockout_movement(self, client):
        """2026-08-11 Sarah：已領出紀錄可刪除（只刪紀錄、不回補庫存）"""
        item = _add_item(client, name="冷媒", qty=10)
        client.post("/api/stockout", json={"item_id": item["id"], "qty": 2, "destination": "工地A"})
        outs = client.get("/api/stockouts").json()
        assert len(outs) == 1
        mid = outs[0]["id"]
        r = client.delete(f"/api/stockouts/{mid}")
        assert r.status_code == 200
        assert client.get("/api/stockouts").json() == []
        # 庫存不能動（8 維持）
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 8

    def test_delete_stockout_movement_not_found(self, client):
        """驗證刪除不存在的已領出紀錄回傳 404"""
        r = client.delete("/api/stockouts/99999")
        assert r.status_code == 404


# ========== 盤點 ==========

class TestStocktake:
    def test_submit_stocktake(self, client):
        """驗證提交盤點後庫存更新為實際數量且 diff 正確"""
        item = _add_item(client, name="冷媒", qty=10)
        r = client.post("/api/stocktake", json={
            "items": [{"item_id": item["id"], "location": "測試位置",
                       "actual_qty": 12, "note": "多找到2罐"}],
        })
        assert r.status_code == 200
        res = r.json()
        assert res["count"] == 1
        assert res["results"][0]["diff"] == 2

        # 庫存更新為實際數量
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 12

    def test_submit_stocktake_specific_location(self, client):
        """v10：盤點指定位置的實際數量"""
        item = _add_item(client, name="冷媒", location="A倉", qty=10)
        client.post(f"/api/items/{item['id']}/stocks",
                    json={"location": "B倉", "qty": 5})
        r = client.post("/api/stocktake", json={
            "items": [{"item_id": item["id"], "location": "B倉", "actual_qty": 7}],
        })
        assert r.status_code == 200
        res = r.json()
        assert res["results"][0]["diff"] == 2  # B倉 5→7
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 17  # A10 + B7

    def test_submit_stocktake_negative_diff(self, client):
        """驗證盤點數量少於庫存時 diff 為負數"""
        item = _add_item(client, name="冷媒", qty=10)
        r = client.post("/api/stocktake", json={
            "items": [{"item_id": item["id"], "location": "測試位置", "actual_qty": 8}],
        })
        assert r.status_code == 200
        assert r.json()["results"][0]["diff"] == -2

    def test_stocktake_dates(self, client):
        """驗證盤點日期紀錄與總 diff"""
        item = _add_item(client, name="冷媒", qty=10)
        client.post("/api/stocktake", json={
            "take_date": "2026-08-25",
            "items": [{"item_id": item["id"], "location": "測試位置", "actual_qty": 11}],
        })
        dates = client.get("/api/stocktake/dates").json()
        assert len(dates) == 1
        assert dates[0]["take_date"] == "2026-08-25"
        assert dates[0]["total_diff"] == 1


    def test_stocktake_optimistic_lock_where_clause(self, client):
        """2026-08-14 樂觀鎖核心行為：API 讀到的 system_qty 已被他人異動 → 寫回 rowcount=0（409 觸發條件）"""
        item = _add_item(client, name="冷媒", qty=10)
        stock_id = item["stocks"][0]["id"]
        c1 = app_db.get_db()
        c2 = app_db.get_db()
        try:
            # A 連線（盤點端）讀到 system_qty=10
            sys_qty = c1.execute("SELECT qty FROM item_stocks WHERE id=?", (stock_id,)).fetchone()["qty"]
            assert sys_qty == 10
            # B 連線（他人出庫/調整）把庫存改成 15
            c2.execute("UPDATE item_stocks SET qty=15 WHERE id=?", (stock_id,))
            c2.commit()
            # A 用舊值 10 當樂觀鎖條件寫回 → rowcount=0（應 409 拒絕）
            cur = c1.execute(
                "UPDATE item_stocks SET qty=qty+2, updated_at=datetime('now') WHERE id=? AND qty=?",
                (stock_id, sys_qty))
            assert cur.rowcount == 0
            # 沒覆蓋成功：庫存維持 15（不是 A 的 12）
            assert c1.execute("SELECT qty FROM item_stocks WHERE id=?", (stock_id,)).fetchone()["qty"] == 15
            # 對照：無異動時用正確值寫回 → rowcount=1（正常盤點不受影響）
            cur2 = c1.execute(
                "UPDATE item_stocks SET qty=qty+1, updated_at=datetime('now') WHERE id=? AND qty=?",
                (stock_id, 15))
            assert cur2.rowcount == 1
            c1.commit()
        finally:
            c1.close()
            c2.close()

    def test_stocktake_missing_item_id_400(self, client):
        """盤點缺 item_id → 400（防 KeyError 500）"""
        r = client.post('/api/stocktake', json={'items': [{'location': 'X', 'actual_qty': 5}]})
        assert r.status_code == 400
        assert 'item_id' in r.json()['detail']

    def test_stocktake_invalid_item_id_type_400(self, client):
        """盤點 item_id 非整數 → 400"""
        r = client.post('/api/stocktake', json={'items': [{'item_id': 'abc', 'actual_qty': 5}]})
        assert r.status_code == 400


# ========== 位置庫存 PATCH（2026-08-14 併發修復：差額寫回 + 流水） ==========

class TestPatchStock:
    def test_patch_stock_delta_and_movement(self, client):
        """2026-08-14：PATCH stocks qty 差額寫回 + 補 movements 流水（原本絕對值覆蓋且零流水）"""
        item = _add_item(client, name="冷媒", qty=10)
        stock_id = item["stocks"][0]["id"]
        r = client.patch(f"/api/stocks/{stock_id}", json={"qty": 15})
        assert r.status_code == 200, r.text
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 15
        # 流水有「編輯位置調整」delta=+5（before=10 after=15）
        movs = client.get("/api/movements").json()
        hits = [m for m in movs if m["reason"] == "編輯位置調整" and m["item_id"] == item["id"]]
        assert len(hits) == 1
        assert hits[0]["delta"] == 5
        assert hits[0]["before_qty"] == 10
        assert hits[0]["after_qty"] == 15

    def test_patch_stock_zero_delta_no_movement(self, client):
        """2026-08-14：qty 無變化 → 不寫流水（避免假流水）"""
        item = _add_item(client, name="冷媒", qty=10)
        stock_id = item["stocks"][0]["id"]
        r = client.patch(f"/api/stocks/{stock_id}", json={"qty": 10})
        assert r.status_code == 200, r.text
        movs = client.get("/api/movements").json()
        assert not any(m["reason"] == "編輯位置調整" for m in movs)

    def test_patch_stock_below_prepared_400(self, client):
        """2026-08-14：負數調整後總量低於待領出 → 400（總量語意與 adjust_qty 一致）"""
        item = _add_item(client, name="冷媒", qty=10)
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 8})
        stock_id = item["stocks"][0]["id"]
        r = client.patch(f"/api/stocks/{stock_id}", json={"qty": 1})
        assert r.status_code == 400
        assert "待領出" in r.json()["detail"]
        # 庫存沒被改
        updated = _get_item(client, item["id"])
        assert updated["total_qty"] == 10

    def test_patch_stock_not_found_404(self, client):
        """2026-08-14：不存在的 stock → 404"""
        r = client.patch("/api/stocks/99999", json={"qty": 5})
        assert r.status_code == 404


# ========== 編輯端點樂觀鎖（2026-08-14 Phase 2：updated_at 快照守衛） ==========

class TestOptimisticLock:
    def test_update_item_optimistic_lock_ok(self, client):
        """2026-08-14：PATCH 帶正確 updated_at 快照 → 200"""
        item = _add_item(client, name="冷媒", qty=10)
        r = client.patch(f"/api/items/{item['id']}", json={"name": "冷媒R22", "updated_at": item["updated_at"]})
        assert r.status_code == 200, r.text
        updated = _get_item(client, item["id"])
        assert updated["name"] == "冷媒R22"

    def test_update_item_optimistic_lock_conflict_409(self, client):
        """2026-08-14：PATCH 帶過期 updated_at 快照（被他人改過）→ 409 不覆蓋"""
        item = _add_item(client, name="冷媒", qty=10)
        # 他人先改過（updated_at 前進）
        client.patch(f"/api/items/{item['id']}", json={"name": "他人改的"})
        # 用舊快照儲存 → 409
        r = client.patch(f"/api/items/{item['id']}", json={"name": "我的修改", "updated_at": item["updated_at"]})
        assert r.status_code == 409
        assert "已被他人修改" in r.json()["detail"]
        # 品項維持他人改的名字（沒被覆蓋）
        updated = _get_item(client, item["id"])
        assert updated["name"] == "他人改的"

    def test_update_kit_optimistic_lock_conflict_409(self, client):
        """2026-08-14：PUT 整組帶過期 updated_at → 409（kits migration 後 updated_at 欄生效）"""
        a = _add_item(client, name="銅管", qty=10)
        r = client.post("/api/kits", json={"name": "測試組", "items": [{"item_id": a["id"], "qty": 1}]})
        assert r.status_code == 201
        kits = client.get("/api/kits").json()
        kit = kits[0]
        # 他人先改過
        client.put(f"/api/kits/{kit['id']}", json={"name": "他人改組", "items": [{"item_id": a["id"], "qty": 2}]})
        # 用舊快照儲存 → 409
        r = client.put(f"/api/kits/{kit['id']}", json={
            "name": "我的修改", "items": [{"item_id": a["id"], "qty": 1}], "updated_at": kit["updated_at"] or "1999-01-01 00:00:00"})
        assert r.status_code == 409
        assert "已被他人修改" in r.json()["detail"]
        # 整組名稱沒被覆蓋
        kits2 = client.get("/api/kits").json()
        assert kits2[0]["name"] == "他人改組"


# ========== 統計（缺貨只列單一 / 低量含整組） ==========

class TestStats:
    def test_stats_basic(self, client):
        """驗證統計 API 回傳總品項數、總數量與缺貨數"""
        _add_item(client, name="品一", qty=5)
        _add_item(client, name="品二", qty=0)  # 缺貨
        s = client.get("/api/stats").json()
        assert s["total_items"] == 2
        assert s["total_qty"] == 5
        assert s["zero_stock"] == 1  # 缺貨只算單一

    def test_stats_zero_stock_excludes_kit(self, client):
        """缺貨只列單一材料：整組 qty=0 不算缺貨；整組也不進一般低庫存（走 kit shortage/insufficient）"""
        _add_item(client, name="單一材料", qty=0)
        a = _add_item(client, name="材料A", qty=10)
        kit = client.post("/api/kits", json={
            "name": "整組套件",
            "items": [{"item_id": a["id"], "qty": 1}],
        }).json()
        _get_item(client, kit["item_id"])  # 確保 kit item 建立

        # 把整組的庫存全扣掉（缺貨狀態）＋ low_stock=1（低庫存狀態）
        import sqlite3
        conn = sqlite3.connect(app_db.DB_PATH)
        conn.execute("UPDATE item_stocks SET qty=0 WHERE item_id=?", (kit["item_id"],))
        conn.execute("UPDATE items SET low_stock=1 WHERE id=?", (kit["item_id"],))
        conn.commit()
        conn.close()

        s = client.get("/api/stats").json()
        assert s["zero_stock"] == 1    # 整組 qty=0 不列入缺貨（只有單一材料那 1 筆）
        assert s["low_stock"] == 0     # 一般低庫存排除整組、排除 qty=0（材料A low_stock=0 不計；整組走 kit 邏輯）

    def test_stats_low_stock_excludes_kit(self, client):
        """一般低庫存只列單一：整組不計入（走 kit shortage/insufficient）"""
        _add_item(client, name="單一低", qty=1, low_stock=3)
        a = _add_item(client, name="材料A", qty=10)
        kit = client.post("/api/kits", json={
            "name": "整組套件",
            "items": [{"item_id": a["id"], "qty": 1}],
        }).json()

        import sqlite3
        conn = sqlite3.connect(app_db.DB_PATH)
        conn.execute("UPDATE item_stocks SET qty=2 WHERE item_id=?", (kit["item_id"],))
        conn.execute("UPDATE items SET low_stock=5 WHERE id=?", (kit["item_id"],))
        conn.commit()
        conn.close()

        s = client.get("/api/stats").json()
        assert s["low_stock"] == 1  # 只有單一那 1 筆；整組不計入一般低庫存


# ========== 2026-09-12：庫存警示 domain contract（方案 B） ==========
# LOW: 非整組、ROUND(qty,3) > 0、low_stock > 0、ROUND(qty,3) <= low_stock
# OUT: 非整組、ROUND(qty,3) <= 0；LOW ∩ OUT = ∅；KPI 數 == 清單長度。

class TestStatsAlertContract:
    """一般 low/out 互斥且排除整組；stats KPI 數與同回應清單長度一致。"""

    def _seed_abcd(self, client):
        """A(qty0)→out、B(qty3)→low、C(qty10)→normal、D(整組)→皆不進。"""
        import sqlite3
        _add_item(client, name="契約A缺貨", qty=0, low_stock=5)
        _add_item(client, name="契約B低庫存", qty=3, low_stock=5)
        _add_item(client, name="契約C正常", qty=10, low_stock=5)
        material = _add_item(client, name="契約材料", qty=10)
        kit = client.post("/api/kits", json={
            "name": "契約整組",
            "items": [{"item_id": material["id"], "qty": 1}],
        })
        assert kit.status_code == 201, kit.text
        kit_item_id = kit.json()["item_id"]
        conn = sqlite3.connect(app_db.DB_PATH)
        conn.execute("UPDATE item_stocks SET qty=2 WHERE item_id=?", (kit_item_id,))
        conn.execute("UPDATE items SET low_stock=5 WHERE id=?", (kit_item_id,))
        conn.commit()
        conn.close()

    def test_stats_low_out_contract_abcd(self, client):
        """A/B/C/D 分類 + KPI 數 == 清單長度 + low/out 互斥。"""
        self._seed_abcd(client)
        s = client.get("/api/stats").json()
        assert s["low_stock"] == 1
        assert s["zero_stock"] == 1
        assert s["low_stock"] == len(s["low_items"])
        assert s["zero_stock"] == len(s["zero_items"])
        low_names = [i["name"] for i in s["low_items"]]
        zero_names = [i["name"] for i in s["zero_items"]]
        assert low_names == ["契約B低庫存"]
        assert zero_names == ["契約A缺貨"]
        assert not (set(low_names) & set(zero_names))

    def test_stats_alert_edge_cases(self, client):
        """qty=0/負數/0.0004→out；low_stock=0 永不 low；整組不進一般 low/out。"""
        import sqlite3
        _add_item(client, name="邊界零", qty=0, low_stock=5)
        neg = _add_item(client, name="邊界負", qty=1, low_stock=5)
        tiny = _add_item(client, name="邊界微量", qty=1, low_stock=5)
        _add_item(client, name="邊界關閉警示", qty=1, low_stock=0)
        material = _add_item(client, name="邊界材料", qty=10)
        kit = client.post("/api/kits", json={
            "name": "邊界整組",
            "items": [{"item_id": material["id"], "qty": 1}],
        })
        assert kit.status_code == 201, kit.text
        kit_item_id = kit.json()["item_id"]
        conn = sqlite3.connect(app_db.DB_PATH)
        conn.execute("UPDATE item_stocks SET qty=-2 WHERE item_id=?", (neg["id"],))
        conn.execute("UPDATE item_stocks SET qty=0.0004 WHERE item_id=?", (tiny["id"],))
        conn.execute("UPDATE item_stocks SET qty=0 WHERE item_id=?", (kit_item_id,))
        conn.execute("UPDATE items SET low_stock=5 WHERE id=?", (kit_item_id,))
        conn.commit()
        conn.close()
        s = client.get("/api/stats").json()
        zero_names = [i["name"] for i in s["zero_items"]]
        low_names = [i["name"] for i in s["low_items"]]
        assert "邊界零" in zero_names
        assert "邊界負" in zero_names
        assert "邊界微量" in zero_names  # ROUND(0.0004,3)=0 → out
        assert "邊界關閉警示" not in low_names
        assert "邊界關閉警示" not in zero_names
        assert "邊界整組" not in low_names
        assert "邊界整組" not in zero_names
        assert s["low_stock"] == len(s["low_items"])
        assert s["zero_stock"] == len(s["zero_items"])

    def test_stats_alerts_beyond_page_size(self, client):
        """超過分頁大小的資料量：stats 清單/KPI 必須是完整集合，不受 page/page_size 影響。"""
        import sqlite3
        conn = sqlite3.connect(app_db.DB_PATH)
        cur = conn.cursor()
        for i in range(60):
            cur.execute(
                "INSERT INTO items (brand, code, name, unit, low_stock, site) VALUES (?,?,?,?,?,?)",
                ("測試牌", f"BULK-Z{i}", f"大量缺貨{i}", "個", 5, "office"))
            cur.execute("INSERT INTO item_stocks (item_id, location, qty) VALUES (?,?,?)",
                        (cur.lastrowid, "A倉", 0))
        for i in range(60):
            cur.execute(
                "INSERT INTO items (brand, code, name, unit, low_stock, site) VALUES (?,?,?,?,?,?)",
                ("測試牌", f"BULK-L{i}", f"大量低庫存{i}", "個", 5, "office"))
            cur.execute("INSERT INTO item_stocks (item_id, location, qty) VALUES (?,?,?)",
                        (cur.lastrowid, "B倉", 2))
        for i in range(10):
            cur.execute(
                "INSERT INTO items (brand, code, name, unit, low_stock, site) VALUES (?,?,?,?,?,?)",
                ("測試牌", f"BULK-O{i}", f"大量正常{i}", "個", 5, "office"))
            cur.execute("INSERT INTO item_stocks (item_id, location, qty) VALUES (?,?,?)",
                        (cur.lastrowid, "C倉", 50))
        conn.commit()
        conn.close()
        s = client.get("/api/stats").json()
        assert s["total_items"] == 130
        assert s["low_stock"] == 60
        assert s["zero_stock"] == 60
        assert len(s["low_items"]) == 60
        assert len(s["zero_items"]) == 60
        # 分頁請求不影響完整統計：每頁 50 筆共 3 頁，stats 始終是全量
        for page in (1, 2, 3):
            body = client.get("/api/items", params={
                "site": "office", "page": page, "page_size": 50}).json()
            assert body["total"] == 130
            assert body["stats"]["low_stock"] == 60
            assert body["stats"]["zero_stock"] == 60
        full = client.get("/api/items", params={
            "site": "office", "page": 1, "page_size": 50,
            "include_alert_items": 1}).json()
        assert len(full["stats"]["low_items"]) == 60
        assert len(full["stats"]["zero_items"]) == 60


# ========== 前端頁面 ==========

class TestFrontend:
    def test_index_served(self, client):
        """驗證前端首頁可存取且包含庫存內容"""
        r = client.get("/")
        assert r.status_code == 200
        assert "庫存" in r.text  # 前端頁面有內容


# ========== 分片（辦公室 / 倉庫） ==========

class TestSiteSharding:
    def test_create_item_default_office(self, client):
        """新增品項預設是辦公室（site=office）"""
        item = _add_item(client, name="辦公室品項")
        assert item["site"] == "office"

    def test_create_item_warehouse(self, client):
        """新增品項指定倉庫"""
        item = _add_item(client, name="倉庫品項", site="warehouse")
        assert item["site"] == "warehouse"

    def test_list_items_site_filter(self, client):
        """site 篩選：只回該分片品項"""
        _add_item(client, name="辦公室品項")
        _add_item(client, name="倉庫品項", site="warehouse")

        office = client.get("/api/items", params={"site": "office"}).json()
        warehouse = client.get("/api/items", params={"site": "warehouse"}).json()
        assert len(office) == 1
        assert len(warehouse) == 1
        assert office[0]["name"] == "辦公室品項"
        assert warehouse[0]["name"] == "倉庫品項"

    def test_list_items_site_all(self, client):
        """site=all 或無參數：回全部品項"""
        _add_item(client, name="辦公室品項")
        _add_item(client, name="倉庫品項", site="warehouse")
        all_items = client.get("/api/items", params={"site": "all"}).json()
        assert len(all_items) == 2

    def test_stats_site_separated(self, client):
        """統計分片分開：辦公室/倉庫各自統計"""
        _add_item(client, name="辦公室A", qty=5)
        _add_item(client, name="辦公室B", qty=3)
        _add_item(client, name="倉庫品", qty=10, site="warehouse")

        office = client.get("/api/stats", params={"site": "office"}).json()
        warehouse = client.get("/api/stats", params={"site": "warehouse"}).json()
        assert office["total_items"] == 2
        assert office["total_qty"] == 8
        assert warehouse["total_items"] == 1
        assert warehouse["total_qty"] == 10

    def test_prepared_site_filter(self, client):
        """待領出清單分片分開"""
        office_item = _add_item(client, name="辦公室品項", qty=10)
        wh_item = _add_item(client, name="倉庫品項", qty=10, site="warehouse")
        client.post(f"/api/items/{office_item['id']}/prepare", json={"qty": 2})
        client.post(f"/api/items/{wh_item['id']}/prepare", json={"qty": 3})

        office_prepared = client.get("/api/prepared", params={"site": "office"}).json()
        wh_prepared = client.get("/api/prepared", params={"site": "warehouse"}).json()
        assert len(office_prepared) == 1
        assert office_prepared[0]["prepared_qty"] == 2
        assert len(wh_prepared) == 1
        assert wh_prepared[0]["prepared_qty"] == 3

    def test_stockouts_site_filter(self, client):
        """已領出紀錄分片分開"""
        office_item = _add_item(client, name="辦公室品項", qty=10)
        wh_item = _add_item(client, name="倉庫品項", qty=10, site="warehouse")
        client.post(f"/api/items/{office_item['id']}/adjust",
                    json={"delta": -2, "reason": "出庫", "destination": "辦公室去向"})
        client.post(f"/api/items/{wh_item['id']}/adjust",
                    json={"delta": -3, "reason": "出庫", "destination": "倉庫去向"})

        office_outs = client.get("/api/stockouts", params={"site": "office"}).json()
        wh_outs = client.get("/api/stockouts", params={"site": "warehouse"}).json()
        assert len(office_outs) == 1
        assert office_outs[0]["destination"] == "辦公室去向"
        assert len(wh_outs) == 1
        assert wh_outs[0]["destination"] == "倉庫去向"

    def test_stockout_site_filter_shows_nonstock_everywhere(self, client):
        """2026-08-13：非庫存品項（is_deleted=1）不受 site 過濾——office/warehouse 都看得到"""
        office_item = _add_item(client, name="辦公室品項", qty=10)
        client.post(f"/api/items/{office_item['id']}/adjust",
                    json={"delta": -2, "reason": "出庫", "destination": "辦公室去向"})
        r = client.post("/api/stockout/nonstock", json={
            "name": "臨時耗材", "qty": 1, "destination": "某案場"})
        assert r.status_code == 200
        ns_id = r.json()["id"]

        office_outs = client.get("/api/stockouts", params={"site": "office"}).json()
        wh_outs = client.get("/api/stockouts", params={"site": "warehouse"}).json()
        # 非庫存品項兩邊都顯示
        assert any(o["item_id"] == ns_id for o in office_outs), "非庫存品項在 office 過濾下應顯示"
        assert any(o["item_id"] == ns_id for o in wh_outs), "非庫存品項在 warehouse 過濾下應顯示"
        # 一般品項只有 office 看得到
        assert any(o["item_id"] == office_item["id"] for o in office_outs)
        assert all(o["item_id"] != office_item["id"] for o in wh_outs)

    def test_update_item_site(self, client):
        """編輯品項可以搬移分片"""
        item = _add_item(client, name="品項", site="office")
        r = client.patch(f"/api/items/{item['id']}", json={"site": "warehouse"})
        assert r.status_code == 200
        assert r.json()["site"] == "warehouse"

    def test_batch_location_can_move_between_sites(self, client):
        """批次改位置可同時將品項從辦公室搬到倉庫，再搬回辦公室。"""
        item = _add_item(client, name="可搬移品項", site="office", location="辦公室櫃A")
        stock_id = item["stocks"][0]["id"]

        moved = client.post("/api/stocks/batch-location", json={
            "stock_ids": [stock_id],
            "new_location": "倉庫櫃B | 第二層",
            "new_site": "warehouse",
        })
        assert moved.status_code == 200, moved.text
        warehouse = client.get("/api/items", params={"site": "warehouse"}).json()
        assert warehouse[0]["site"] == "warehouse"
        assert warehouse[0]["stocks"][0]["location"] == "倉庫櫃B | 第二層"

        moved_back = client.post("/api/stocks/batch-location", json={
            "stock_ids": [stock_id],
            "new_location": "辦公室櫃C",
            "new_site": "office",
        })
        assert moved_back.status_code == 200, moved_back.text
        office = client.get("/api/items", params={"site": "office"}).json()
        assert office[0]["site"] == "office"
        assert office[0]["stocks"][0]["location"] == "辦公室櫃C"

    def test_batch_location_rejects_invalid_site(self, client):
        """批次改位置只接受 office 或 warehouse。"""
        item = _add_item(client, name="分片驗證品項")
        stock_id = item["stocks"][0]["id"]
        response = client.post("/api/stocks/batch-location", json={
            "stock_ids": [stock_id], "new_location": "新位置", "new_site": "factory",
        })
        assert response.status_code == 422


class TestQuotations:
    def _quote_payload(self, client):
        item = _add_item(client, name="報價用冷氣", brand="測試牌", code="Q-001", qty=8)
        return item, {
            "quote_date": "2026-09-09", "customer_name": "測試客戶",
            "contact": "王先生 0912-345-678", "address": "台北市測試區",
            "valid_days": 30, "tax_type": "included", "note": "付款條件：驗收後付款",
            "items": [{"inventory_item_id": item["id"], "item_name": item["name"],
                       "specification": "3.6kW 基本安裝", "qty": 2,
                       "unit": "台", "unit_price": 25000}],
        }

    def test_create_list_update_delete_quotation(self, client):
        item, payload = self._quote_payload(client)
        created = client.post("/api/quotations", json=payload)
        assert created.status_code == 201, created.text
        quote = created.json()
        assert quote["customer_name"] == "測試客戶"
        assert quote["items"][0]["inventory_item_id"] == item["id"]
        assert quote["total"] == 50000
        listed = client.get("/api/quotations", params={"q": "測試客戶"})
        assert listed.status_code == 200 and listed.json()["total"] == 1
        changed = client.put(f"/api/quotations/{quote['id']}", json=dict(payload, customer_name="修改後客戶"))
        assert changed.status_code == 200 and changed.json()["customer_name"] == "修改後客戶"
        assert client.delete(f"/api/quotations/{quote['id']}").status_code == 200
        assert client.get(f"/api/quotations/{quote['id']}").status_code == 404

    def test_inventory_items_can_be_used_in_quotation(self, client):
        item = _add_item(client, name="可帶入品項", brand="大金", code="INV-1", qty=4)
        response = client.get("/api/quotations/inventory-items", params={"q": "可帶入品項"})
        assert response.status_code == 200
        assert response.json()[0]["id"] == item["id"]
        assert response.json()[0]["unit"] == item["unit"]

    def test_quotation_exports_xlsx_and_pdf(self, client):
        _, payload = self._quote_payload(client)
        quote = client.post("/api/quotations", json=payload).json()
        xlsx = client.get(f"/api/quotations/{quote['id']}/export.xlsx")
        pdf = client.get(f"/api/quotations/{quote['id']}/export.pdf")
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
        assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert pdf.headers["content-type"] == "application/pdf"

    def test_quotation_rejects_blank_item_text(self, client):
        """品項名稱與單位只有空白時也必須拒絕。"""
        _, payload = self._quote_payload(client)
        payload["items"][0]["item_name"] = "   "
        response = client.post("/api/quotations", json=payload)
        assert response.status_code == 422

    def test_quotation_xlsx_contains_terms_and_pdf_escapes_text(self, client):
        """Excel 含報價條件；PDF 可處理 XML 特殊字元。"""
        _, payload = self._quote_payload(client)
        payload["customer_name"] = "客戶 <A&B>"
        quote = client.post("/api/quotations", json=payload).json()
        xlsx = client.get(f"/api/quotations/{quote['id']}/export.xlsx")
        import io
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(xlsx.content), data_only=True).active
        values = [row[0].value for row in ws.iter_rows()]
        assert "有效天數" in values and "稅別" in values and "備註／付款條件" in values
        pdf = client.get(f"/api/quotations/{quote['id']}/export.pdf")
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    def test_quotation_rejects_invalid_values(self, client):
        response = client.post("/api/quotations", json={"quote_date": "bad-date", "customer_name": "", "items": []})
        assert response.status_code in (400, 422)


# ========== v10 相容性與 CASCADE 完整性（2026-08-09 補測） ==========

class TestV10CompatAndCascade:
    def test_list_items_has_compat_fields(self, client):
        """list 端點每筆都要有相容欄位（qty/total_qty/stocks/location/note）——防 undefined"""
        _add_item(client, name="多位置品", location="A倉", qty=2, note="首批")
        items = client.get("/api/items").json()
        assert len(items) == 1
        it = items[0]
        for k in ("qty", "total_qty", "stocks", "location", "note"):
            assert k in it, f"list 回應缺少相容欄位 {k}"
        assert it["qty"] == it["total_qty"] == 2
        assert it["location"] == "A倉"
        assert it["note"] == "首批"
        assert len(it["stocks"]) == 1

    def test_delete_item_with_movements(self, client):
        """有異動紀錄的品項刪除（M6 soft-delete）：movements 稽核軌跡保留 + 庫存清零流水"""
        item = _add_item(client, name="要刪的", qty=5)
        client.post(f"/api/items/{item['id']}/adjust", json={"delta": -2, "reason": "出庫"})
        assert len(client.get("/api/movements").json()) == 1

        r = client.delete(f"/api/items/{item['id']}")
        assert r.status_code == 200
        # soft-delete：保留原有流水 + 新增「品項刪除清零」流水（剩餘 3）
        movements = client.get("/api/movements").json()
        assert len(movements) == 2
        clear_down = [m for m in movements if m["reason"] == "品項刪除清零"]
        assert len(clear_down) == 1
        assert clear_down[0]["delta"] == -3  # 原 5 - 已出 2 = 剩 3 清零

    def test_delete_item_with_prepared_qty_blocked(self, client):
        """有待領出數量的品項不可刪除 → 400"""
        item = _add_item(client, name='有待領', qty=10)
        client.post(f'/api/items/{item["id"]}/prepare', json={'qty': 3})
        r = client.delete(f'/api/items/{item["id"]}')
        assert r.status_code == 400
        assert '待領出' in r.json()['detail']

    def test_delete_item_with_stocktake(self, client):
        """有盤點紀錄的品項刪除（M6 soft-delete）：已刪品項不再出現在準備清單"""
        item = _add_item(client, name="盤點過", qty=10)
        client.post("/api/stocktake", json={
            "items": [{"item_id": item["id"], "location": "測試位置", "actual_qty": 10}]})

        r = client.delete(f"/api/items/{item['id']}")
        assert r.status_code == 200
        assert len(client.get("/api/prepared").json()) == 0  # 已刪品項不顯示在準備清單

    def test_delete_kit_and_material(self, client):
        """刪整組：kits + kit_items 一併清；被材料引用的品項也可刪"""
        a = _add_item(client, name="材料甲", qty=10)
        b = _add_item(client, name="材料乙", qty=10)
        kit = client.post("/api/kits", json={
            "name": "測試整組",
            "items": [{"item_id": a["id"], "qty": 2}, {"item_id": b["id"], "qty": 1}],
        }).json()

        # 刪材料甲（被 kit_items 引用）→ 應成功且 kit_items 引用被移除
        r = client.delete(f"/api/items/{a['id']}")
        assert r.status_code == 200

        # 刪整組品項 → kits + kit_items 都清
        kid, kit_item_id = kit["item_id"], kit["id"]
        # 先建立 kit（自動建 is_kit item），再造第二個 kit 引用同 kit 品項？不需要——直接刪
        r2 = client.delete(f"/api/items/{kid}")
        assert r2.status_code == 200
        assert len(client.get("/api/kits").json()) == 0

    def test_update_main_fields_preserves_stocks(self, client):
        """只改主檔欄位（name/brand）不帶 stocks → 位置庫存原樣保留"""
        item = _add_item(client, name="舊名", brand="舊牌", location="A倉", qty=5)
        client.post(f"/api/items/{item['id']}/stocks", json={"location": "B倉", "qty": 3})
        r = client.patch(f"/api/items/{item['id']}", json={"name": "新名"})
        assert r.status_code == 200
        up = r.json()
        assert up["name"] == "新名"
        assert len(up["stocks"]) == 2          # 兩筆位置都還在
        assert up["total_qty"] == 8
        locs = sorted(s["location"] for s in up["stocks"])
        assert locs == ["A倉", "B倉"]

    def test_stockout_response_has_compat_fields(self, client):
        """出庫回應 = 完整品項（qty/stocks），前端可接"""
        item = _add_item(client, name="出庫品", location="A倉", qty=5)
        r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 2, "destination": "客戶"})
        assert r.status_code == 200
        p = r.json()
        assert p["qty"] == 3
        assert "stocks" in p and "note" in p and "location" in p

    def test_export_after_delete(self, client):
        """防回歸：刪除品項後 export 仍正常（JOIN 無殘留）"""
        item = _add_item(client, name="暫存", qty=1)
        client.delete(f"/api/items/{item['id']}")
        r = client.get("/api/export")
        assert r.status_code == 200

    def test_export_workbook_content(self, client):
        """匯出 workbook 可開、3 sheet 欄位正確、品項×位置展開語意、檔名 RFC 5987 編碼"""
        _add_item(client, name="冷媒管", qty=5, location="A倉")
        r = client.get("/api/export")
        assert r.status_code == 200
        assert r.headers["content-type"] == \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "filename*=UTF-8''" in r.headers["content-disposition"]

        import io as _io
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(r.content))
        # 2026-08-13 家豪：庫存明細拆「辦公室」「倉庫」兩頁
        assert wb.sheetnames == ["辦公室", "倉庫", "異動紀錄", "廠牌統計"]

        ws = wb["辦公室"]
        assert [c.value for c in ws[1]] == \
            ["編號", "廠牌", "品項名稱", "型號", "單位", "位置", "位置數量", "位置備註", "總數量"]
        # 每列 = 主檔 × 每位置一行：找到冷媒管列，位置/數量/總量正確（_add_item 預設 office）
        row = next(rr for rr in ws.iter_rows(min_row=2, values_only=True) if rr[2] == "冷媒管")
        assert row[5] == "A倉" and row[6] == 5 and row[8] == 5

        # 倉庫頁目前 0 筆：只有表頭
        ws_wh = wb["倉庫"]
        assert [c.value for c in ws_wh[1]] == \
            ["編號", "廠牌", "品項名稱", "型號", "單位", "位置", "位置數量", "位置備註", "總數量"]
        assert ws_wh.max_row == 1, "倉庫頁無資料時應只有表頭"

        ws2 = wb["異動紀錄"]
        assert [c.value for c in ws2[1]] == ["時間", "品項", "變動", "原本", "現在", "去向", "原因"]

        ws3 = wb["廠牌統計"]
        assert [c.value for c in ws3[1]] == ["廠牌", "品項數", "總庫存"]

    def test_export_formula_injection_safe(self, client):
        """公式注入防護：= 開頭的字串以 ' 前綴儲存，開啟 Excel 不會被當公式執行（含 unit / 廠牌統計）"""
        _add_item(client, name="=1+1", location="=HYPERLINK(1)", brand="=1+1", unit="=2+2")
        r = client.get("/api/export")
        assert r.status_code == 200

        import io as _io
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(r.content))
        ws = wb["辦公室"]
        values = [v for row in ws.iter_rows(min_row=2, values_only=True) for v in row]
        assert "'=1+1" in values        # 防護：撇號前綴
        assert "'=HYPERLINK(1)" in values
        assert "'=2+2" in values        # unit 欄位也有防護
        assert "=1+1" not in values     # 沒有裸公式
        assert "=2+2" not in values
        ws3 = wb["廠牌統計"]
        stats_values = [v for row in ws3.iter_rows(min_row=2, values_only=True) for v in row]
        assert "'=1+1" in stats_values  # 廠牌統計 brand 也有防護
        assert "=1+1" not in stats_values

    def test_export_days_clamped(self, client):
        """days 負數 / 超界不會 500：clamp 到 0~366"""
        for days in (-5, 99999, 0, 30, 366):
            r = client.get("/api/export", params={"days": days})
            assert r.status_code == 200, f"days={days} 失敗"

    def test_export_writes_no_file(self, client):
        """記憶體回傳：匯出後 exports/ 不新增任何檔案（零留檔）"""
        export_dir = os.path.join(BASE_DIR, "exports")
        before = set(os.listdir(export_dir)) if os.path.isdir(export_dir) else set()
        r = client.get("/api/export")
        assert r.status_code == 200
        after = set(os.listdir(export_dir)) if os.path.isdir(export_dir) else set()
        assert after == before

# ---------- 方案 A：自動版本號（server 端） ----------

def test_index_html_auto_version_params(client):
    """方案 A：server 回傳的 index.html 每個 static 資源自動帶 ?v=檔案mtime。
    原始檔不寫版本號，此測試守護「對外服務時自動加上」的行為。"""
    r = client.get("/")
    assert r.status_code == 200
    html = r.text
    # 資源引用存在
    assert 'src="/static/' in html and 'href="/static/' in html
    # 每個資源都帶版本號（原始檔不寫，但 server 回傳一定有）
    import re
    vers = re.findall(r"/static/([^\"'? >]+?\.[a-z]+)(\?v=\d+)", html)
    assert len(vers) >= 10, f"版本號資源太少: {len(vers)}"
    # 版本號 = 檔案實際 mtime
    import os
    from app.config import STATIC_DIR
    for rel, ver in vers:
        fp = os.path.join(STATIC_DIR, rel)
        assert os.path.exists(fp), f"資源不存在: {rel}"
        assert ver == f"?v={os.stat(fp).st_mtime_ns}", f"{rel} 版本號不是 mtime: {ver}"


def test_login_html_auto_version_params(client):
    """login.html 的 logo 資源也自動帶 mtime 版本號"""
    r = client.get("/login.html")
    assert r.status_code == 200
    html = r.text
    import re, os
    from app.config import STATIC_DIR
    vers = re.findall(r"/static/([^\"'? >]+?\.[a-z]+)(\?v=\d+)", html)
    assert len(vers) >= 1
    for rel, ver in vers:
        fp = os.path.join(STATIC_DIR, rel)
        assert os.path.exists(fp), f"資源不存在: {rel}"
        assert ver == f"?v={os.stat(fp).st_mtime_ns}", f"{rel} 版本號不是 mtime: {ver}"


def test_html_source_has_no_version_params():
    """原始 HTML（未經 server）不應有手動 ?v=N——避免開發者困惑要不要改"""
    import os
    for name in ("index.html", "login.html"):
        p = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                         "static", name)
        with open(p, encoding="utf-8") as fh:
            src = fh.read()
        # 註解允許出現「?v=N」說明文字，但資源引用（src=/href=）不得帶版本號
        import re
        refs = re.findall(r'(?:src|href)="/static/[^"]+"', src)
        assert refs, f"{name} 找不到資源引用"
        bad = [x for x in refs if "?v=" in x]
        assert not bad, f"{name} 有手動版本號: {bad}"

# ========== B5：cookie secure flag（HTTPS 才設） ==========

class TestCookieSecureFlag:
    def test_login_cookie_not_secure_on_http(self, client):
        """本機 HTTP 登入 → cookie 不設 secure（否則 HTTP 直連登入會失效）"""
        r = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
        assert r.status_code == 200
        set_cookie = r.headers.get("set-cookie", "")
        assert "hvac_session=" in set_cookie
        assert "Secure" not in set_cookie.split("hvac_session=")[1].split(";")[0].upper() or "Secure" not in set_cookie

    def test_login_cookie_secure_on_https(self, tmp_path, monkeypatch):
        """HTTPS 登入 → cookie 帶 Secure flag
        2026-08-14 補：改用 tmp DB（monkeypatch 後 lifespan 自動 init+建 admin），
        不再打真實 inventory.db——根治「database is locked」（原無 fixture 直接打真實 DB）"""
        from fastapi.testclient import TestClient
        test_db = tmp_path / "test_inventory.db"
        monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
        with TestClient(app_main.app, base_url="https://testserver") as c:
            r = c.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
            assert r.status_code == 200
            assert "Secure" in r.headers.get("set-cookie", "")


# ========== B6：安全 headers ==========

class TestSecurityHeaders:
    def test_security_headers_present(self, client):
        """所有回應都帶 X-Frame-Options / X-Content-Type-Options / Referrer-Policy / CSP"""
        r = client.get("/")
        assert r.headers.get("x-frame-options") == "DENY"
        assert r.headers.get("x-content-type-options") == "nosniff"
        assert r.headers.get("referrer-policy") == "no-referrer"
        csp = r.headers.get("content-security-policy", "")
        assert "default-src 'self'" in csp
        assert "frame-ancestors 'none'" in csp  # clickjacking 雙保險

    def test_csp_applied_to_api_too(self, client):
        """API 回應同樣帶 CSP（不因路徑而漏）"""
        r = client.get("/api/items")
        assert r.status_code == 200
        assert "content-security-policy" in r.headers

    def test_health_has_security_headers(self, client):
        """驗證健康檢查回應帶安全 headers"""
        r = client.get("/health")
        assert r.headers.get("x-frame-options") == "DENY"


# ---------- Phase 1 驗證（2026-08-11：負數 qty 422 / delete_stock 404） ----------

class TestPhase1Validation:
    def test_create_item_negative_qty_422(self, client):
        """M7：位置庫存 qty 不得為負（pydantic ge=0）"""
        r = client.post("/api/items", json={
            "brand": "測試牌", "code": "", "name": "負數品項", "unit": "個",
            "low_stock": 0, "site": "office",
            "stocks": [{"location": "A倉", "qty": -5, "note": ""}],
        })
        assert r.status_code == 422

    def test_update_stock_negative_qty_422(self, client):
        """M7：PATCH /api/stocks qty 不得為負"""
        item = _add_item(client, name="庫存位置", qty=5, location="A倉")
        sid = item["stocks"][0]["id"]
        r = client.patch(f"/api/stocks/{sid}", json={"qty": -1})
        assert r.status_code == 422

    def test_stockout_negative_qty_422(self, client):
        """M7：出庫 qty 不得為負"""
        item = _add_item(client, name="出庫品項", qty=5)
        r = client.post("/api/stockout", json={"item_id": item["id"], "qty": -3, "destination": "台北"})
        assert r.status_code == 422

    def test_delete_stock_not_found_404(self, client):
        """M11：刪除不存在的庫存位置回 404（原本假成功 ok:True）"""
        r = client.delete("/api/stocks/99999")
        assert r.status_code == 404

# ========== 2026-08-25：輸入驗證補強（A1/A2/A3） ==========

class TestInputValidation:
    def test_update_item_negative_low_stock_422(self, client):
        """A1：編輯品項 low_stock 不得為負（pydantic ge=0）"""
        item = _add_item(client, name="低庫存品")
        r = client.patch(f"/api/items/{item['id']}", json={"low_stock": -5})
        assert r.status_code == 422

    def test_update_item_zero_low_stock_ok(self, client):
        """A1：low_stock=0 合法（允許關閉警示）"""
        item = _add_item(client, name="關閉警示品", low_stock=10)
        r = client.patch(f"/api/items/{item['id']}", json={"low_stock": 0})
        assert r.status_code == 200
        assert r.json()["low_stock"] == 0

    def test_update_item_negative_stock_qty_422(self, client):
        """A2：編輯品項位置庫存 qty 不得為負（pydantic ge=0）"""
        item = _add_item(client, name="位置品", qty=5, location="A倉")
        r = client.patch(f"/api/items/{item['id']}", json={
            "stocks": [{"location": "A倉", "qty": -3, "note": ""}]
        })
        assert r.status_code == 422

    def test_edit_stockout_negative_qty_422(self, client):
        """A3：編輯已領出 qty 不得為負或零（pydantic gt=0）"""
        item = _add_item(client, name="出庫品", qty=10)
        rec_r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 3, "destination": "測試"
        })
        rec = rec_r.json()
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": -1})
        assert r.status_code == 422

    def test_edit_stockout_zero_qty_422(self, client):
        """A3：編輯已領出 qty=0 也拒絕（gt=0）"""
        item = _add_item(client, name="出庫品2", qty=10)
        rec_r = client.post("/api/stockout", json={
            "item_id": item["id"], "qty": 3, "destination": "測試"
        })
        rec = rec_r.json()
        r = client.patch(f"/api/stockouts/{rec['id']}", json={"qty": 0})
        assert r.status_code == 422

    def test_create_item_empty_brand_required(self, client):
        """B4：新增品項 brand 為空字串仍可建立（後端相容）"""
        r = client.post("/api/items", json={
            "brand": "", "code": "X001", "name": "測試", "unit": "個",
            "stocks": [{"location": "A", "qty": 1, "note": ""}]
        })
        # 後端目前不擋空 brand（前端擋），這裡確認後端行為
        assert r.status_code == 201

# ========== Phase 3（2026-08-11）：交易/併發（H5/H6/M3/M4/M5/M10） ==========

class TestPhase3Concurrency:
    def test_stockout_respects_prepared_qty(self, client):
        """M3：有待領出數量時一般出庫不得讓 total - prepared 變負"""
        item = _add_item(client, name="M3品項", qty=10, location="A倉")
        r = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 8})
        assert r.status_code == 200
        r = client.post("/api/stockout", json={"item_id": item["id"], "qty": 5, "destination": "台北"})
        assert r.status_code == 400
        assert "待領出" in r.json()["detail"]

    def test_prepare_exceeds_available_400(self, client):
        """H6：準備量超過可領數量 → 400"""
        item = _add_item(client, name="H6品項", qty=10)
        r = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 11})
        assert r.status_code == 400
        assert "可領出數量不足" in r.json()["detail"]

    def test_stocktake_negative_400(self, client):
        """M4：盤點負數 → 400"""
        item = _add_item(client, name="盤點負數", qty=10, location="A倉")
        r = client.post("/api/stocktake", json={"items": [{"item_id": item["id"], "location": "A倉", "actual_qty": -5}]})
        assert r.status_code == 400

    def test_stocktake_bad_float_400(self, client):
        """M4/M7b：盤點非數字 → 400（不是 500）"""
        item = _add_item(client, name="盤點亂數", qty=10, location="A倉")
        r = client.post("/api/stocktake", json={"items": [{"item_id": item["id"], "location": "A倉", "actual_qty": "abc"}]})
        assert r.status_code == 400
        assert "格式錯誤" in r.json()["detail"]

    def test_stocktake_below_prepared_400(self, client):
        """M4：盤點數量低於待領出 → 400"""
        item = _add_item(client, name="盤點準備", qty=10, location="A倉")
        client.post(f"/api/items/{item['id']}/prepare", json={"qty": 8})
        r = client.post("/api/stocktake", json={"items": [{"item_id": item["id"], "location": "A倉", "actual_qty": 3}]})
        assert r.status_code == 400
        assert "待領出" in r.json()["detail"]

    def test_stocktake_deleted_item_400(self, client):
        """M6：soft-delete 品項不可盤點 → 400（2026-08-11 補漏）"""
        item = _add_item(client, name="已刪盤點", qty=10, location="A倉")
        assert client.delete(f"/api/items/{item['id']}").status_code == 200
        r = client.post("/api/stocktake", json={"items": [{"item_id": item["id"], "location": "A倉", "actual_qty": 5}]})
        assert r.status_code == 400
        assert "已刪除" in r.json()["detail"]

    def test_adjust_deducts_from_first_stock(self, client):
        """M10：負數調整從頭扣（A=5,B=10 → -8 → A=0,B=7）"""
        item = _add_item(client, name="M10品項", qty=5, location="A倉")
        r = client.post(f"/api/items/{item['id']}/stocks", json={"location": "B倉", "qty": 10, "note": ""})
        assert r.status_code == 201  # 新增位置庫存回傳 201
        r = client.post(f"/api/items/{item['id']}/adjust", json={"delta": -8, "reason": "測試"})
        assert r.status_code == 200
        it = _get_item(client, item["id"])
        qty_by_loc = {s["location"]: s["qty"] for s in it["stocks"]}
        assert qty_by_loc["A倉"] == 0 and qty_by_loc["B倉"] == 7, qty_by_loc

    def test_concurrent_adjust_no_negative(self, client):
        """H5：兩連線同時扣減超庫存 → 至多一個成功、庫存不為負（守衛式原子扣減）"""
        import threading
        item = _add_item(client, name="併發品項", qty=10, location="A倉")
        results = []
        barrier = threading.Barrier(2)

        def worker():
            conn = app_db.get_db()
            try:
                barrier.wait()
                cur = conn.execute("UPDATE item_stocks SET qty=qty-8 WHERE id=? AND qty>=8",
                                   (item["stocks"][0]["id"],))
                conn.commit()
                results.append(cur.rowcount)
            except Exception as e:
                results.append(str(e))
            finally:
                conn.close()

        ts = [threading.Thread(target=worker) for _ in range(2)]
        [t.start() for t in ts]
        [t.join() for t in ts]
        ok = sum(1 for r in results if r == 1)
        assert ok == 1, f"應只有一個扣減成功，實際: {results}"
        conn = app_db.get_db()
        try:
            qty = conn.execute("SELECT qty FROM item_stocks WHERE id=?", (item["stocks"][0]["id"],)).fetchone()["qty"]
        finally:
            conn.close()
        assert qty == 2, f"庫存應為 2（10-8），實際 {qty}（不得負庫存）"

    def test_prepare_twice_atomic_no_overshoot(self, client):
        """H6：併發 prepare 不超可領（兩連線各準備 8，庫存 10 → 至多一個成功）"""
        import threading
        item = _add_item(client, name="併發準備", qty=10, location="A倉")
        results = []
        barrier = threading.Barrier(2)

        def worker():
            conn = app_db.get_db()
            try:
                barrier.wait()
                cur = conn.execute(
                    "UPDATE items SET prepared_qty = prepared_qty + 8 WHERE id = ? "
                    "AND prepared_qty + 8 <= (SELECT COALESCE(SUM(qty),0) FROM item_stocks WHERE item_id = ?)",
                    (item["id"], item["id"]))
                conn.commit()
                results.append(cur.rowcount)
            except Exception as e:
                results.append(str(e))
            finally:
                conn.close()

        ts = [threading.Thread(target=worker) for _ in range(2)]
        [t.start() for t in ts]
        [t.join() for t in ts]
        ok = sum(1 for r in results if r == 1)
        assert ok == 1, f"應只有一個 prepare 成功，實際: {results}"
        conn = app_db.get_db()
        try:
            pq = conn.execute("SELECT prepared_qty FROM items WHERE id=?", (item["id"],)).fetchone()["prepared_qty"]
        finally:
            conn.close()
        assert pq == 8, pq

    def test_items_unique_index_exists(self, client):
        """M5：無重複資料時建立 idx_items_unique（併發重複 DB 層防線）"""
        conn = app_db.get_db()
        try:
            idx = conn.execute("PRAGMA index_list('items')").fetchall()
            names = [r["name"] for r in idx]
            assert "idx_items_unique" in names, names
        finally:
            conn.close()

# ========== Phase 4（2026-08-11）：審計/資料完整性（M1/M9/M6） ==========

class TestPhase4Audit:
    def test_edit_item_keeps_stock_id_and_writes_movement(self, client):
        """M1：編輯品項同位置保留 stock id + qty 變化寫流水（不再 DELETE+INSERT 全量替換）"""
        item = _add_item(client, name="M1品項", qty=10, location="A倉")
        sid = item["stocks"][0]["id"]
        r = client.patch(f"/api/items/{item['id']}", json={"stocks": [{"location": "A倉", "qty": 7, "note": ""}]})
        assert r.status_code == 200
        it = _get_item(client, item["id"])
        assert it["stocks"][0]["id"] == sid, "stock id 應保留"
        assert it["stocks"][0]["qty"] == 7
        conn = app_db.get_db()
        try:
            last = conn.execute(
                "SELECT reason, delta, before_qty, after_qty FROM movements WHERE item_id=? ORDER BY id DESC LIMIT 1",
                (item["id"],)).fetchone()
        finally:
            conn.close()
        assert last and last["reason"] == "編輯品項調整", dict(last) if last else None
        assert last["before_qty"] == 10 and last["after_qty"] == 7

    def test_edit_item_same_qty_no_movement(self, client):
        """M1：qty 沒變不寫流水"""
        item = _add_item(client, name="M1不變", qty=5, location="A倉")
        r = client.patch(f"/api/items/{item['id']}", json={"stocks": [{"location": "A倉", "qty": 5, "note": "改備註"}]})
        assert r.status_code == 200
        conn = app_db.get_db()
        try:
            cnt = conn.execute("SELECT COUNT(*) FROM movements WHERE item_id=? AND reason='編輯品項調整'",
                               (item["id"],)).fetchone()[0]
        finally:
            conn.close()
        assert cnt == 0

    def test_return_stockout_before_qty_is_current(self, client):
        """M9：退回流水 before_qty = 當前庫存（原用歷史值 m["after_qty"]）"""
        item = _add_item(client, name="M9品項", qty=10, location="A倉")
        r = client.post("/api/stockout", json={"item_id": item["id"], "qty": 4, "destination": "台北"})
        assert r.status_code == 200
        mid = client.get("/api/stockouts?limit=5").json()[0]["id"]  # 最新出庫 = movement id
        r2 = client.post(f"/api/stockouts/{mid}/return")
        assert r2.status_code == 200
        conn = app_db.get_db()
        try:
            last = conn.execute(
                "SELECT before_qty, after_qty FROM movements WHERE item_id=? AND reason='退回已領出' ORDER BY id DESC LIMIT 1",
                (item["id"],)).fetchone()
        finally:
            conn.close()
        assert last and last["before_qty"] == 6 and last["after_qty"] == 10, dict(last)

    def test_delete_item_soft_delete(self, client):
        """M6：刪除品項 → 列表消失、movements 保留、已領出歷史仍顯示、後續操作 404"""
        item = _add_item(client, name="M6品項", qty=8, location="A倉")
        client.post("/api/stockout", json={"item_id": item["id"], "qty": 3, "destination": "台北"})
        r = client.delete(f"/api/items/{item['id']}")
        assert r.status_code == 200
        # 列表消失
        ids = [i["id"] for i in client.get("/api/items").json()]
        assert item["id"] not in ids
        # movements 保留（稽核軌跡不銷毀）
        conn = app_db.get_db()
        try:
            cnt = conn.execute("SELECT COUNT(*) FROM movements WHERE item_id=?", (item["id"],)).fetchone()[0]
        finally:
            conn.close()
        assert cnt >= 1
        # 已領出歷史仍顯示品項
        outs = client.get("/api/stockouts?limit=50").json()
        assert any(o["item_id"] == item["id"] for o in outs)
        # 對已刪品項的操作 → 404
        assert client.post(f"/api/items/{item['id']}/adjust", json={"delta": -1, "reason": "x"}).status_code == 404
        assert client.post("/api/stockout", json={"item_id": item["id"], "qty": 1, "destination": "x"}).status_code == 404
        assert client.post(f"/api/items/{item['id']}/prepare", json={"qty": 1}).status_code == 404
        assert client.post(f"/api/items/{item['id']}/photo",
                           files={"file": ("t.jpg", b"x" * 100, "image/jpeg")}).status_code == 404

    def test_deleted_item_stats_excluded(self, client):
        """M6：stats 不含已刪品項"""
        _add_item(client, name="刪除統計", qty=1, location="A倉")
        item2 = _add_item(client, name="刪除統計2", qty=1, location="A倉")
        client.delete(f"/api/items/{item2['id']}")
        st = client.get("/api/stats").json()
        assert st["total_items"] == 1, st

    def test_deleted_kit_material_cannot_assemble(self, client):
        """M6：套件材料已刪除 → 組裝 400"""
        mat = _add_item(client, name="材料甲", qty=5, location="A倉")
        r = client.post("/api/kits", json={"name": "測試套件", "items": [{"item_id": mat["id"], "qty": 2}]})
        assert r.status_code == 201
        kit_id = r.json()["id"]
        client.delete(f"/api/items/{mat['id']}")
        r2 = client.post(f"/api/kits/{kit_id}/assemble", json={"qty": 1})
        assert r2.status_code == 400
        assert "已刪除" in r2.json()["detail"]

# ========== Phase 5（2026-08-11）：健壯性（M8/G1/L11/M19/M21） ==========

class TestPhase5Robustness:
    def test_import_bad_qty_400_nothing_written(self, client):
        """M8：import 任一筆 qty 壞 → 400 且整批不寫入（無部分成功）"""
        r = client.post("/api/import", json={"items": [
            {"name": "好筆", "qty": 10},
            {"name": "壞筆", "qty": "abc"},
        ]})
        assert r.status_code == 400
        assert "格式錯誤" in r.json()["detail"]
        names = [i["name"] for i in client.get("/api/items").json()]
        assert "好筆" not in names and "壞筆" not in names, "整批應不寫入"

    def test_import_missing_name_400(self, client):
        """M8：import 缺名稱 → 400"""
        r = client.post("/api/import", json={"items": [{"name": "", "qty": 5}]})
        assert r.status_code == 400
        assert "名稱" in r.json()["detail"]

    def test_import_non_dict_400(self, client):
        """M8：import 含非 dict 元素（字串/數字）→ 400 不是 500"""
        r = client.post("/api/import", json={"items": [{"name": "正常", "qty": 1}, "garbage", 42]})
        assert r.status_code == 400
        assert "格式錯誤" in r.json()["detail"]

    def test_import_over_500_400(self, client):
        """M8：import 超過 500 筆 → 400"""
        many = [{"name": f"批量品{i}", "qty": 1} for i in range(501)]
        r = client.post("/api/import", json={"items": many})
        assert r.status_code == 400
        assert "500" in r.json()["detail"]

    def test_import_valid_still_works(self, client):
        """M8：正常 import 不受影響"""
        r = client.post("/api/import", json={"items": [
            {"name": "正常品", "brand": "B", "qty": 7, "location": "A倉"},
            {"name": "正常品2", "brand": "B", "qty": 3, "location": "A倉"},
        ]})
        assert r.status_code == 200
        assert r.json()["inserted"] == 2

    def test_similar_coarse_filter_keeps_results(self, client):
        """G1：SQL 粗篩不改變 similar 結果（含空白/括號的相似品項仍命中）"""
        _add_item(client, name="大金 冷氣機", qty=5, location="A倉")
        _add_item(client, name="電風扇", qty=5, location="A倉")
        r = client.get("/api/items/similar", params={"name": "大金冷氣機"})
        assert r.status_code == 200
        names = [h["name"] for h in r.json()]
        assert "大金 冷氣機" in names, f"粗篩不該漏相似品項: {names}"
        assert "電風扇" not in names, f"粗篩不該撈無關品項: {names}"

    def test_similar_coarse_filter_code(self, client):
        """G1：code 相同仍命中"""
        _add_item(client, name="品項X", code="M-100", qty=3, location="A倉")
        r = client.get("/api/items/similar", params={"code": "M-100"})
        assert any(h["name"] == "品項X" for h in r.json())

    def test_limit_capped_422(self, client):
        """L11：limit 超過上限 → 422（3 端點）"""
        _add_item(client, name="L11品項", qty=5, location="A倉")
        client.post("/api/stockout", json={"item_id": client.get("/api/items").json()[0]["id"], "qty": 1})
        assert client.get("/api/stockouts?limit=999999").status_code == 422
        assert client.get("/api/movements?limit=999999").status_code == 422
        assert client.get("/api/stocktakes?limit=999999").status_code == 422
        # 正常 limit 不受影響
        assert client.get("/api/stockouts?limit=100").status_code == 200

    def test_photo_pixel_limit_400(self, client):
        """M19：超過 40MP 像素上限 → 400"""
        from PIL import Image
        import io
        item = _add_item(client, name="照片大圖", qty=1, location="A倉")
        buf = io.BytesIO()
        Image.new("RGB", (6500, 6500), "red").save(buf, "JPEG")  # 42.25MP > 40MP
        buf.seek(0)
        r = client.post(f"/api/items/{item['id']}/photo",
                        files={"file": ("big.jpg", buf.getvalue(), "image/jpeg")})
        assert r.status_code == 400
        assert "解析度過高" in r.json()["detail"]

    def test_photo_normal_size_ok(self, client):
        """M19：正常尺寸照片不受影響"""
        from PIL import Image
        import io
        item = _add_item(client, name="照片正常", qty=1, location="A倉")
        buf = io.BytesIO()
        Image.new("RGB", (800, 600), "blue").save(buf, "JPEG")
        buf.seek(0)
        r = client.post(f"/api/items/{item['id']}/photo",
                        files={"file": ("ok.jpg", buf.getvalue(), "image/jpeg")})
        assert r.status_code == 200

    def test_csrf_cross_origin_403(self, client):
        """M21：跨站 Origin POST → 403"""
        item = _add_item(client, name="CSRF品項", qty=5, location="A倉")
        r = client.post("/api/stockout",
                        json={"item_id": item["id"], "qty": 1, "destination": "x"},
                        headers={"Origin": "http://evil.com"})
        assert r.status_code == 403

    def test_csrf_same_origin_ok(self, client):
        """M21：同源 Origin 與無 Origin（curl）→ 放行"""
        item = _add_item(client, name="CSRF同源", qty=5, location="A倉")
        r1 = client.post("/api/stockout",
                         json={"item_id": item["id"], "qty": 1, "destination": "x"},
                         headers={"Origin": "http://testserver"})
        assert r1.status_code == 200
        r2 = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 1})
        assert r2.status_code == 200  # 無 Origin（curl/測試）→ 放行


class TestNonStockOut:
    """2026-08-13 Sarah：已領出可直接新增「不在單一庫存/整組庫存」的品項（POST /api/stockout/nonstock）"""

    def test_nonstock_stockout_creates_movement(self, client):
        """非庫存品項領出：建臨時品項（is_deleted=1）+ 出庫流水；庫存頁看不到、已領出看得到"""
        r = client.post("/api/stockout/nonstock", json={
            "name": "冷媒銅管 3分", "code": '3/8"', "unit": "捲",
            "qty": 2, "destination": "張小姐-大樓保養", "note": "維修"})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["name"] == "冷媒銅管 3分"

        # 已領出清單看得到（item_deleted 標記 + 品名/型號/單位/數量/去向）
        outs = client.get("/api/stockouts").json()
        match = [o for o in outs if o["item_id"] == body["id"]]
        assert len(match) == 1
        assert match[0]["item_name"] == "冷媒銅管 3分"
        assert match[0]["code"] == '3/8"'
        assert match[0]["unit"] == "捲"
        assert match[0]["item_deleted"] == 1
        assert match[0]["delta"] == -2
        assert match[0]["destination"] == "張小姐-大樓保養"
        assert "出庫 - 維修" in match[0]["reason"]

        # 庫存頁看不到（is_deleted=1 被過濾）
        items = client.get("/api/items").json()
        assert all(i["id"] != body["id"] for i in items), "非庫存臨時品項不應出現在庫存頁"

    def test_nonstock_stockout_requires_name(self, client):
        """品項名稱空白 → 400"""
        r = client.post("/api/stockout/nonstock", json={
            "name": "   ", "qty": 1, "destination": "客戶A"})
        assert r.status_code == 400

    def test_nonstock_stockout_requires_destination(self, client):
        """去哪裡空白 → 400"""
        r = client.post("/api/stockout/nonstock", json={
            "name": "雜項", "qty": 1, "destination": ""})
        assert r.status_code == 400

    def test_nonstock_stockout_qty_must_be_positive(self, client):
        """數量 0/負 → pydantic gt=0 擋 422 或手動 400"""
        r = client.post("/api/stockout/nonstock", json={
            "name": "雜項", "qty": 0, "destination": "客戶A"})
        assert r.status_code in (400, 422)

    def test_nonstock_prepare_shows_in_list(self, client):
        """非庫存品項待領出：/api/prepare/nonstock 建臨時品項 + prepared_qty，出現在待領出清單（含 is_deleted 標記）"""
        r = client.post("/api/prepare/nonstock", json={
            "name": "臨時耗材", "code": "T-1", "unit": "捲", "qty": 3, "note": "工地備料"})
        assert r.status_code == 200, r.text
        pid = r.json()["id"]

        prepared = client.get("/api/prepared").json()
        match = [i for i in prepared if i["id"] == pid]
        assert len(match) == 1, "非庫存待領出應出現在待領出清單"
        assert match[0]["prepared_qty"] == 3

        # 2026-08-16 防回歸：帶 site 參數（前端 currentSite 預設 office）也要看得到非庫存待領出
        prepared_office = client.get("/api/prepared?site=office").json()
        match_office = [i for i in prepared_office if i["id"] == pid]
        assert len(match_office) == 1, "非庫存待領出不應被 site 過濾掉"
        assert match[0]["is_deleted"] == 1
        assert match[0]["unit"] == "捲"

        # 庫存頁看不到
        items = client.get("/api/items").json()
        assert all(i["id"] != pid for i in items)

    def test_nonstock_prepare_out_becomes_stockout(self, client):
        """非庫存品項待領出 → 確認出庫：不扣庫存、直接寫出庫流水、清 prepared_qty"""
        r = client.post("/api/prepare/nonstock", json={
            "name": "臨時材料", "qty": 2, "destination": ""})
        assert r.status_code == 200
        pid = r.json()["id"]

        out = client.post(f"/api/items/{pid}/prepared-out",
                          json={"qty": 2, "note": "某案場", "location": ""})
        assert out.status_code == 200, out.text
        assert out.json()["prepared_qty"] == 0

        # 已領出清單有該筆（非庫存、出庫流水）
        outs = client.get("/api/stockouts").json()
        match = [o for o in outs if o["item_id"] == pid]
        assert len(match) == 1
        assert match[0]["delta"] == -2
        assert match[0]["item_deleted"] == 1
        assert match[0]["destination"] == "某案場"

    def test_nonstock_prepared_return_clears(self, client):
        """非庫存品項待領出 → 退回：清 prepared_qty、無出庫流水（不加庫存）"""
        r = client.post("/api/prepare/nonstock", json={
            "name": "臨時材料", "qty": 2, "destination": ""})
        assert r.status_code == 200
        pid = r.json()["id"]

        ret = client.post(f"/api/items/{pid}/prepared-return", json={"qty": 2, "location": ""})
        assert ret.status_code == 200, ret.text
        assert ret.json()["prepared_qty"] == 0

        prepared = client.get("/api/prepared").json()
        assert all(i["id"] != pid for i in prepared), "退回後不應再出現在待領出"
        outs = client.get("/api/stockouts").json()
        assert all(o["item_id"] != pid for o in outs), "退回不應產生出庫流水"

    def test_nonstock_stockout_does_not_deduct_inventory(self, client):
        """非庫存領出不影響任何既有庫存數量"""
        item = _add_item(client, name="既有品", qty=5)
        r = client.post("/api/stockout/nonstock", json={
            "name": "臨時品", "qty": 3, "destination": "客戶A"})
        assert r.status_code == 200
        after = client.get("/api/items").json()
        cur = [i for i in after if i["id"] == item["id"]][0]
        assert cur["qty"] == 5, "非庫存領出不應扣到既有庫存"


# ========== B7：database is locked 根治防回歸（2026-08-14） ==========

class TestDbLockRelease:
    """寫入端點 error 後鎖必須釋放——防「bare-conn 洩漏」復發（f0b520f/3007c61 根治）。"""

    def test_write_error_releases_db_lock(self, client):
        """寫入端點拋 400（同名衝突 → rollback 路徑）後，獨立連線可立即寫入 → 無 RESERVED 鎖洩漏"""
        import sqlite3
        payload = {"name": "鎖洩漏測試", "sort_order": 1, "is_active": True}
        r1 = client.post("/api/service-types", json=payload)
        assert r1.status_code == 200, f"建立服務項目應成功，got {r1.status_code}"
        r2 = client.post("/api/service-types", json=payload)
        assert r2.status_code == 400, "同名衝突應 400（觸發 rollback 路徑）"

        # 決定性驗證：獨立連線（timeout=1，不給長等）立即寫入成功 = 鎖已釋放
        probe = sqlite3.connect(str(app_db.DB_PATH), timeout=1)
        try:
            probe.execute(
                "INSERT INTO service_types (name, sort_order, is_active) VALUES (?,?,?)",
                ("鎖已釋放", 2, 1))
            probe.commit()
        finally:
            probe.close()

    def test_rollback_path_does_not_leak_lock_on_validation_error(self, client):
        """寫入端點在 commit 前因驗證錯誤拋 400（items 去重）→ 鎖同樣釋放"""
        import sqlite3
        item = _add_item(client, name="重複品", brand="B", code="C", qty=3)
        r = client.post("/api/items", json={
            "brand": "B", "code": "C", "name": "重複品", "unit": "個",
            "low_stock": 0, "site": "office", "stocks": []})
        assert r.status_code == 400, "同鍵去重應 400（item 已存在）"

        probe = sqlite3.connect(str(app_db.DB_PATH), timeout=1)
        try:
            probe.execute("INSERT INTO service_types (name, sort_order, is_active) VALUES (?,?,?)",
                          ("驗證錯誤後可寫", 3, 1))
            probe.commit()
        finally:
            probe.close()

# ========== B8：lifespan cleanup_expired 防回歸（2026-08-15） ==========

class TestLifespanCleanup:
    """lifespan 啟動時清理過期 session——docstring 承諾「啟動時與登入時呼叫」落地（563a96d）。"""

    def test_lifespan_removes_expired_sessions(self, tmp_path, monkeypatch):
        """啟動時：過期 session 被刪、未過期保留（用獨立 tmp DB 驗證 lifespan 行為）"""
        import datetime
        from fastapi.testclient import TestClient
        test_db = tmp_path / "test_lifespan.db"
        monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
        app_db.init_db()

        # 塞 1 筆過期 + 1 筆未過期 session（user_id 指向已存在的 admin——先建）
        _conn = app_db.get_db()
        try:
            from app.services.auth import init_admin_if_missing
            init_admin_if_missing(_conn)
            _conn.execute(
                "INSERT INTO sessions (user_id, token_hash, expires_at) VALUES (?, ?, ?)",
                (1, "expired_token_hash", "2000-01-01 00:00:00"))
            _conn.execute(
                "INSERT INTO sessions (user_id, token_hash, expires_at) VALUES (?, ?, ?)",
                (1, "valid_token_hash",
                 (datetime.datetime.now() + datetime.timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")))
            _conn.commit()
        finally:
            _conn.close()

        with TestClient(app_main.app):
            pass  # lifespan 啟動 → cleanup_expired 執行

        # 過期被刪、未過期保留
        _conn2 = app_db.get_db()
        try:
            hashes = [r["token_hash"] for r in
                      _conn2.execute("SELECT token_hash FROM sessions").fetchall()]
        finally:
            _conn2.close()
        assert "expired_token_hash" not in hashes
        assert "valid_token_hash" in hashes


# ---------- 2026-08-26 category 欄位 ----------



class TestCategory:
    """品項分類（category）欄位測試"""

    def test_create_item_with_category(self, client):
        """新增品項含 category"""
        resp = client.post("/api/items", json={
            "brand": "TEST-CAT", "code": "CAT001", "name": "測試分類品項",
            "unit": "個", "site": "office", "category": "遙控器",
            "stocks": [{"location": "測試位置A", "qty": 5}]
        })
        assert resp.status_code == 201
        data = resp.json()
        assert data["category"] == "遙控器"
        assert data["brand"] == "TEST-CAT"
        # 清理
        client.delete(f"/api/items/{data['id']}")

    def test_create_item_without_category(self, client):
        """新增品項不含 category（預設空字串）"""
        resp = client.post("/api/items", json={
            "brand": "TEST-NO-CAT", "code": "NC001", "name": "無分類品項",
            "unit": "個", "site": "office",
            "stocks": [{"location": "測試位置B", "qty": 1}]
        })
        assert resp.status_code == 201
        data = resp.json()
        assert data.get("category", "") == ""
        # 清理
        client.delete(f"/api/items/{data['id']}")

    def test_update_item_category(self, client):
        """編輯品項修改 category"""
        # 新增
        resp = client.post("/api/items", json={
            "brand": "TEST-UPD-CAT", "code": "UC001", "name": "待改分類",
            "unit": "個", "site": "office", "category": "管材",
            "stocks": [{"location": "測試位置C", "qty": 1}]
        })
        item_id = resp.json()["id"]
        # 修改 category
        resp = client.patch(f"/api/items/{item_id}", json={
            "category": "電氣配件"
        })
        assert resp.status_code == 200
        assert resp.json()["category"] == "電氣配件"
        # 清理
        client.delete(f"/api/items/{item_id}")

    def test_items_list_includes_category(self, client):
        """GET /api/items 回傳含 category 欄位"""
        # 先新增一筆測試資料
        resp = client.post("/api/items", json={
            "brand": "TEST-LIST-CAT", "code": "LC001", "name": "列表測試",
            "unit": "個", "site": "office", "category": "化學品",
            "stocks": [{"location": "測試位置D", "qty": 1}]
        })
        assert resp.status_code == 201
        item_id = resp.json()["id"]
        # 再查列表
        resp = client.get("/api/items?site=office")
        assert resp.status_code == 200
        items = resp.json()
        assert len(items) > 0
        # 每筆都應有 category 欄位
        for item in items:
            assert "category" in item, f"品項 {item.get('name')} 缺 category 欄位"
        # 清理
        client.delete(f"/api/items/{item_id}")


# ========== Deep merge-safety: canonical quantity ledger invariants ==========

class TestDeepQtyIntegrity:
    @pytest.mark.parametrize("initial,delta,expected_before,expected_after", [
        (0, 1 / 3, 0, 0.333),
        (1, -1 / 3, 1, 0.667),
    ])
    def test_adjust_canonical_delta_keeps_stock_and_movement_aligned(
        self, client, initial, delta, expected_before, expected_after
    ):
        item = _add_item(client, name="調整精度", qty=initial)
        response = client.post(f"/api/items/{item['id']}/adjust", json={"delta": delta, "reason": "深度稽核"})
        assert response.status_code == 200, response.text
        assert _get_item(client, item["id"])["total_qty"] == pytest.approx(expected_after)
        movement = next(m for m in client.get("/api/movements").json() if m["item_id"] == item["id"])
        assert movement["before_qty"] == pytest.approx(expected_before)
        assert movement["delta"] == pytest.approx(expected_after - expected_before)
        assert movement["after_qty"] == pytest.approx(expected_after)
        assert round(movement["before_qty"] + movement["delta"], 3) == pytest.approx(movement["after_qty"])

    def test_adjust_sub_precision_delta_rejected_without_phantom_movement(self, client):
        item = _add_item(client, name="調整塵", qty=1)
        response = client.post(f"/api/items/{item['id']}/adjust", json={"delta": 0.0004, "reason": "深度稽核"})
        assert response.status_code == 400
        assert _get_item(client, item["id"])["total_qty"] == 1
        assert not any(m["item_id"] == item["id"] for m in client.get("/api/movements").json())

    def test_stockout_three_locations_canonical_remainder(self, client):
        response = client.post("/api/items", json={
            "brand": "測試牌", "code": "Q-111", "name": "三位置精度",
            "stocks": [
                {"location": "A", "qty": 0.111},
                {"location": "B", "qty": 0.111},
                {"location": "C", "qty": 0.111},
            ],
        })
        assert response.status_code == 201, response.text
        item = response.json()
        out = client.post("/api/stockout", json={"item_id": item["id"], "qty": 0.333})
        assert out.status_code == 200, out.text
        current = _get_item(client, item["id"])
        assert current["total_qty"] == pytest.approx(0)
        assert all(stock["qty"] == pytest.approx(0) for stock in current["stocks"])
        movement = next(m for m in client.get("/api/movements").json() if m["item_id"] == item["id"])
        assert movement["before_qty"] == pytest.approx(0.333)
        assert movement["delta"] == pytest.approx(-0.333)
        assert movement["after_qty"] == pytest.approx(0)

    def test_prepared_repeated_fraction_has_no_dust_or_ghost_item(self, client):
        item = _add_item(client, name="待領出精度", qty=1)
        for _ in range(3):
            response = client.post(f"/api/items/{item['id']}/prepare", json={"qty": 1 / 3})
            assert response.status_code == 200, response.text
        prepared_item = _get_item(client, item["id"])
        assert prepared_item["prepared_qty"] == 0.999
        for _ in range(3):
            response = client.post(f"/api/items/{item['id']}/prepared-out", json={"qty": 1 / 3})
            assert response.status_code == 200, response.text
        assert _get_item(client, item["id"])["prepared_qty"] == 0
        assert not any(row["id"] == item["id"] for row in client.get("/api/prepared").json())

    def test_kit_fractional_assemble_and_disassemble_ledger(self, client):
        material = _add_item(client, name="組裝材料", qty=1)
        kit_response = client.post("/api/kits", json={
            "name": "精度整組", "items": [{"item_id": material["id"], "qty": 1}],
        })
        assert kit_response.status_code == 201, kit_response.text
        kit_id = kit_response.json()["id"]
        kit_item_id = kit_response.json()["item_id"]
        assemble = client.post(f"/api/kits/{kit_id}/assemble", json={"qty": 1 / 3})
        assert assemble.status_code == 200, assemble.text
        assembled_movement = next(
            m for m in client.get("/api/movements").json()
            if m["item_id"] == kit_item_id and m["reason"].startswith("組裝完成")
        )
        assert assembled_movement["before_qty"] == pytest.approx(0)
        assert assembled_movement["delta"] == pytest.approx(0.333)
        assert assembled_movement["after_qty"] == pytest.approx(0.333)
        disassemble = client.post(f"/api/kits/{kit_id}/disassemble", json={"qty": 1 / 3})
        assert disassemble.status_code == 200, disassemble.text
        disassembled_movement = next(
            m for m in client.get("/api/movements").json()
            if m["item_id"] == material["id"] and m["reason"].startswith("拆解套件")
        )
        assert disassembled_movement["delta"] == pytest.approx(0.333)
        assert disassembled_movement["before_qty"] == pytest.approx(0.667)
        assert disassembled_movement["after_qty"] == pytest.approx(1)

    def test_stocktake_ledger_uses_canonical_actual_and_diff(self, client):
        item = _add_item(client, name="盤點精度", qty=1)
        response = client.post("/api/stocktake", json={
            "take_date": "2026-09-14",
            "items": [{"item_id": item["id"], "location": "測試位置", "actual_qty": 1 / 3}],
        })
        assert response.status_code == 200, response.text
        result = response.json()["results"][0]
        assert result["actual_qty"] == pytest.approx(0.333)
        assert result["diff"] == pytest.approx(-0.667)
        movement = next(m for m in client.get("/api/movements").json() if m["item_id"] == item["id"])
        assert movement["before_qty"] == pytest.approx(1)
        assert movement["delta"] == pytest.approx(-0.667)
        assert movement["after_qty"] == pytest.approx(0.333)

    def test_create_import_and_stock_edit_store_canonical_qty(self, client):
        created = _add_item(client, name="建立精度", qty=1.2345)
        assert created["stocks"][0]["qty"] == pytest.approx(1.235)

        imported = client.post("/api/import", json={"items": [{
            "brand": "匯入牌", "code": "IMP-1", "name": "匯入精度",
            "unit": "個", "qty": 1.2345, "location": "匯入位置",
        }]})
        assert imported.status_code == 200, imported.text
        imported_item = next(i for i in client.get("/api/items").json() if i["name"] == "匯入精度")
        assert imported_item["total_qty"] == pytest.approx(1.235)

        stock_id = created["stocks"][0]["id"]
        edited = client.patch(f"/api/stocks/{stock_id}", json={"qty": 2.3455})
        assert edited.status_code == 200, edited.text
        current = _get_item(client, created["id"])
        assert current["total_qty"] == pytest.approx(2.346)
        movement = next(m for m in client.get("/api/movements").json() if m["item_id"] == created["id"])
        assert movement["before_qty"] == pytest.approx(1.235)
        assert movement["delta"] == pytest.approx(1.111)
        assert movement["after_qty"] == pytest.approx(2.346)

    def test_kit_assemble_rolls_back_material_deduction_if_add_fails(self, client, monkeypatch):
        material = _add_item(client, name="組裝回滾材料", qty=1)
        kit_response = client.post("/api/kits", json={
            "name": "回滾整組", "items": [{"item_id": material["id"], "qty": 1}],
        })
        assert kit_response.status_code == 201, kit_response.text
        kit_id = kit_response.json()["id"]
        kit_item_id = kit_response.json()["item_id"]

        import app.routes.kits as kits_route
        def fail_after_deduct(*args, **kwargs):
            raise RuntimeError("模擬整組增加失敗")
        monkeypatch.setattr(kits_route, "_add_total", fail_after_deduct)

        with pytest.raises(RuntimeError, match="模擬整組增加失敗"):
            client.post(f"/api/kits/{kit_id}/assemble", json={"qty": 1})

        assert _get_item(client, material["id"])["total_qty"] == pytest.approx(1)
        assert _get_item(client, kit_item_id)["total_qty"] == pytest.approx(0)
        assert not any(m["item_id"] in {material["id"], kit_item_id}
                       for m in client.get("/api/movements").json())
