# -*- coding: utf-8 -*-
"""庫存 Excel 匯出改版回歸測試。"""
import io
import os
import sys
import zipfile
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

import app.database as app_db  # noqa: E402
import main as app_main  # noqa: E402
from app.services import movement_time  # noqa: E402
from app.routes.export import _movement_type  # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    test_db = tmp_path / "export.db"
    monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
    app_db.init_db()
    from app.services.auth import SESSION_COOKIE, create_session, init_admin_if_missing
    conn = app_db.get_db()
    try:
        init_admin_if_missing(conn)
        admin_id = conn.execute("SELECT id FROM users WHERE username='admin'").fetchone()["id"]
        token = create_session(conn, admin_id)
    finally:
        conn.close()
    with TestClient(app_main.app) as c:
        c.cookies.set(SESSION_COOKIE, token)
        yield c


def add_item(client, *, name="測試品", brand="測試牌", code="T-1", site="office", category="電子零件", unit="個", qty=10, low_stock=2):
    response = client.post("/api/items", json={
        "name": name, "brand": brand, "code": code, "site": site,
        "category": category, "unit": unit, "low_stock": low_stock,
        "stocks": [{"location": "A櫃", "qty": qty, "note": "位置備註"}],
    })
    assert response.status_code == 201, response.text
    return response.json()


def insert_movement(item_id, created_at, *, delta=1, reason="庫存調整"):
    conn = app_db.get_db()
    try:
        conn.execute(
            "INSERT INTO movements(item_id, delta, before_qty, after_qty, reason, destination, created_at) VALUES(?,?,?,?,?,?,?)",
            (item_id, delta, 0, delta, reason, "測試去向", created_at),
        )
        conn.commit()
    finally:
        conn.close()


def export_book(client, **params):
    response = client.get("/api/export", params=params)
    assert response.status_code == 200, response.text
    return load_workbook(filename=__import__("io").BytesIO(response.content), data_only=False)


def test_export_has_expected_sheets_in_order(client):
    book = export_book(client)
    assert book.sheetnames == ["庫存總表(單一庫存)", "位置明細(單一庫存)", "異動紀錄(單一庫存)"]
    for worksheet in book.worksheets:
        assert "A1:D1" in {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
        assert worksheet["A1"].value
        assert worksheet["A1"].alignment.horizontal == "center"


def test_export_month_filters_movements_and_past_month_excludes_next_month(client):
    item = add_item(client)
    insert_movement(item["id"], "2026-09-05 10:00:00")
    insert_movement(item["id"], "2026-10-01 00:00:00", delta=2)
    book = export_book(client, month="2026-09")
    rows = list(book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True))
    assert len([r for r in rows if r[0]]) == 1
    assert rows[0][0] == "2026-09-05 10:00:00"


def test_export_invalid_date_range_returns_400(client):
    response = client.get("/api/export", params={"start_date": "2026-09-20", "end_date": "2026-09-01"})
    assert response.status_code == 400
    response = client.get("/api/export", params={"month": "2026/09"})
    assert response.status_code == 400


def test_export_formula_contract_and_tables(client):
    item = add_item(client)
    insert_movement(item["id"], "2026-09-05 10:00:00")
    book = export_book(client, month="2026-09")
    inventory = book["庫存總表(單一庫存)"]
    assert set(inventory.tables) == {"tblInventory"}
    assert set(book["位置明細(單一庫存)"].tables) == {"tblPosition"}
    assert set(book["異動紀錄(單一庫存)"].tables) == {"tblMovement"}
    headers = [c.value for c in inventory[5]]
    cells = {name: inventory.cell(6, headers.index(name) + 1).value for name in headers}
    assert cells["總庫存"] == "=SUMIFS(tblPosition[位置數量],tblPosition[品項編號(系統編號)],A6)"
    assert cells["可用庫存"] == "=I6-H6"
    assert cells["位置數"] == "=COUNTIFS(tblPosition[品項編號(系統編號)],A6)"
    assert cells["庫存狀態"].startswith("=IF(J6<0")


def test_export_multi_location_has_one_inventory_row_and_two_position_rows(client):
    item = add_item(client, name="多位置", code="MULTI", qty=1)
    response = client.post(f"/api/items/{item['id']}/stocks", json={"location": "B櫃", "qty": 2})
    assert response.status_code == 201, response.text
    book = export_book(client)
    inventory_rows = [r for r in book["庫存總表(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    position_rows = [r for r in book["位置明細(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert len(inventory_rows) == 1
    assert len(position_rows) == 2
    assert inventory_rows[0][8].startswith("=")
    assert all(len(r) == 9 for r in position_rows)


def test_export_prepared_soft_delete_and_raw_text_safety(client):
    active = add_item(client, name="=危險名稱", brand="+危險牌", code="@CODE", qty=10)
    deleted = add_item(client, name="刪除品", code="DEL", qty=4)
    conn = app_db.get_db()
    try:
        conn.execute("UPDATE items SET prepared_qty=7 WHERE id=?", (active["id"],))
        conn.execute("UPDATE items SET is_deleted=1 WHERE id=?", (deleted["id"],))
        conn.commit()
    finally:
        conn.close()
    book = export_book(client)
    rows = [r for r in book["庫存總表(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert len(rows) == 1
    assert rows[0][2].startswith("'")
    assert rows[0][3].startswith("'")
    assert rows[0][7] == 7
    assert isinstance(rows[0][8], str) and rows[0][8].startswith("=")


def test_export_decimal_qty_stays_numeric_and_uses_unit_format(client):
    add_item(client, name="小數品", code="DEC", unit="米", qty=0.333, low_stock=0)
    book = export_book(client)
    position = book["位置明細(單一庫存)"]
    qty_cell = position.cell(6, 8)
    assert isinstance(qty_cell.value, (int, float))
    assert qty_cell.value == pytest.approx(0.333)
    assert qty_cell.number_format == "#,##0.###"


def test_export_custom_date_range_and_filename(client):
    item = add_item(client)
    insert_movement(item["id"], "2026-09-17 12:00:00")
    insert_movement(item["id"], "2026-09-18 00:00:00")
    response = client.get("/api/export", params={"start_date": "2026-09-17", "end_date": "2026-09-17"})
    assert response.status_code == 200
    assert "20260917-20260917_" in response.headers["content-disposition"]
    book = load_workbook(filename=__import__("io").BytesIO(response.content), data_only=False)
    rows = [r for r in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert len(rows) == 1


def test_export_overview_and_stats_remain_optional_not_default(client):
    """回歸測試：總覽與統計可手動選取，但預設不產生。"""
    item = add_item(client, name="分類統計", category="冷媒零件", qty=6)
    conn = app_db.get_db()
    try:
        conn.execute("UPDATE items SET prepared_qty=2 WHERE id=?", (item["id"],))
        conn.commit()
    finally:
        conn.close()
    default_book = export_book(client)
    assert "01 總覽" not in default_book.sheetnames
    assert "06 統計" not in default_book.sheetnames

    optional_book = export_book(client, sections="overview,inventory,stats")
    assert optional_book.sheetnames == ["01 總覽", "庫存總表(單一庫存)", "06 統計"]
    assert optional_book["01 總覽"]["A1"].value == "庫存管理報表"
    assert optional_book["06 統計"]["A4"].value == "庫存區統計"
    assert "tblInventory" in optional_book["01 總覽"]["B6"].value
    assert optional_book["06 統計"]["J6"].value == "冷媒零件"
    assert optional_book["06 統計"]["K6"].value == 1
    assert optional_book["06 統計"]["L6"].value == 6
    assert optional_book["06 統計"]["M6"].value == 4

    for title in ("庫存總表(單一庫存)", "位置明細(單一庫存)"):
        headers = [cell.value for cell in default_book[title][5] if cell.value is not None]
        assert "分類" not in headers


def test_export_stats_have_no_24_row_ceiling_for_brands_and_categories(client):
    """統計表選取後，分類與廠牌摘要完整輸出超過 24 筆的項目。"""
    for index in range(30):
        add_item(client, name=f"統計品項{index}", brand=f"品牌{index:02d}", code=f"STAT-{index}", category=f"分類{index:02d}", qty=index + 1)
    stats = export_book(client, sections="inventory,stats")["06 統計"]
    assert [stats.cell(row, 10).value for row in range(6, 36)] == [f"分類{index:02d}" for index in range(30)]
    assert [stats.cell(row, 15).value for row in range(6, 36)] == [f"品牌{index:02d}" for index in range(30)]
    assert stats.cell(35, 11).value == 1
    assert stats.cell(35, 12).value == 30
    assert stats.cell(35, 13).value == 30
    assert stats.cell(35, 16).value.startswith("=COUNTIF(tblInventory[廠牌]")
    assert stats.cell(35, 17).value.startswith("=SUMIF(tblInventory[廠牌]")
    assert stats.cell(35, 18).value.startswith("=SUMIF(tblInventory[廠牌]")


def test_export_display_period_does_not_show_exclusive_end_as_inclusive(client):
    past = export_book(client, month="2026-08")
    assert "2026/08/01 ～ 2026/08/31" in past["庫存總表(單一庫存)"]["A2"].value
    single = export_book(client, start_date="2026-09-17", end_date="2026-09-17")
    assert "2026/09/17 ～ 2026/09/17" in single["庫存總表(單一庫存)"]["A2"].value
    assert "2026/09/18" not in single["庫存總表(單一庫存)"]["A2"].value


def test_export_contains_no_dynamic_array_functions_or_xlfn_in_formulas_and_xml(client):
    add_item(client, name="警示", code="F5-1", qty=0, low_stock=2)
    book_response = client.get("/api/export")
    assert book_response.status_code == 200
    book = load_workbook(io.BytesIO(book_response.content), data_only=False)
    forbidden = ("FILTER(", "SORT(", "UNIQUE(", "SORTBY(", "HSTACK(", "VSTACK(", "LET(", "SEQUENCE(", "TAKE(", "DROP(", "CHOOSECOLS(", "TOCOL(", "TOROW(", "XLOOKUP(", "_XLFN", "_XLWS")
    formulas = [cell.value.upper() for ws in book.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert not [formula for formula in formulas if any(token in formula for token in forbidden)]
    with zipfile.ZipFile(io.BytesIO(book_response.content)) as archive:
        xml = "\n".join(archive.read(name).decode("utf-8") for name in archive.namelist() if name.startswith("xl/worksheets/"))
    assert not any(token in xml.upper() for token in ("FILTER", "_XLFN", "_XLWS"))


def test_export_alerts_use_traditional_index_match_formulas(client):
    add_item(client, name="缺貨", code="F5-2", qty=0, low_stock=2)
    book = export_book(client, sections="inventory,positions,alerts,movements")
    inventory = book["庫存總表(單一庫存)"]
    alert = book["庫存警示(單一庫存)"]
    assert inventory.column_dimensions["M"].hidden is True
    assert inventory.cell(5, 13).value == "警示序號"
    assert "IF(" in inventory.cell(6, 13).value and "COUNTIF(" in inventory.cell(6, 13).value
    for column in range(1, 12):
        formula = alert.cell(6, column).value
        assert isinstance(formula, str) and formula.startswith("=IFERROR(INDEX(")
        assert "MATCH(ROW()-5,tblInventory[警示序號],0)" in formula
    assert "FILTER(" not in (alert.cell(6, 1).value or "")


def test_export_alert_candidates_cover_multiple_and_zero_alerts_without_errors(client):
    for index in range(5):
        add_item(client, name=f"警示{index}", code=f"F5-A{index}", qty=0, low_stock=1)
    for index in range(3):
        add_item(client, name=f"正常{index}", code=f"F5-N{index}", qty=10, low_stock=1)
    book = export_book(client, sections="inventory,positions,alerts,movements")
    alert = book["庫存警示(單一庫存)"]
    assert all(isinstance(alert.cell(row, 1).value, str) and alert.cell(row, 1).value.startswith("=IFERROR(") for row in range(6, 14))
    assert alert.max_row == 13
    empty_book = export_book(client, sites="van", sections="alerts")
    empty_alert = empty_book["庫存警示(單一庫存)"]
    assert empty_alert.cell(6, 1).value == "目前沒有資料"


def test_export_movement_integer_and_decimal_number_formats(client):
    integer_item = add_item(client, name="整數異動", code="MOVE-INT", qty=1)
    decimal_item = add_item(client, name="小數異動", code="MOVE-DEC", unit="米", qty=1.333)
    insert_movement(integer_item["id"], "2026-09-17 10:00:00", delta=-1, reason="出庫")
    insert_movement(decimal_item["id"], "2026-09-17 09:00:00", delta=0.333, reason="庫存調整")

    book = export_book(client, month="2026-09")
    movement = book["異動紀錄(單一庫存)"]
    rows = [row for row in movement.iter_rows(min_row=6) if row[0].value]
    integer_row = next(row for row in rows if row[2].value == integer_item["id"])
    decimal_row = next(row for row in rows if row[2].value == decimal_item["id"])

    assert [integer_row[col - 1].value for col in (8, 9, 10)] == [-1, 0, -1]
    assert all(integer_row[col - 1].number_format == "#,##0" for col in (8, 9, 10))
    assert isinstance(decimal_row[7].value, (int, float))
    assert decimal_row[7].value == pytest.approx(0.333)
    assert decimal_row[7].number_format == "#,##0.###"


def test_export_includes_assembled_kit_inventory(client):
    material = add_item(client, name="組裝材料", code="KIT-MAT", qty=10)
    kit_response = client.post("/api/kits", json={
        "name": "測試整組", "brand": "測試牌", "code": "KIT-001",
        "site": "office", "items": [{"item_id": material["id"], "qty": 2}],
    })
    assert kit_response.status_code == 201, kit_response.text
    kit = kit_response.json()
    assembled = client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 2})
    assert assembled.status_code == 200, assembled.text

    book = export_book(client)
    inventory_rows = [r for r in book["庫存總表(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    position_rows = [r for r in book["位置明細(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert any(row[0] == kit["item_id"] for row in inventory_rows)
    kit_positions = [row for row in position_rows if row[0] == kit["item_id"]]
    assert kit_positions and kit_positions[0][7] == 2
    kit_inventory = next(row for row in inventory_rows if row[0] == kit["item_id"])
    assert isinstance(kit_inventory[8], str) and "tblPosition" in kit_inventory[8]


def test_export_kit_assemble_disassemble_movements_are_preserved(client):
    material = add_item(client, name="拆解材料", code="KIT-MAT-2", qty=10)
    kit_response = client.post("/api/kits", json={
        "name": "組拆整組", "brand": "測試牌", "code": "KIT-002",
        "site": "office", "items": [{"item_id": material["id"], "qty": 1}],
    })
    assert kit_response.status_code == 201, kit_response.text
    kit = kit_response.json()
    assert client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 1}).status_code == 200
    assert client.post(f"/api/kits/{kit['id']}/disassemble", json={"qty": 1}).status_code == 200

    book = export_book(client)
    reasons = [row[11] for row in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    assert any(str(reason).startswith("組裝套件:") for reason in reasons)
    assert any(str(reason).startswith("組裝完成:") for reason in reasons)
    assert any(str(reason).startswith("拆解:") for reason in reasons)
    assert any(str(reason).startswith("拆解套件:") for reason in reasons)


def test_export_movements_preserve_soft_deleted_item_history(client):
    item = add_item(client, name="歷史刪除品", code="DEL-HISTORY", qty=3)
    deleted = client.delete(f"/api/items/{item['id']}")
    assert deleted.status_code == 200, deleted.text

    book = export_book(client)
    rows = [row for row in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    deleted_rows = [row for row in rows if row[2] == item["id"]]
    assert deleted_rows
    assert any(row[11] == "品項刪除清零" and row[1] == "刪除清零" for row in deleted_rows)
    assert not any(row[0] == item["id"] for row in book["庫存總表(單一庫存)"].iter_rows(min_row=6, values_only=True))


def test_export_movements_include_nonstock_stockout(client):
    response = client.post("/api/stockout/nonstock", json={
        "name": "非庫存歷史品", "code": "NONSTOCK-1", "unit": "個",
        "qty": 2, "destination": "測試案場",
    })
    assert response.status_code == 200, response.text
    item_id = response.json()["id"]

    book = export_book(client)
    rows = [row for row in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    nonstock_rows = [row for row in rows if row[2] == item_id]
    assert nonstock_rows and nonstock_rows[0][11] == "出庫"


@pytest.mark.parametrize("reason, expected", [
    ("領出準備", "待領出"), ("出庫", "出庫"), ("出庫 - 案場", "出庫"),
    ("退回已領出", "退回"), ("盤點調整", "盤點"), ("庫存調撥", "調撥"),
    ("組裝套件:測試整組", "整組組裝"), ("組裝完成:測試整組", "整組組裝"),
    ("拆解:測試整組", "整組拆解"), ("拆解套件:測試整組", "整組拆解"),
    ("品項刪除清零", "刪除清零"), ("已領出編輯調整", "庫存調整"),
])
def test_movement_type_matches_production_reason_contract(reason, expected):
    assert _movement_type(reason, -1) == expected


def test_export_custom_range_accepts_exactly_366_days_and_rejects_367(client):
    accepted = client.get("/api/export", params={"start_date": "2024-01-01", "end_date": "2024-12-31"})
    rejected = client.get("/api/export", params={"start_date": "2024-01-01", "end_date": "2025-01-01"})
    assert accepted.status_code == 200, accepted.text
    assert rejected.status_code == 400


def test_export_custom_range_uses_left_closed_right_open_boundaries(client):
    item = add_item(client, name="邊界品", code="DATE-BOUNDARY")
    insert_movement(item["id"], "2026-09-01 00:00:00", reason="庫存調整")
    insert_movement(item["id"], "2026-10-01 00:00:00", delta=2, reason="庫存調整")
    book = export_book(client, start_date="2026-09-01", end_date="2026-09-30")
    rows = [row for row in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    assert [row[0] for row in rows] == ["2026-09-01 00:00:00"]


def test_export_movement_site_prefers_return_site_over_source_site(client):
    item = add_item(client, name="退回來源品", code="RETURN-SITE")
    conn = app_db.get_db()
    try:
        conn.execute(
            "INSERT INTO movements(item_id, delta, before_qty, after_qty, reason, destination, "
            "created_at, source_site, return_site) VALUES(?,?,?,?,?,?,?,?,?)",
            (item["id"], 1, 0, 1, "退回已領出", "倉庫", "2026-09-15 12:00:00", "office", "warehouse"),
        )
        conn.commit()
    finally:
        conn.close()
    book = export_book(client, month="2026-09", sites="warehouse")
    rows = [row for row in book["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    assert rows and rows[0][6] == "倉庫"


@pytest.mark.parametrize("raw, expected", [
    ("2026-10-01", "2026-10-01 00:00:00"),
    ("2026-10-01 08:30", "2026-10-01 08:30:00"),
    ("2026-10-01 08:30:45", "2026-10-01 08:30:45"),
    ("2026-10-01T08:30:45", "2026-10-01 08:30:45"),
])
def test_normalize_user_datetime(raw, expected):
    assert movement_time.normalize_user_datetime(raw) == expected


@pytest.mark.parametrize("raw", ["2026-02-31", "2026-10-01 25:00", "2026-10-01T08:30:00+08:00", "garbage"])
def test_normalize_user_datetime_rejects_invalid_or_timezone(raw):
    with pytest.raises(ValueError):
        movement_time.normalize_user_datetime(raw)


def test_stockout_created_at_uses_taipei_business_time_and_export_boundary(client, monkeypatch):
    monkeypatch.setattr(movement_time, "now_sql", lambda: "2026-10-01 00:30:00")
    item = add_item(client, name="跨月出庫", code="TZ-001", qty=10)
    response = client.post("/api/stockout", json={
        "item_id": item["id"], "qty": 1, "destination": "測試案場", "location": "A櫃",
    })
    assert response.status_code == 200, response.text
    conn = app_db.get_db()
    try:
        movement = conn.execute(
            "SELECT created_at FROM movements WHERE item_id=? AND reason LIKE '出庫%' ORDER BY id DESC LIMIT 1",
            (item["id"],),
        ).fetchone()
    finally:
        conn.close()
    assert movement["created_at"] == "2026-10-01 00:30:00"
    october = export_book(client, start_date="2026-10-01", end_date="2026-10-01")
    september = export_book(client, start_date="2026-09-01", end_date="2026-09-01")
    oct_rows = [r for r in october["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    sep_rows = [r for r in september["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if r[0]]
    assert any(r[2] == item["id"] and r[0] == "2026-10-01 00:30:00" for r in oct_rows)
    assert not any(r[2] == item["id"] and r[0] == "2026-10-01 00:30:00" for r in sep_rows)


def test_stockout_user_entered_created_at_is_normalized_and_validated(client):
    item = add_item(client, name="日期編輯品", code="TZ-EDIT", qty=2)
    response = client.post("/api/stockout", json={
        "item_id": item["id"], "qty": 1, "destination": "測試案場", "location": "A櫃",
    })
    assert response.status_code == 200, response.text
    conn = app_db.get_db()
    try:
        movement_id = conn.execute(
            "SELECT id FROM movements WHERE item_id=? AND reason LIKE '出庫%' ORDER BY id DESC LIMIT 1",
            (item["id"],),
        ).fetchone()["id"]
    finally:
        conn.close()
    updated = client.patch(f"/api/stockouts/{movement_id}", json={"created_at": "2026-10-01T08:30"})
    assert updated.status_code == 200, updated.text
    assert updated.json()["created_at"] == "2026-10-01 08:30:00"
    invalid = client.patch(f"/api/stockouts/{movement_id}", json={"created_at": "2026-02-31"})
    assert invalid.status_code == 400


def test_transfer_uses_single_timestamp_for_both_movements(client, monkeypatch):
    item = add_item(client, name="跨午夜調撥", code="TRANSFER-TIME", site="office", qty=10)
    timestamps = iter([
        "2026-09-30 23:59:59",
        "2026-10-01 00:00:00",
        "2026-10-01 00:00:01",
    ])
    original_now_sql = movement_time.now_sql
    monkeypatch.setattr(movement_time, "now_sql", lambda: next(timestamps))
    response = client.post("/api/inventory/transfers", json={
        "item_id": item["id"], "qty": 2, "target_site": "warehouse",
        "source_location": "A櫃", "target_location": "B櫃",
    })
    assert response.status_code == 201, response.text
    conn = app_db.get_db()
    try:
        rows = conn.execute(
            "SELECT item_id, delta, reason, created_at FROM movements "
            "WHERE reason='庫存調撥' ORDER BY id"
        ).fetchall()
    finally:
        conn.close()
    assert len(rows) == 2
    assert {row["created_at"] for row in rows} == {"2026-09-30 23:59:59"}
    monkeypatch.setattr(movement_time, "now_sql", original_now_sql)

    september = export_book(client, start_date="2026-09-30", end_date="2026-09-30")
    october = export_book(client, start_date="2026-10-01", end_date="2026-10-01")
    sep_rows = [row for row in september["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    oct_rows = [row for row in october["異動紀錄(單一庫存)"].iter_rows(min_row=6, values_only=True) if row[0]]
    assert len([row for row in sep_rows if row[11] == "庫存調撥"]) == 2
    assert not [row for row in oct_rows if row[11] == "庫存調撥"]


def test_export_rejects_empty_or_unknown_sections(client):
    """拒絕空的工作表選取內容及未知的工作表識別碼。"""
    assert client.get("/api/export", params={"sections": ""}).status_code == 400
    assert client.get("/api/export", params={"sections": "inventory,unknown"}).status_code == 400


def test_export_sections_default_to_single_inventory_and_keep_alerts_optional(client):
    """預設匯出三張單一庫存表，警示只在明確選取時輸出。"""
    default_book = export_book(client)
    assert default_book.sheetnames == [
        "庫存總表(單一庫存)",
        "位置明細(單一庫存)",
        "異動紀錄(單一庫存)",
    ]

    alert_book = export_book(client, sections="alerts")
    assert alert_book.sheetnames == ["庫存警示(單一庫存)"]


def test_export_body_cells_use_text_format_and_center_alignment(client):
    """文字格式與置中對齊適用於 A3 以後，不覆寫標題列或數量格式。"""
    add_item(client, name="格式對齊", code="ALIGN", qty=2)
    book = export_book(client, sections="inventory,positions,movements")
    for sheet in book.worksheets:
        assert sheet["A1"].font.name == "Microsoft JhengHei"
        assert sheet["A2"].font.name == "Microsoft JhengHei"
        assert sheet["A1"].number_format == "General"
        assert sheet["A2"].number_format == "General"
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                if cell.value is None:
                    continue
                assert cell.alignment.horizontal == "center"
                assert cell.alignment.vertical == "center"
                assert cell.font.name == "Microsoft JhengHei"
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    assert cell.number_format == "@"
                if cell.number_format == "General":
                    raise AssertionError(f"{sheet.title}!{cell.coordinate} kept General format")
    assert book["位置明細(單一庫存)"].cell(6, 8).number_format == "#,##0"
    assert book["庫存總表(單一庫存)"].cell(6, 12).number_format == "@"


def test_export_single_inventory_headers_styles_and_columns(client):
    """單一庫存匯出移除分類並保留 Excel 可用的欄名、格式與欄寬。"""
    add_item(client, name="格式驗證", code="W-7", qty=12)
    book = export_book(client, sections="inventory,positions,alerts,movements")
    inventory = book["庫存總表(單一庫存)"]
    positions = book["位置明細(單一庫存)"]
    alerts = book["庫存警示(單一庫存)"]
    movements = book["異動紀錄(單一庫存)"]

    assert inventory["A1"].value == "單一庫存總表"
    assert positions["A1"].value == "單一庫存位置明細"
    assert alerts["A1"].value == "單一庫存警示"
    assert movements["A1"].value == "單一庫存異動紀錄"
    assert "庫存快照" not in " ".join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value)

    inventory_headers = [cell.value for cell in inventory[5] if cell.value is not None]
    position_headers = [cell.value for cell in positions[5] if cell.value is not None]
    alert_headers = [cell.value for cell in alerts[5] if cell.value is not None]
    movement_headers = [cell.value for cell in movements[5] if cell.value is not None]
    assert inventory_headers[0] == "品項編號(系統編號)"
    assert "分類" not in inventory_headers
    assert "分類" not in position_headers
    assert "分類" not in alert_headers
    assert position_headers[0] == "品項編號(系統編號)"
    assert alert_headers[1] == "品項編號(系統編號)"
    assert movement_headers[2] == "品項編號(系統編號)"
    assert "tblPosition[品項編號(系統編號)]" in inventory.cell(6, 9).value
    assert inventory.cell(6, 13).value.startswith("=IF(")

    assert inventory["A1"].font.name == "Microsoft JhengHei"
    assert inventory["A1"].font.bold
    assert all(inventory.cell(5, col).font.bold for col in range(1, len(inventory_headers) + 1))
    assert inventory["B6"].number_format == "@"
    assert inventory["G6"].number_format == "#,##0"
    assert inventory.column_dimensions["A"].width < 15
    assert inventory["B6"].value == "公司"
    assert "庫存快照" not in str(inventory["A2"].value)
