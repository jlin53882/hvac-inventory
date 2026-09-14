# -*- coding: utf-8 -*-
"""Dynamic petty-cash Excel rendering regression tests."""
from pathlib import Path

from openpyxl import load_workbook

from app.services.engineering_petty_cash import build_engineering_report
from app.services.petty_cash_report import build_petty_cash_report


ROOT = Path(__file__).parents[1]


def _load_bytes(builder, report, tmp_path, name):
    path = tmp_path / name
    path.write_bytes(builder(report).getvalue())
    return load_workbook(path, data_only=False).active


def _engineering_report(categories, start="2026-09-01", end="2026-09-04"):
    total = sum(
        receipt["amount"]
        for category in categories
        for group in category.get("groups", [])
        for receipt in group.get("receipts", [])
    )
    return {
        "start_date": start,
        "end_date": end,
        "categories": categories,
        "total_amount": total,
    }


def test_engineering_excel_has_period_title_and_dynamic_body(tmp_path):
    """工程報表期間標題與 body 列數不應受原始範本固定列數影響。"""
    categories = [
        {
            "name": "交通費",
            "groups": [
                {
                    "name": "油資",
                    "receipts": [
                        {"tax_id_mark": "V", "receipt_number": "DH-1", "amount": 100, "details": ["九二無鉛"]},
                        {"tax_id_mark": "12345678", "receipt_number": "DH-2", "amount": 725, "details": ["九二無鉛"]},
                    ],
                }
            ],
        },
        {
            "name": "工程材料費",
            "groups": [
                {
                    "name": "五金/工具",
                    "receipts": [
                        {"tax_id_mark": "V", "receipt_number": "CB 20011829", "amount": 1004, "details": [f"材料 {i}" for i in range(6)]},
                    ],
                }
            ],
        },
    ]
    ws = _load_bytes(build_engineering_report, _engineering_report(categories), tmp_path, "engineering.xlsx")

    assert ws.title == "0901-0904"
    assert ws["A1"].value == "09/01~09/04工程零用金明細表"
    assert [ws.cell(2, col).value for col in range(1, 7)] == ["類別", "項目", "統編", "發票號碼", "細項", "金額"]
    assert ws["F5"].value == 825
    assert ws["F12"].value == 1004
    assert ws["F14"].value == 1829
    assert "C6:C11" in [str(value) for value in ws.merged_cells.ranges]
    assert "D6:D11" in [str(value) for value in ws.merged_cells.ranges]
    assert "F6:F11" in [str(value) for value in ws.merged_cells.ranges]
    assert ws.calculate_dimension() == "A1:F16"
    assert str(ws.print_area).endswith("$A$1:$F$16")
    assert all(ws.cell(row, col).value is None for row in range(17, ws.max_row + 1) for col in range(1, 7))


def test_engineering_excel_single_day_period_is_not_repeated(tmp_path):
    """單日工程報表只顯示一次日期，sheet tab 也使用單日 token。"""
    categories = [{
        "name": "交通費",
        "groups": [{"name": "油資", "receipts": [{"amount": 100, "details": ["九二無鉛"]}]}],
    }]
    ws = _load_bytes(
        build_engineering_report,
        _engineering_report(categories, "2026-09-04", "2026-09-04"),
        tmp_path,
        "engineering-single.xlsx",
    )
    assert ws.title == "0904"
    assert ws["A1"].value == "09/04工程零用金明細表"
    assert "09/04~09/04" not in ws["A1"].value


def test_general_excel_rebuilds_body_after_more_than_template_capacity(tmp_path):
    """一般零用金超過原本 15 列時仍只保留實際 body/footer。"""
    entries = [
        {
            "entry_date": f"2026-09-{day:02d}",
            "entry_type": "expense",
            "description": f"支出 {day}",
            "amount": day,
            "category": "測試",
            "items": [],
        }
        for day in range(1, 21)
    ]
    report = {
        "start_date": "2026-09-01",
        "end_date": "2026-09-20",
        "opening_balance": 1000,
        "prepared_by": "王小明",
        "entries": entries,
    }
    ws = _load_bytes(build_petty_cash_report, report, tmp_path, "general.xlsx")

    assert ws.title == "0901-0920"
    assert ws["A1"].value == "09/01~09/20零用金收支明細表"
    assert ws["D24"].value == "本期餘額"
    assert ws["E24"].value == 790
    assert ws["E27"].value == "王小明"
    assert ws.calculate_dimension() == "A1:F27"
    assert str(ws.print_area).endswith("$A$1:$F$27")
