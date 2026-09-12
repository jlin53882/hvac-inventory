# -*- coding: utf-8 -*-
"""
安全回歸守門測試（2026-08-12 新增）
====================================
防止「每次新加功能就帶進安全漏洞」的掃描式 + 行為式測試。

為什麼需要：
  - Phase 0-6 + af13870 修復後，f3c8880（手機 UI）/ 36306ea（整組編輯）/
    行事曆新功能仍各帶進 stored XSS / 負 qty 假流水 / 公式注入——
    因為既有測試只守「功能行為」，不守「安全模式」（esc 使用 / 輸入驗證存在性）。
  - 本檔不看 git 歷史、掃「目前檔案系統」→ 新 render 檔、新端點自動被檢查。

執行：env -u PYTHONPATH .venv/Scripts/python.exe -m pytest tests/test_security_regression.py -v
"""
import glob
import os
import re
import sqlite3
import sys

import pytest
from fastapi.testclient import TestClient

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

import app.database as app_db  # noqa: E402
import main as app_main  # noqa: E402


# ---------- fixture（比照 test_main.py：每個測試獨立 tmp DB + admin 登入） ----------
@pytest.fixture()
def client(tmp_path, monkeypatch):
    test_db = tmp_path / "test_inventory.db"
    monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
    app_db.init_db()

    from app.services.auth import init_admin_if_missing
    _conn = app_db.get_db()
    try:
        init_admin_if_missing(_conn)
    finally:
        _conn.close()

    with TestClient(app_main.app) as c:
        r = c.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
        assert r.status_code == 200, f"測試 admin 登入失敗: {r.status_code} {r.text}"
        # 清掉 fixture 登入可能殘留的 per-IP 失敗計數（TestClient 共用 IP "testclient"）
        from app.services.auth import clear_ip_fail
        clear_ip_fail("testclient")
        yield c

    try:
        if test_db.exists():
            test_db.unlink()
    except PermissionError:
        pass


# ============================================================
# 1. 掃描式：所有 HTML 模板內插必須 esc/jsStr（或明確安全）
# ============================================================


# 已人工審核的「安全內插」白名單（2026-08-12 baseline 全檔審核）：
# - 數字/算術/布林/常數三元/程式內變數/內部 HTML 參數（呼叫端已消毒）
# - 新檔案/新內插若不在清單 → 測試紅 → 人工審核（安全則加這裡，否則補 esc()）
REVIEWED_SAFE_BODIES = {
    # 報價單內部分頁（2026-09-09）：active 只由固定模式傳入，輸出皆為固定 class/文字。
    "active === 'quotation' ? 'active' : ''", "active === 'upload' ? 'active' : ''", "quoteModeTabs('quotation')", "quoteModeTabs('upload')",
    # bottomsheet.js（動作選單：icon/label 為開發者傳入常數；items 為內部 map HTML）
    "a.icon", "icon", "a.label", "items",
    # 數字欄位（qty/id/統計）
    "s.qty ?? 0", "s.qty", "itemId", "i.qty", "row.qty || 1", "it.qty",
    "c.need_qty", "c.stock", "k.stock_qty", "kits.length", "items.length",
    "totalPrepared", "i.prepared_qty", "list.length", "m", "totalOut",
    "o.item_id", "absNum(o.delta)", "absNum(s.qty)", "val",
    # kits.js 整組待領出/已領出按鈕（2026-08-14 工作區既有改動）：k.item_id 為 DB 數字主鍵（同 c.item_id/o.item_id）
    "k.item_id",
    "prevTotal - i + 1", "locItems.length", "counts[b] || ALL_ITEMS.length",
    "last.diff_count", "last.item_count", "last.total_diff",
    "d.diff_count", "d.item_count", "d.total_diff", "ALL_ITEMS.length",
    "s.sort_order", "totalQtyStr", "low", "zero", "s.qty", "id",
    # 布林三元（常數輸出或含 esc 分支）
    "u.is_active ? '⏸ 帳號停用' : '▶️ 帳號啟用'", "u.is_active ? '✅ 啟用' : '⛔ 停用'",
    # sync_status（2026-08-27 gcal sync：server 端產生的固定字串，非使用者輸入）
    "e.sync_status === 'synced' ? '✅' : e.sync_status === 'pending' ? '⏳' : e.sync_status === 'failed' ? '❌' : ''",
    "e.sync_status === 'synced' ? '已同步到 Google 行事曆' : e.sync_status === 'pending' ? '等待同步' : e.sync_status === 'failed' ? '同步失敗' : '未綁定同步 Key'",
    "e.sync_status === 'synced' ? '已同步到 Google 行事曆' : e.sync_status === 'partial_failed' ? '部分同步失敗' : e.sync_status === 'pending' ? '等待同步' : e.sync_status === 'failed' ? '同步失敗' : '未綁定同步 Key'",
    "e.sync_status === 'synced' ? '✅' : e.sync_status === 'partial_failed' ? '⚠️' : e.sync_status === 'pending' ? '⏳' : e.sync_status === 'failed' ? '❌' : ''",
    # settings.js（2026-08-16 設定中心）：u.is_active 為 DB bool 常數輸出（同帳號頁模式）；u.count 為 COUNT(*) 數字（已 esc）
    "u.is_active ? 'checked' : ''",
    # settings.js 方案 B 逐筆收編（2026-08-16）：it.item_id 為 DB 整數主鍵（同 c.item_id 先例）、
    # absNum(it.total_qty) 為格式化數字字串、g.items.length 為陣列長度（同 items.length 先例）；
    # 其餘內插（esc/布林三元含 HTML 常數）由掃描器自動放行
    "it.item_id",
    "absNum(it.total_qty)",
    "g.items.length",
    "isLow ? '🎉 沒有低庫存品項' : '🎉 沒有缺貨品項'",
    # card.js mobileCardShell cardClass（2026-09-07 Phase 5：CSS class 為開發者傳入常數）
    "p.cardClass ? ' ' + p.cardClass : ''", "isLow ? '警示值' : '位置'",
    "isLow ? '#92400e' : '#991b1b'", "isLow ? '#fffbeb' : '#fef2f2'",
    "isLow ? '⚠️ 低庫存品項' : '⛔ 缺貨品項'", "isLow ? (i.low_stock || 0) : esc(locStr || '—')",
    "enough ? 'ok' : 'low'", "enough ? '#16a34a' : '#dc2626'",
    "locCollapsed ? ' collapsed' : ''", "i.site === 'warehouse' ? ' 🏭' : ''",
    "reverted ? ' style=\"opacity:0.55\"' : ''", "p.reverted ? ' reverted' : ''",
    # stockout.js 退回紀錄列（2026-09-07 Sarah：退回顯示在已領出頁）
    "isReturn ? ' style=\"background:#f6ffed\"' : ''",
    "isReturn ? 'qty-pos' : 'qty-neg'",
    "isReturn ? '+' : '-'",
    "f && f.user_ids.includes(p.id) ? 'checked' : ''",
    "f && f.service_type_id === s.id ? 'selected' : ''",
    "s.is_active ? 'on' : ''", "low ? 'warn' : ''", "zero ? 'danger' : ''",
    # 內部變數/函式回傳（內部已消毒或格式化）
    "roleClass", "roleLabel", "ROLE_LABELS[u.role] || u.role",
    "calFmtCreatedAt(e.created_at)", "calModalHtml(isAdmin)", "calSettingsHtml(isAdmin)",
    "_fmtTW(new Date())", "todayStr()", "who", "locHtml", "pPhoto", "soPhoto",
    "displayLoc || '未標示'", "display", "cls ? ' ' + cls : ''",
    # stocktake.js 盤點輸入表 tab（2026-08-13）：stkGroupByLoc 內部所有使用者資料已 esc()
    "stkGroupByLoc(kitRows)", "stkGroupByLoc(singleRows)",
    # calendar.js 24 制時間下拉（2026-08-13）：${hh}/${mm} 為 String(h/mm).padStart 產生的純數字（00-23/00-55），不可控
    "hh", "mm",
    # perms.js 權限頁（2026-08-13）：class 三元 / ROLE_LABELS・GROUP_LABELS 常數輸出，使用者資料已 esc()
    "u.id === curUid ? 'active' : ''", "badge", "u.is_active ? '啟用中' : '已停用'",
    "u.is_active ? '' : 'off'", "u.is_active ? 'on' : 'off'", "isMe ? '（自己）' : ''",
    "u.is_active ? '✅ 啟用中' : '⏸ 已停用'", "GROUP_LABELS[mod] || mod", "locked ? 'locked' : ''",
    "srcCls", "srcLabel", "checked", "disabled ? 'disabled' : ''", "isMe ? 'disabled' : ''",
    "Object.keys(permChanges).length ? 'changed' : ''", "u.is_active ? 0 : 1",
    "u.is_active ? '⏸ 停用帳號' : '▶️ 啟用帳號'",
    "x.status === 'ok' ? 'ok' : 'err'", "x.status === 'ok' ? '✔' : '✘'",
    # card.js 共用元件參數（呼叫端傳入已消毒 HTML）
    "p.moreBtnHTML || ''", "p.nameHTML", "p.thumb", "p.actionsHTML || ''",
    "p.extraHTML || ''", "p.qtyHTML",
    # 工程／一般零用金：helper 內部對使用者資料已 esc，回傳固定 HTML 結構
    "engOptionSelect('category', c.name)", "engOptionSelect('group', g.name)", "pcGeneralCategoryOptions(src.category || '')",
    "seq", "receiptCells", "totalReceipts",
    # 含 esc 的組合內插
    "i.code ? ' · ' + esc(i.code) : ''", "sel ? esc(sel.brand) + ' ' + esc(sel.name) : ''",
    "s.note ? ' · 📝 ' + esc(s.note) : ''",
    # stocktake.js 多行三元提示文字（isLow ? '常數提示' : '常數提示'）
    "isLow\n        ? '💡 庫存數量已低於（或等於）警示值，建議盡快補貨。點品項可直接編輯警示值。'\n        : '💡 庫存為 0 或以下的品項，需要補貨或盤點確認。'",
    # 每日簽名報表（2026-09-07）：ic.* 只由 _dsrIconFor 固定映射產生；note 已在同一行以 esc(r.note) 處理，空值是固定 HTML。
    "ic.bg", "ic.icon", "note", "fileVisual",
    # stocktake.js 整組盤點展開組成材料（2026-08-16）：c.item_id 為 DB 數字主鍵（同 k.item_id/o.item_id）、
    # c.has_photo 為布林控制縮圖/佔位三元、kitCompsHTML 為內部已 esc 的組裝 HTML（同 pPhoto/soPhoto 模式）
    "c.item_id", "c.has_photo", "kitCompsHTML",
    # stocktake status modal：buildThumb/extra 由既有安全 helper 與 esc(item.in_kits) 建構 HTML
    "buildThumb(item.id, item.has_photo)", "extra",
    # stockout.js：row/card/group fragments are assembled from escaped fields and fixed numeric IDs
    "returned", "destination", "returnLocation", "renderStockoutActions(o, isViewer)",
    "records.map(function(o) { return renderStockoutDesktopRow(o,", "desktopRows", "body", "stockoutBar",
    # stocktake.js 組成品項實際數量輸入框（2026-08-16）：mKey/mVal 為 itemId:location 鍵與輸入值（數字或空）、
    # mSysQty 為 absNum 格式化的數字字串（純數字不可控）
    "mKey", "mVal", "mSysQty",
    # stocktake.js 盤點表分組渲染（2026-08-26 搜尋過濾）：stkGroupByLoc 回傳內部已 esc 的 HTML，
    # filteredKitRows/filteredSingleRows 為 kitRows/singleRows 過濾後陣列（同 kitCompsHTML 安全模式）
    "stkGroupByLoc(filteredKitRows)", "stkGroupByLoc(filteredSingleRows)",
    # stocktake.js row renderer（2026-09-09）：fragment 先逐欄 esc/jsStr，class/data-kind 為固定值。
    "materialPhoto", "materialName", "stocktakeInput(materialKey, materialSystemQty)",
    "rowClass", "item.is_kit ? 'assembly' : 'single'", "displayLoc",
    "stock.note ? ' · 📝 ' + esc(stock.note) : ''",
    "stocktakeInput(key, systemQty)",
    "stocktakeValues[key] === undefined || stocktakeValues[key] === '' ? 'pending' : 'zero'",
    "stocktakeValues[key] === undefined || stocktakeValues[key] === '' ? '—' : '0'",
    "materials",
    # edit.js 兩段式位置（2026-09-06）：_cabinetOptions 從固定清單產生 select options，
    # c 為固定 cabs 陣列元素（編號A~F/鐵架/二樓），selected 為屬性三元，均非使用者輸入
    "_cabinetOptions(cabinet)", "_cabinetOptions('')", "c",
    "c === selected ? 'selected' : ''", "c || '— 請選擇 —'",
    # calendar.js Desktop dashboard（2026-09-09）：service/sync/updated/range 是由已 esc 的資料組成的內部 HTML fragment；icon 是固定映射。
    "calSyncStatusIcon(e.sync_status)", "service", "sync", "updated", "range",
    # calendar.js right-panel search (2026-09-09): fragments are assembled only from esc() values or fixed markup.
    "assignees", "assigneeHtml", "addressHtml", "noteHtml", "keyword", "body", "meta",
    # card.js 批次改位置 checkbox（2026-09-06）：selectedStockIds 為內部 Set、i.id 為 DB 數字主鍵
    "selectedStockIds.has(i.stocks && i.stocks.length ? i.stocks[0].id : 0) ? 'checked' : ''",
    "batchMode ? `<input type=\"checkbox\"",
    "p.checkboxHTML || ''",
    "_allSelected() ? '☐ 取消全選' : '☑ 全選'",
    "batchMode ? 'style=\"padding-left:32px\"' : ''",
    # media URL helper：id 來自 DB 數字主鍵，variant/thumbnailUrl 僅為內部固定變體。
    "photoSrc(itemId, 'thumbnail')", "photoSrc(itemId, 'preview')",
    "photoSrc(c.item_id, 'thumbnail')", "photoSrc(i.id, 'thumbnail')", "photoSrc(item.id, 'thumbnail')",
    "photoSrc(o.item_id, 'thumbnail')", "thumbnailUrl || photoSrc(id, 'thumbnail')",
    # inventory.js desktop redesign（2026-09-09）：條件文字為固定 UI；數值由格式化 helper 產生。
    # kits.js dashboard（2026-09-09）：以下為 renderer 內部已 esc 的 HTML 組合或 DB 數字主鍵。
    "renderKitStatusBadge(status.status)", "renderKitActionButtons(k, isViewer, isM, status)",
    "components.map(renderKitComponentRow).join('')",
    "hasFilter ? '沒有符合條件的庫存品項' : '目前沒有庫存品項'",
    "hasFilter ? '可以嘗試清除篩選或調整搜尋條件。' : '新增品項後，庫存與位置會在這裡集中管理。'",
    "clearButton", "addButton", "formatInventoryQuantity(totalQty)", "lowCount", "zeroCount",
    "formatInventoryQuantity(item.low_stock)", "isOut ? 'is-out' : 'is-low'", "thumb",
    "item.code ? ' · 型號 ' + esc(item.code) : ''", "item.code ? '型號： ' + esc(item.code) : ''", "formatInventoryQuantity(status.qty)",
    "threshold", "editAction", "empty", "title", "intro", "listHTML",
    # prepared.js desktop redesign：photo/nonStock/actions 是由數字主鍵與已跳脫欄位組成的內部 HTML。
    "itemCount", "absNum(totalPrepared)", "item.prepared_qty", "photo", "nonStock",
    "absNum(item.prepared_qty)", "absNum(item.qty)", "actions",
    # 共用 status-list renderer：rows/locationFilter/extraHTML 是已 esc 的內部 fragment；buildThumb 統一處理 URL/placeholder。
    "locationFilter", "rows", "columnHeadings", "buildThumb(kit.item_id, !!source.has_photo, kit.name, '🔧', source.thumbnail_url)",
    "buildThumb(item.id, item.has_photo, item.name, '📦', item.thumbnail_url)", "config.extraHTML || ''",
    "missingHTML", "statusListFormatQuantity(stock)", "statusListFormatQuantity(status.qty)",
    # 2026-09-12 數量系統：Qty.disp 輸出僅數字/分數字元（0-9 . / - 空格），無 HTML metachars；
    # stocktakeInput 內部對 key/value/sysqty/unit 全 esc()/jsStr()（stocktake.js）
    "(typeof Qty !== 'undefined') ? Qty.disp(s.qty, item.unit) : s.qty",
    "(typeof Qty !== 'undefined') ? Qty.disp(item.prepared_qty, item.unit) : absNum(item.prepared_qty)",
    "(typeof Qty !== 'undefined') ? Qty.disp(item.qty, item.unit) : absNum(item.qty)",
    "stocktakeInput(materialKey, materialSystemQty, c.unit)",
    "stocktakeInput(key, systemQty, item.unit)",
    # 零用金月報（2026-09-12）：以下皆為內部已 esc 的 HTML fragment、固定映射或常數三元——
    # _pcPeriodText/_pcFileLabel/pcItemText 回傳 esc() 組裝字串；amt/incomeCell/expenseCell 為 esc(_pcMoney()) 數字或空字串；
    # ops/editBtn/delBtn/rowsHtml/itemsHtml 由 esc() 欄位 + DB 數字主鍵 + 固定 markup 組成；
    # pcStatusBadge 為 draft/completed 固定映射；id/pcEntryType 三元只輸出固定文字/class（id 為數字主鍵）。
    "_pcPeriodText(r)", "_pcFileLabel(r)", "pcStatusBadge(r.status)",
    "ops", "editBtn", "delBtn", "rowsHtml", "itemsHtml",
    "amt", "incomeCell", "expenseCell",
    "pcItemText(first)", "pcItemText(items[i])",
    "id ? '✏️ 編輯零用金月報' : '＋ 新增零用金月報'",
    "pcEntryType === 'income' ? ' active' : ''", "pcEntryType === 'expense' ? ' active' : ''",
    "pcEntryType === 'income' ? 'display:none' : ''",
}


def test_js_html_templates_interpolations_escaped():
    """掃描 static/js/**：HTML 模板（含 < 或 >）內的內插必須 esc()/jsStr() 或
    在已審核白名單（REVIEWED_SAFE_BODIES）。

    新 render 檔一進 repo 就被掃——bottomsheet.js（f3c8880）與
    calendar.js（行事曆）類型的 stored XSS 回歸會被本測試擋下。"""
    js_root = os.path.join(BASE_DIR, "static", "js")
    bad = []
    for f in sorted(glob.glob(os.path.join(js_root, "**", "*.js"), recursive=True)):
        rel = os.path.relpath(f, BASE_DIR).replace("\\", "/")
        src = open(f, encoding="utf-8").read()
        for m in re.finditer(r"`([^`]*)`", src):
            seg = m.group(1)
            if "<" not in seg and ">" not in seg:
                continue  # 非 HTML 模板（textContent 等）不檢查
            for im in re.finditer(r"\$\{([^}]*)\}", seg):
                body = im.group(1).replace("\r", "").strip()  # CRLF 檔的 \r 先去掉，白名單統一 LF
                if body.startswith(("esc(", "jsStr(")):
                    continue  # 已跳脫
                if body in REVIEWED_SAFE_BODIES:
                    continue  # 已審核安全
                if re.fullmatch(r"[\d\s+\-*/().\[\]]+", body):
                    continue  # 純算術/數字
                if re.fullmatch(r"\w+\.id", body):
                    continue  # 數字主鍵（e.id、p.id、u.id…）
                if re.fullmatch(r"\w+\.length", body):
                    continue  # 陣列長度（數字）
                if body in ("n", "d", "cls", "idx", "i", "m", "val", "low", "zero"):
                    continue  # 單字變數（索引/計數）
                if "e.message" in body:
                    continue  # 本專案僅含 HTTP 狀態碼/瀏覽器原生訊息
                if "<" in body or ">" in body:
                    continue  # 三元輸出 HTML 常數（內部巢狀內插會被當獨立模板掃描）
                line = src.count("\n", 0, m.start()) + 1
                bad.append(f"  {rel}:{line}  ${{{body[:60]}}}")
    assert not bad, "發現未跳脫且未審核的 HTML 內插點（stored XSS 風險）：\n" + "\n".join(bad)


# ============================================================
# 2. 掃描式：uploads 公開封鎖矩陣（大小寫/穿越/前綴）
# ============================================================

def test_uploads_block_path_matrix():
    """_is_upload_path 必須封鎖：大小寫變體（Windows 不分大小寫）/ 反斜線 / 穿越。"""
    from main import _is_upload_path
    blocked = [
        "uploads/4.jpg", "static/uploads/4.jpg",
        "UPLOADS/4.jpg", "STATIC/UPLOADS/4.JPG", "Static/Uploads/4.jpg",
        "uploads\\4.jpg", "static\\uploads\\4.jpg",
        "uploads/../secret.txt", "UPLOADS/../4.jpg",
    ]
    allowed = ["js/app.js", "img/logo.png", "css/style.css", "login.html"]
    for p in blocked:
        assert _is_upload_path(p), f"照片路徑應封鎖卻放行: {p}"
    for p in allowed:
        assert not _is_upload_path(p), f"一般靜態資源應放行卻封鎖: {p}"


# ============================================================
# 3. 行為式：本次修復的輸入驗證/守衛（防回歸）
# ============================================================

def test_delete_stockout_rejects_non_outbound(client):
    """DELETE /api/stockouts/{id} 只允許刪出庫流水（手動調整/盤點等稽核軌跡不可刪）。"""
    r = client.post("/api/items", json={"name": "測試品", "stocks": [{"location": "A", "qty": 10}]})
    assert r.status_code in (200, 201), f"建品項失敗: {r.status_code}"
    item_id = r.json()["id"]
    r = client.post(f"/api/items/{item_id}/adjust", json={"delta": 5, "reason": "手動調整"})
    assert r.status_code == 200

    conn = sqlite3.connect(app_db.DB_PATH)
    try:
        mid = conn.execute("SELECT id FROM movements WHERE reason='手動調整' ORDER BY id DESC LIMIT 1").fetchone()[0]
    finally:
        conn.close()

    r = client.delete(f"/api/stockouts/{mid}")
    assert r.status_code == 400, f"手動調整流水應拒絕刪除，卻 {r.status_code}: {r.text}"


def test_kit_negative_qty_rejected(client):
    """整組材料 qty 必須 >0（負 qty 會讓組裝寫出 +delta 假流水）。"""
    r = client.post("/api/items", json={"name": "材料", "stocks": [{"location": "A", "qty": 10}]})
    item_id = r.json()["id"]
    r = client.post("/api/kits", json={"name": "測試組", "items": [{"item_id": item_id, "qty": -1}]})
    assert r.status_code == 400, f"負材料數量應 400，卻 {r.status_code}: {r.text}"


def test_kit_missing_item_rejected(client):
    """整組材料品項必須存在（不存在 → 400，不 500）。"""
    r = client.post("/api/kits", json={"name": "測試組", "items": [{"item_id": 99999, "qty": 1}]})
    assert r.status_code == 400, f"不存在材料品項應 400，卻 {r.status_code}: {r.text}"


def test_kit_non_dict_comp_rejected(client):
    """整組材料非 dict → 400（不 500）。"""
    r = client.post("/api/kits", json={"name": "測試組", "items": ["x"]})
    assert r.status_code == 400, f"非 dict 材料應 400，卻 {r.status_code}: {r.text}"


def test_import_negative_qty_rejected(client):
    """import 數量不得為負（負庫存入庫）。"""
    r = client.post("/api/import", json={"items": [{"name": "負數品", "qty": -5}]})
    assert r.status_code == 400, f"負數數量應 400，卻 {r.status_code}: {r.text}"


def test_stocktake_non_dict_rejected(client):
    """盤點 items 元素非 dict → 400（不 500）。"""
    r = client.post("/api/stocktake", json={"items": ["x"]})
    assert r.status_code == 400, f"非 dict 盤點項目應 400，卻 {r.status_code}: {r.text}"


def test_change_password_rate_limited(client):
    """改密碼端點套 per-IP rate limit：連續錯舊密碼超過上限 → 429。"""
    from app.services.auth import clear_ip_fail
    try:
        for i in range(10):
            r = client.put("/api/auth/password", json={"old_password": "WrongPass1", "new_password": "NewPass123"})
            assert r.status_code == 400, f"第 {i+1} 次錯誤舊密碼應 400，卻 {r.status_code}"
        r = client.put("/api/auth/password", json={"old_password": "WrongPass1", "new_password": "NewPass123"})
        assert r.status_code == 429, f"超過上限應 429，卻 {r.status_code}: {r.text}"
    finally:
        clear_ip_fail("testclient")


def test_appointment_time_format_validated(client):
    """行事曆時間必須 HH:MM（XSS payload 不可入庫）。"""
    r = client.post("/api/appointments", json={
        "client_name": "測試客戶", "date": "2026-08-12",
        "start_time": '"><img src=x onerror=alert(1)>', "end_time": "23:00",
    })
    assert r.status_code == 400, f"非法時間格式應 400，卻 {r.status_code}: {r.text}"


def test_appointment_year_out_of_range(client):
    """月曆 year 越界 → 400（datetime.date ValueError 500 已擋）。"""
    r = client.get("/api/appointments", params={"year": -5, "month": 1})
    assert r.status_code == 400, f"year=-5 應 400，卻 {r.status_code}: {r.text}"


def test_appointment_bad_service_type(client):
    """service_type_id 不存在 → 400（FK IntegrityError 500 已擋）。"""
    r = client.post("/api/appointments", json={
        "client_name": "測試客戶", "date": "2026-08-12",
        "start_time": "09:00", "end_time": "10:00", "service_type_id": 99999,
    })
    assert r.status_code == 400, f"不存在服務項目應 400，卻 {r.status_code}: {r.text}"


def test_user_color_hex_validated(client):
    """使用者 color 只接受 #RRGGBB（屬性逃逸 stored XSS 已擋）。"""
    r = client.put("/api/users/1", json={"color": 'red" onmouseover="alert(1)'})
    assert r.status_code == 400, f"非法 color 應 400，卻 {r.status_code}: {r.text}"


def test_daily_report_time_formula_guarded():
    """日報表時間欄也套 _safe()：= 開頭時間 → 輸出含撇號（公式注入已擋）。"""
    from app.services.report import build_daily_report
    buf, mmdd = build_daily_report("2026-08-12", [{
        "client_name": "客戶", "address": "", "service_type_id": 1, "service_name": "保養",
        "start_time": "=1+1", "end_time": "2:00", "note": "",
    }])
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    b4 = ws["B4"].value
    assert isinstance(b4, str) and b4.startswith("'"), f"時間欄應被 _safe 防護，卻拿到裸值: {b4!r}"
def test_preview_csp_is_frame_ancestors_only_for_pdf_viewer(client):
    """preview 的 CSP 不可含 default-src（Chrome PDF 閱讀器走 chrome-extension:// 會被擋）。

    回歸：手機 iframe 開 PDF 顯示「這項內容已遭到封鎖」。
    framing 保護保留（SAMEORIGIN + frame-ancestors 'self'）；非 preview 維持全站嚴格 CSP。
    """
    up = client.post(
        "/api/signed-reports",
        data={"report_date": "2026-09-14", "uploader_name": "測試", "note": ""},
        files={"file": ("t.pdf", b"%PDF-1.4 test", "application/pdf")},
    )
    assert up.status_code == 200, up.text
    rid = up.json()["id"]

    preview = client.get(f"/api/signed-reports/{rid}/preview")
    assert preview.status_code == 200
    csp = preview.headers["content-security-policy"]
    assert "frame-ancestors 'self'" in csp
    assert "default-src" not in csp
    assert preview.headers["x-frame-options"] == "SAMEORIGIN"
    assert preview.headers["cache-control"] == "no-store"

    other = client.get("/api/auth/me")
    assert other.status_code == 200
    other_csp = other.headers["content-security-policy"]
    assert "default-src 'self'" in other_csp
    assert "frame-ancestors 'none'" in other_csp
