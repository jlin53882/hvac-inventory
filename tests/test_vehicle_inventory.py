# -*- coding: utf-8 -*-
"""廂型車／貨車獨立庫存區：四 site 契約與整組同區測試。"""
import os

import pytest
from fastapi.testclient import TestClient

import app.config as app_config
import app.database as app_db
import main as app_main
from app.services.auth import SESSION_COOKIE, create_session, init_admin_if_missing


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """每個測試使用獨立 SQLite，避免碰正式 inventory.db。"""
    test_db = tmp_path / "test_vehicle_inventory.db"
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()
    monkeypatch.setattr(app_db, "DB_PATH", str(test_db))
    monkeypatch.setattr(app_config, "UPLOAD_DIR", str(upload_dir))
    app_db.init_db()
    conn = app_db.get_db()
    try:
        init_admin_if_missing(conn)
        admin_id = conn.execute("SELECT id FROM users WHERE username='admin'").fetchone()["id"]
        token = create_session(conn, admin_id)
    finally:
        conn.close()
    with TestClient(app_main.app) as test_client:
        test_client.cookies.set(SESSION_COOKIE, token)
        yield test_client
    try:
        if test_db.exists():
            test_db.unlink()
    except PermissionError:
        pass


def add_item(client, *, site="office", code="R410A", name="冷媒 R410A", qty=10,
             location="測試位置", low_stock=0):
    response = client.post("/api/items", json={
        "brand": "大金",
        "code": code,
        "name": name,
        "unit": "個",
        "low_stock": low_stock,
        "site": site,
        "stocks": [{"location": location, "qty": qty, "note": ""}],
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_export_contains_four_inventory_site_sheets(client):
    for site in ("office", "warehouse", "van", "truck"):
        add_item(client, site=site, code=f"EXPORT-{site}", qty=1, location="車內" if site in ("van", "truck") else "櫃位")
    response = client.get("/api/export?days=30")
    assert response.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert {"辦公室", "倉庫", "廂型車", "貨車"} <= set(workbook.sheetnames)


def test_vehicle_sites_accept_duplicate_item_identity_per_site(client):
    """相同物料可在 van/truck 各有獨立 item row 與數量。"""
    van = add_item(client, site="van", qty=2, location="車內")
    truck = add_item(client, site="truck", qty=3, location="車內")

    assert van["site"] == "van"
    assert truck["site"] == "truck"
    assert van["id"] != truck["id"]
    assert client.get("/api/items", params={"site": "van"}).json()[0]["total_qty"] == 2
    assert client.get("/api/items", params={"site": "truck"}).json()[0]["total_qty"] == 3


def test_vehicle_site_rejects_unknown_value(client):
    response = client.post("/api/items", json={
        "brand": "大金", "code": "BAD-SITE", "name": "非法區域",
        "site": "factory", "stocks": [{"location": "A", "qty": 1}],
    })
    assert response.status_code == 422
    assert client.get("/api/items", params={"site": "factory"}).status_code == 422


def test_stats_summary_exposes_four_inventory_sites(client):
    add_item(client, site="office", code="OFFICE", qty=1)
    add_item(client, site="warehouse", code="WAREHOUSE", qty=10)
    add_item(client, site="van", code="VAN", qty=2, location="車內")
    add_item(client, site="truck", code="TRUCK", qty=3, location="車內")

    response = client.get("/api/stats/summary")
    assert response.status_code == 200
    body = response.json()
    assert set(("all", "office", "warehouse", "van", "truck")) <= body.keys()
    assert body["van"]["total_qty"] == 2
    assert body["truck"]["total_qty"] == 3
    assert body["all"]["total_qty"] == 16


def test_kit_uses_selected_site_and_rejects_cross_site_component(client):
    van_component = add_item(client, site="van", code="VAN-MAT", name="廂型車材料", qty=4, location="車內")
    warehouse_component = add_item(client, site="warehouse", code="WH-MAT", name="倉庫材料", qty=4, location="鐵架")

    created = client.post("/api/kits", json={
        "name": "廂型車工具組",
        "site": "van",
        "items": [{"item_id": van_component["id"], "qty": 1}],
    })
    assert created.status_code == 201, created.text
    kits = client.get("/api/kits", params={"site": "van"})
    assert kits.status_code == 200
    assert kits.json()[0]["site"] == "van"

    cross_site = client.post("/api/kits", json={
        "name": "錯誤跨區整組",
        "site": "van",
        "items": [{"item_id": warehouse_component["id"], "qty": 1}],
    })
    assert cross_site.status_code == 400


def test_stocktake_site_isolation(client):
    van = add_item(client, site="van", code="TAKE-VAN", qty=2, location="車內")
    truck = add_item(client, site="truck", code="TAKE-TRUCK", qty=4, location="車內")
    cross = client.post("/api/stocktake", json={
        "site": "van",
        "items": [{"item_id": truck["id"], "location": "車內", "actual_qty": 3}],
    })
    assert cross.status_code == 400
    ok = client.post("/api/stocktake", json={
        "site": "van",
        "items": [{"item_id": van["id"], "location": "車內", "actual_qty": 1}],
    })
    assert ok.status_code == 200
    rows = client.get("/api/stocktakes", params={"site": "van"}).json()
    assert len(rows) == 1
    assert rows[0]["item_id"] == van["id"]


def test_inventory_transfer_creates_target_site_item_and_moves_stock(client):
    source = add_item(client, site="office", code="TRANSFER", qty=5, location="辦公室櫃A")
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"],
        "target_site": "van",
        "qty": 2,
        "source_location": "辦公室櫃A",
        "target_location": "車內工具櫃｜左側",
    })
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["source_item_id"] == source["id"]
    assert body["target_site"] == "van"
    assert body["qty"] == 2

    office = client.get("/api/items", params={"site": "office"}).json()
    van = client.get("/api/items", params={"site": "van"}).json()
    assert office[0]["total_qty"] == 3
    assert van[0]["total_qty"] == 2
    assert van[0]["stocks"][0]["location"] == "車內工具櫃｜左側"


def test_inventory_transfer_rejects_same_site_and_insufficient_stock(client):
    source = add_item(client, site="truck", code="TRANSFER-ERR", qty=1, location="車內")
    same_site = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "truck", "qty": 1,
    })
    assert same_site.status_code == 400
    insufficient = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 2,
    })
    assert insufficient.status_code == 400


def test_movements_can_be_filtered_by_inventory_site(client):
    source = add_item(client, site="office", code="MOV-SITE", qty=2, location="辦公室")
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 1,
    })
    assert response.status_code == 201
    van_movements = client.get("/api/movements", params={"site": "van"}).json()
    office_movements = client.get("/api/movements", params={"site": "office"}).json()
    assert len(van_movements) == 1
    assert van_movements[0]["site"] == "van"
    assert len(office_movements) == 1
    assert all(row["site"] == "office" for row in office_movements)


def test_assembled_kit_transfer_copies_definition_to_target_site(client):
    component = add_item(client, site="van", code="KIT-MAT", name="車用材料", qty=2, location="車內")
    created = client.post("/api/kits", json={
        "name": "車用工具組", "site": "van",
        "items": [{"item_id": component["id"], "qty": 1}],
    })
    assert created.status_code == 201, created.text
    kit = created.json()
    assembled = client.post(f"/api/kits/{kit['id']}/assemble", json={"qty": 1})
    assert assembled.status_code == 200, assembled.text

    transferred = client.post("/api/inventory/transfers", json={
        "item_id": kit["item_id"], "target_site": "truck", "qty": 1,
    })
    assert transferred.status_code == 201, transferred.text
    target_kits = client.get("/api/kits", params={"site": "truck"}).json()
    assert len(target_kits) == 1
    assert target_kits[0]["name"] == "車用工具組"
    assert target_kits[0]["stock_qty"] == 1
    assert target_kits[0]["components"][0]["name"] == "車用材料"


def test_item_patch_rejects_cross_site_change(client):
    item = add_item(client, site="office", code="PATCH-SITE", qty=3)
    response = client.patch(f"/api/items/{item['id']}", json={"site": "van"})
    assert response.status_code == 400
    assert client.get("/api/items", params={"site": "office"}).json()[0]["id"] == item["id"]
    assert client.get("/api/items", params={"site": "van"}).json() == []


def test_batch_location_rejects_cross_site_move(client):
    item = add_item(client, site="office", code="BATCH-SITE", qty=2, location="A")
    stock_id = item["stocks"][0]["id"]
    response = client.post("/api/stocks/batch-location", json={
        "stock_ids": [stock_id], "new_location": "B", "new_site": "van",
    })
    assert response.status_code == 400
    current = client.get("/api/items", params={"site": "office"}).json()[0]
    assert current["site"] == "office"
    assert current["stocks"][0]["location"] == "A"


def test_batch_location_does_not_move_unselected_stock_to_other_site(client):
    item = add_item(client, site="office", code="BATCH-SELECT", qty=2, location="A")
    extra_response = client.post(f"/api/items/{item['id']}/stocks", json={"location": "B", "qty": 3})
    assert extra_response.status_code == 201
    response = client.post("/api/stocks/batch-location", json={
        "stock_ids": [item["stocks"][0]["id"]], "new_location": "C", "new_site": "van",
    })
    assert response.status_code == 400
    current = client.get("/api/items", params={"site": "office"}).json()[0]
    assert {s["location"] for s in current["stocks"]} == {"A", "B"}


def test_transfer_rejects_when_qty_would_drop_below_prepared(client):
    source = add_item(client, site="office", code="PREP-TRANSFER", qty=10, location="A")
    prepared = client.post(f"/api/items/{source['id']}/prepare", json={"qty": 8})
    assert prepared.status_code == 200, prepared.text
    allowed = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 2, "source_location": "A",
    })
    assert allowed.status_code == 201, allowed.text
    rejected = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "truck", "qty": 1, "source_location": "A",
    })
    assert rejected.status_code == 400
    office = client.get("/api/items", params={"site": "office"}).json()[0]
    assert office["total_qty"] == 8
    assert office["prepared_qty"] == 8


def _create_kit(client, *, site, component, name):
    response = client.post("/api/kits", json={
        "name": name, "site": site, "items": [{"item_id": component["id"], "qty": 1}],
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_transfer_existing_normal_item_can_reuse_target(client):
    source = add_item(client, site="van", code="REUSE-NORMAL", qty=2, location="車內")
    target = add_item(client, site="truck", code="REUSE-NORMAL", qty=0, location="車內")
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "truck", "qty": 1,
    })
    assert response.status_code == 201, response.text
    assert response.json()["target_item_id"] == target["id"]
    target_item = next(i for i in client.get("/api/items", params={"site": "truck"}).json()
                       if i["id"] == target["id"])
    assert target_item["total_qty"] == 1


def test_transfer_rejects_target_when_kit_flag_mismatch(client):
    component = add_item(client, site="van", code="FLAG-COMP", name="組件", qty=1, location="車內")
    source_kit = _create_kit(client, site="van", component=component, name="旗標衝突組")
    assembled = client.post(f"/api/kits/{source_kit['id']}/assemble", json={"qty": 1})
    assert assembled.status_code == 200
    normal_target = client.post("/api/items", json={
        "brand": "", "code": "", "name": "旗標衝突組", "unit": "組", "site": "truck",
        "stocks": [{"location": "車內", "qty": 0}],
    })
    assert normal_target.status_code == 201
    rejected = client.post("/api/inventory/transfers", json={
        "item_id": source_kit["item_id"], "target_site": "truck", "qty": 1,
    })
    assert rejected.status_code == 409
    source_item = client.get("/api/items", params={"site": "van"}).json()
    assert source_item[0]["total_qty"] == 1
    assert client.get("/api/items", params={"site": "truck"}).json()[0]["total_qty"] == 0


def test_transfer_existing_identical_kit_reuses_target(client):
    van_component = add_item(client, site="van", code="IDENTICAL-COMP", name="同材料", qty=2, location="車內")
    truck_component = add_item(client, site="truck", code="IDENTICAL-COMP", name="同材料", qty=0, location="車內")
    source_kit = _create_kit(client, site="van", component=van_component, name="相同 BOM 組")
    target_kit = _create_kit(client, site="truck", component=truck_component, name="相同 BOM 組")
    assert client.post(f"/api/kits/{source_kit['id']}/assemble", json={"qty": 1}).status_code == 200
    transferred = client.post("/api/inventory/transfers", json={
        "item_id": source_kit["item_id"], "target_site": "truck", "qty": 1,
    })
    assert transferred.status_code == 201, transferred.text
    assert transferred.json()["target_item_id"] == target_kit["item_id"]
    target_item = next(i for i in client.get("/api/items", params={"site": "truck"}).json()
                       if i["id"] == target_kit["item_id"])
    assert target_item["total_qty"] == 1


def test_transfer_rejects_existing_kit_with_different_bom(client):
    source_component = add_item(client, site="van", code="BOM-A", name="材料 A", qty=2, location="車內")
    target_component = add_item(client, site="truck", code="BOM-B", name="材料 B", qty=2, location="車內")
    source_kit = _create_kit(client, site="van", component=source_component, name="不同 BOM 組")
    _create_kit(client, site="truck", component=target_component, name="不同 BOM 組")
    assert client.post(f"/api/kits/{source_kit['id']}/assemble", json={"qty": 1}).status_code == 200
    before = client.get("/api/movements", params={"site": "van"}).json()
    rejected = client.post("/api/inventory/transfers", json={
        "item_id": source_kit["item_id"], "target_site": "truck", "qty": 1,
    })
    assert rejected.status_code == 409
    source_items = client.get("/api/items", params={"site": "van"}).json()
    assert any(i["id"] == source_kit["item_id"] and i["total_qty"] == 1 for i in source_items)
    assert len(client.get("/api/movements", params={"site": "van"}).json()) == len(before)


def test_transfer_from_empty_location_only_deducts_empty_location(client):
    source = add_item(client, site="office", code="EMPTY-LOCATION", qty=2, location="")
    client.post(f"/api/items/{source['id']}/stocks", json={"location": "A", "qty": 5})
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 2, "source_location": "",
    })
    assert response.status_code == 201, response.text
    current = client.get("/api/items", params={"site": "office"}).json()[0]
    assert {s["location"]: s["qty"] for s in current["stocks"]} == {"": 0, "A": 5}


def test_transfer_with_source_location_none_can_deduct_across_locations(client):
    source = add_item(client, site="office", code="ALL-LOCATIONS", qty=2, location="A")
    client.post(f"/api/items/{source['id']}/stocks", json={"location": "B", "qty": 3})
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 4, "source_location": None,
    })
    assert response.status_code == 201, response.text
    current = client.get("/api/items", params={"site": "office"}).json()[0]
    assert current["total_qty"] == 1


def test_import_rejects_unknown_site(client):
    response = client.post("/api/import", json={"items": [{
        "brand": "大金", "code": "BAD-IMPORT", "name": "未知分片", "site": "factory", "qty": 1,
    }]})
    assert response.status_code == 400


def test_import_accepts_van_and_truck(client):
    response = client.post("/api/import", json={"items": [
        {"brand": "大金", "code": "IMPORT-VAN", "name": "車料一", "site": "van", "qty": 1},
        {"brand": "大金", "code": "IMPORT-TRUCK", "name": "車料二", "site": "truck", "qty": 1},
    ]})
    assert response.status_code == 200, response.text
    assert client.get("/api/items", params={"site": "van"}).json()[0]["site"] == "van"
    assert client.get("/api/items", params={"site": "truck"}).json()[0]["site"] == "truck"


def test_transfer_rejects_qty_that_canonicalizes_to_zero(client):
    source = add_item(client, site="office", code="ZERO-CANON", qty=1, location="A")
    response = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 0.0004,
    })
    assert response.status_code == 400
    assert client.get("/api/items", params={"site": "van"}).json() == []
    assert client.get("/api/items", params={"site": "office"}).json()[0]["total_qty"] == 1


def test_transfer_fraction_boundary_equal_prepared_is_allowed(client):
    source = add_item(client, site="office", code="TRANSFER-FRACTION-PREPARED", qty=0.3, location="A")
    prepared = client.post(f"/api/items/{source['id']}/prepare", json={"qty": 0.2})
    assert prepared.status_code == 200, prepared.text
    transferred = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 0.1, "source_location": "A",
    })
    assert transferred.status_code == 201, transferred.text
    office = next(i for i in client.get("/api/items", params={"site": "office"}).json()
                  if i["id"] == source["id"])
    assert office["total_qty"] == 0.2
    assert office["prepared_qty"] == 0.2


def test_transfer_fraction_boundary_rejects_below_prepared(client):
    source = add_item(client, site="office", code="TRANSFER-FRACTION-REJECT", qty=0.3, location="A")
    prepared = client.post(f"/api/items/{source['id']}/prepare", json={"qty": 0.2})
    assert prepared.status_code == 200, prepared.text
    rejected = client.post("/api/inventory/transfers", json={
        "item_id": source["id"], "target_site": "van", "qty": 0.101, "source_location": "A",
    })
    assert rejected.status_code == 400
    office = next(i for i in client.get("/api/items", params={"site": "office"}).json()
                  if i["id"] == source["id"])
    assert office["total_qty"] == 0.3
    assert office["prepared_qty"] == 0.2


def test_soft_deleted_inventory_stockout_stays_in_original_site(client):
    item = add_item(client, site="office", code="DELETED-SITE", qty=3, location="A")
    out = client.post("/api/stockout", json={
        "item_id": item["id"], "qty": 1, "destination": "測試工地",
    })
    assert out.status_code == 200, out.text
    deleted = client.delete(f"/api/items/{item['id']}")
    assert deleted.status_code == 200, deleted.text
    office = client.get("/api/stockouts", params={"site": "office"}).json()
    van = client.get("/api/stockouts", params={"site": "van"}).json()
    truck = client.get("/api/stockouts", params={"site": "truck"}).json()
    assert any(row["item_id"] == item["id"] for row in office)
    assert all(row["item_id"] != item["id"] for row in van)
    assert all(row["item_id"] != item["id"] for row in truck)


def test_nonstock_stockout_visible_in_all_inventory_sites(client):
    response = client.post("/api/stockout/nonstock", json={
        "name": "臨時耗材四區", "qty": 1, "destination": "某案場",
    })
    assert response.status_code == 200, response.text
    item_id = response.json()["id"]
    conn = app_db.get_db()
    try:
        row = conn.execute("SELECT is_deleted, site FROM items WHERE id=?", (item_id,)).fetchone()
        assert row["is_deleted"] == 1
        assert row["site"] == ""
    finally:
        conn.close()
    for site in ("office", "warehouse", "van", "truck"):
        rows = client.get("/api/stockouts", params={"site": site}).json()
        assert any(row["item_id"] == item_id for row in rows), site
