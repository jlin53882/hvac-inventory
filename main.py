# -*- coding: utf-8 -*-
"""
冷凍空調庫存管理系統 - 後端 API
================================
技術：FastAPI + SQLite（單檔資料庫，免安裝）
功能：
  - GET  /api/items          查詢品項（支援 brand / search / location 篩選）
  - POST /api/items          新增品項
  - PATCH /api/items/{id}    修改品項資料
  - POST /api/items/{id}/adjust  加減庫存數量（delta 正負）
  - POST /api/import         從 JSON 匯入（首次建立資料用）
  - GET  /api/export         匯出 Excel (.xlsx)
  - GET  /api/movements      查詢庫存異動紀錄
  - GET  /                   前端網頁
  - GET  /health             健康檢查

啟動：.venv/Scripts/python.exe -m uvicorn main:app --host 0.0.0.0 --port 8000
"""
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import Optional
import sqlite3, json, os, io, datetime, shutil

# ---------- 路徑設定 ----------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "inventory.db")
STATIC_DIR = os.path.join(BASE_DIR, "static")
EXPORT_DIR = os.path.join(BASE_DIR, "exports")

os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)

app = FastAPI(title="冷凍空調庫存系統", version="1.0.0")


# ---------- 資料庫 ----------
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS items (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        brand       TEXT NOT NULL DEFAULT '',
        code        TEXT DEFAULT '',
        name        TEXT NOT NULL,
        qty         REAL NOT NULL DEFAULT 0,
        prepared_qty REAL NOT NULL DEFAULT 0,
        unit        TEXT NOT NULL DEFAULT '個',
        location    TEXT DEFAULT '',
        note        TEXT DEFAULT '',
        low_stock   REAL DEFAULT 0,
        is_kit      INTEGER NOT NULL DEFAULT 0,
        site        TEXT NOT NULL DEFAULT 'office',
        created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS movements (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        item_id     INTEGER NOT NULL REFERENCES items(id),
        delta       REAL NOT NULL,
        before_qty  REAL NOT NULL DEFAULT 0,
        after_qty   REAL NOT NULL DEFAULT 0,
        reason      TEXT DEFAULT '',
        destination TEXT DEFAULT '',
        created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS stocktakes (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        take_date   TEXT NOT NULL,
        item_id     INTEGER NOT NULL REFERENCES items(id),
        system_qty  REAL NOT NULL DEFAULT 0,
        actual_qty  REAL NOT NULL DEFAULT 0,
        diff        REAL NOT NULL DEFAULT 0,
        note        TEXT DEFAULT '',
        created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS kits (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        item_id     INTEGER NOT NULL REFERENCES items(id),
        name        TEXT NOT NULL,
        note        TEXT DEFAULT '',
        created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS kit_items (
        id          INTEGER PRIMARY KEY AUTOINCREMENT,
        kit_id      INTEGER NOT NULL REFERENCES kits(id),
        item_id     INTEGER NOT NULL REFERENCES items(id),
        qty         REAL NOT NULL DEFAULT 1,
        created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    CREATE INDEX IF NOT EXISTS idx_items_brand ON items(brand);
    CREATE INDEX IF NOT EXISTS idx_items_location ON items(location);
    CREATE INDEX IF NOT EXISTS idx_movements_item ON movements(item_id);
    """)

    # 舊資料庫遷移：補上新增欄位
    item_cols = [r[1] for r in conn.execute("PRAGMA table_info(items)").fetchall()]
    if "prepared_qty" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN prepared_qty REAL NOT NULL DEFAULT 0")
        print("[migrate] items.prepared_qty 欄位已新增")
    if "is_kit" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN is_kit INTEGER NOT NULL DEFAULT 0")
        print("[migrate] items.is_kit 欄位已新增")
    if "site" not in item_cols:
        conn.execute("ALTER TABLE items ADD COLUMN site TEXT NOT NULL DEFAULT 'office'")
        print("[migrate] items.site 欄位已新增")
    mov_cols = [r[1] for r in conn.execute("PRAGMA table_info(movements)").fetchall()]
    if "destination" not in mov_cols:
        conn.execute("ALTER TABLE movements ADD COLUMN destination TEXT DEFAULT ''")
        print("[migrate] movements.destination 欄位已新增")
    conn.commit()
    conn.close()


init_db()


# ---------- Pydantic 模型 ----------
class ItemCreate(BaseModel):
    brand: str = ""
    code: str = ""
    name: str
    qty: float = 0
    unit: str = "個"
    location: str = ""
    note: str = ""
    low_stock: float = 0
    site: str = "office"  # office=辦公室 / warehouse=倉庫


class ItemUpdate(BaseModel):
    brand: Optional[str] = None
    code: Optional[str] = None
    name: Optional[str] = None
    unit: Optional[str] = None
    location: Optional[str] = None
    note: Optional[str] = None
    low_stock: Optional[float] = None
    site: Optional[str] = None


class AdjustRequest(BaseModel):
    delta: float
    reason: str = ""
    destination: str = ""  # 出庫去向（客戶/案場/工地）


# ---------- API ----------
@app.get("/health")
def health():
    return {"status": "ok", "time": datetime.datetime.now().isoformat()}


@app.get("/api/items")
def list_items(
    brand: Optional[str] = None,
    search: Optional[str] = None,
    location: Optional[str] = None,
    sort: str = "brand",
    site: Optional[str] = None,
):
    conn = get_db()
    sql = "SELECT * FROM items WHERE 1=1"
    params = []
    if site and site != "all":
        sql += " AND site = ?"
        params.append(site)
    if brand and brand != "全部":
        sql += " AND brand = ?"
        params.append(brand)
    if location:
        sql += " AND location = ?"
        params.append(location)
    if search:
        sql += " AND (name LIKE ? OR code LIKE ? OR brand LIKE ? OR note LIKE ?)"
        like = f"%{search}%"
        params += [like, like, like, like]

    sort_map = {
        "brand": "brand COLLATE NOCASE, name",
        "location": "location, name",
        "qty": "qty DESC",
        "created": "id DESC",
    }
    sql += f" ORDER BY {sort_map.get(sort, 'brand COLLATE NOCASE, name')}"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.post("/api/items", status_code=201)
def create_item(item: ItemCreate):
    conn = get_db()
    cur = conn.execute(
        "INSERT INTO items (brand, code, name, qty, unit, location, note, low_stock, site) VALUES (?,?,?,?,?,?,?,?,?)",
        (item.brand, item.code, item.name, item.qty, item.unit, item.location, item.note, item.low_stock, item.site),
    )
    conn.commit()
    row = conn.execute("SELECT * FROM items WHERE id=?", (cur.lastrowid,)).fetchone()
    conn.close()
    return dict(row)


@app.patch("/api/items/{item_id}")
def update_item(item_id: int, upd: ItemUpdate):
    conn = get_db()
    fields = {k: v for k, v in upd.model_dump().items() if v is not None}
    if not fields:
        raise HTTPException(400, "沒有要更新的欄位")
    fields["updated_at"] = datetime.datetime.now().isoformat()
    sets = ", ".join(f"{k}=?" for k in fields)
    conn.execute(f"UPDATE items SET {sets} WHERE id=?", (*fields.values(), item_id))
    conn.commit()
    row = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "品項不存在")
    return dict(row)


@app.post("/api/items/{item_id}/adjust")
def adjust_qty(item_id: int, req: AdjustRequest):
    conn = get_db()
    row = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "品項不存在")

    before = row["qty"]
    after = max(0, before + req.delta)

    conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                 (after, datetime.datetime.now().isoformat(), item_id))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (item_id, req.delta, before, after, req.reason, req.destination),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    return dict(updated)


@app.post("/api/import")
def import_items(data: dict):
    """data: {"items": [...], "mode": "replace"|"append"}"""
    items = data.get("items", [])
    mode = data.get("mode", "replace")
    conn = get_db()

    if mode == "replace":
        conn.execute("DELETE FROM movements")
        conn.execute("DELETE FROM items")

    count = 0
    for it in items:
        cur = conn.execute(
            "INSERT INTO items (brand, code, name, qty, unit, location, note, low_stock) VALUES (?,?,?,?,?,?,?,?)",
            (
                it.get("brand", ""),
                it.get("code", ""),
                it.get("name", ""),
                it.get("qty", 0),
                it.get("unit", "個"),
                it.get("location", ""),
                it.get("note", ""),
                it.get("low_stock", 0),
            ),
        )
        # 同步建立異動紀錄（初始庫存）
        conn.execute(
            "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
            (cur.lastrowid, it.get("qty", 0), 0, it.get("qty", 0), "初始匯入", ""),
        )
        count += 1

    conn.commit()
    total = conn.execute("SELECT COUNT(*) FROM items").fetchone()[0]
    conn.close()
    return {"imported": count, "total_items": total}


@app.get("/api/export")
def export_excel():
    """匯出完整庫存 Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(500, "缺少 openpyxl，請先安裝")

    conn = get_db()
    rows = conn.execute("SELECT * FROM items ORDER BY brand COLLATE NOCASE, name").fetchall()
    conn.close()

    wb = openpyxl.Workbook()

    # Sheet 1: 庫存明細
    ws = wb.active
    ws.title = "庫存明細"
    headers = ["編號", "廠牌", "品項名稱", "數量", "單位", "位置", "備註"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="2E5C8A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r["id"], r["brand"], r["name"], r["qty"], r["unit"], r["location"], r["note"]])
    widths = [8, 14, 40, 10, 8, 30, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Sheet 2: 異動紀錄
    ws2 = wb.create_sheet("異動紀錄")
    ws2.append(["時間", "品項", "變動", "原本", "現在", "去向", "原因"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    conn = get_db()
    movs = conn.execute("""
        SELECT m.created_at, i.name, m.delta, m.before_qty, m.after_qty, m.destination, m.reason
        FROM movements m JOIN items i ON i.id = m.item_id
        ORDER BY m.id DESC LIMIT 500
    """).fetchall()
    conn.close()
    for m in movs:
        ws2.append(list(m))

    # Sheet 3: 廠牌統計
    ws3 = wb.create_sheet("廠牌統計")
    ws3.append(["廠牌", "品項數", "總庫存"])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    conn = get_db()
    stats = conn.execute("SELECT brand, COUNT(*), SUM(qty) FROM items GROUP BY brand ORDER BY COUNT(*) DESC").fetchall()
    conn.close()
    for s in stats:
        ws3.append(list(s))

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(EXPORT_DIR, f"庫存報表_{ts}.xlsx")
    wb.save(out_path)

    return FileResponse(out_path, filename=f"庫存報表_{ts}.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/movements")
def list_movements(limit: int = 100):
    conn = get_db()
    rows = conn.execute("""
        SELECT m.*, i.name as item_name, i.brand
        FROM movements m JOIN items i ON i.id = m.item_id
        ORDER BY m.id DESC LIMIT ?
    """, (limit,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.get("/api/stats")
def stats(site: Optional[str] = None):
    conn = get_db()
    where = ""
    params = ()
    if site and site != "all":
        where = " WHERE site = ?"
        params = (site,)
    total = conn.execute(f"SELECT COUNT(*) FROM items{where}", params).fetchone()[0]
    total_qty = conn.execute(f"SELECT COALESCE(SUM(qty),0) FROM items{where}", params).fetchone()[0]
    # 缺貨只統計單一材料；低庫存整組與單一都算
    zero_where = where + (" AND " if where else " WHERE ") + "is_kit = 0 AND qty <= 0"
    low_where = where + (" AND " if where else " WHERE ") + "qty <= low_stock AND low_stock > 0"
    low = conn.execute(f"SELECT COUNT(*) FROM items{low_where}", params).fetchone()[0]
    zero = conn.execute(f"SELECT COUNT(*) FROM items{zero_where}", params).fetchone()[0]
    brands = conn.execute(f"SELECT COUNT(DISTINCT brand) FROM items{where}", params).fetchone()[0]
    conn.close()
    return {"total_items": total, "total_qty": total_qty, "low_stock": low, "zero_stock": zero, "brands": brands}


# ---------- 出庫紀錄（帶去向） ----------
class StockOutRequest(BaseModel):
    item_id: int
    qty: float
    destination: str = ""
    note: str = ""


@app.post("/api/stockout")
def stock_out(req: StockOutRequest):
    """出庫：扣庫存 + 記錄去向（客戶/案場/工地）"""
    if req.qty <= 0:
        raise HTTPException(400, "出庫數量必須大於 0")
    conn = get_db()
    row = conn.execute("SELECT * FROM items WHERE id=?", (req.item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "品項不存在")
    if row["qty"] < req.qty:
        raise HTTPException(400, f"庫存不足！目前只剩 {row['qty']} {row['unit']}")

    before = row["qty"]
    after = before - req.qty
    reason = "出庫"
    if req.note:
        reason = f"出庫 - {req.note}"

    conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                 (after, datetime.datetime.now().isoformat(), req.item_id))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (req.item_id, -req.qty, before, after, reason, req.destination),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM items WHERE id=?", (req.item_id,)).fetchone()
    conn.close()
    return dict(updated)


@app.get("/api/stockouts")
def list_stock_outs(limit: int = 100, search: str = "", site: Optional[str] = None):
    """出庫紀錄（含去向）"""
    conn = get_db()
    sql = """
        SELECT m.*, i.name as item_name, i.brand, i.code, i.unit
        FROM movements m JOIN items i ON i.id = m.item_id
        WHERE m.delta < 0
    """
    params = []
    if site and site != "all":
        sql += " AND i.site = ?"
        params.append(site)
    if search:
        sql += " AND (m.destination LIKE ? OR i.name LIKE ? OR i.brand LIKE ?)"
        like = f"%{search}%"
        params += [like, like, like]
    sql += " ORDER BY m.id DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------- 領出準備（兩階段出庫） ----------
class PrepareRequest(BaseModel):
    qty: float
    note: str = ""


@app.post("/api/items/{item_id}/prepare")
def prepare_item(item_id: int, req: PrepareRequest):
    """領出準備：把東西拿出來準備（庫存不扣，只標記 prepared_qty）"""
    if req.qty <= 0:
        raise HTTPException(400, "數量必須大於 0")
    conn = get_db()
    row = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "品項不存在")
    available = row["qty"] - row["prepared_qty"]
    if req.qty > available:
        raise HTTPException(400, f"可領出數量不足！可用 {available} {row['unit']}")

    new_prepared = row["prepared_qty"] + req.qty
    conn.execute("UPDATE items SET prepared_qty=?, updated_at=? WHERE id=?",
                 (new_prepared, datetime.datetime.now().isoformat(), item_id))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (item_id, 0, row["qty"], row["qty"], "領出準備", ""),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    return dict(updated)


@app.post("/api/items/{item_id}/prepared-out")
def prepared_out(item_id: int, req: PrepareRequest):
    """確認出庫：從準備中的數量真正出庫（此時才扣庫存）+ 記錄去向"""
    if req.qty <= 0:
        raise HTTPException(400, "數量必須大於 0")
    conn = get_db()
    row = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "品項不存在")
    if req.qty > row["prepared_qty"]:
        raise HTTPException(400, f"準備中的數量只有 {row['prepared_qty']} {row['unit']}")

    before = row["qty"]
    after = before - req.qty
    new_prepared = row["prepared_qty"] - req.qty
    dest = req.note  # note 欄位當去向用（相容前端）

    conn.execute("UPDATE items SET qty=?, prepared_qty=?, updated_at=? WHERE id=?",
                 (after, new_prepared, datetime.datetime.now().isoformat(), item_id))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (item_id, -req.qty, before, after, "出庫", dest),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    return dict(updated)


@app.post("/api/items/{item_id}/prepared-return")
def prepared_return(item_id: int, req: PrepareRequest):
    """退回：把準備中的數量退回（取消領出）"""
    if req.qty <= 0:
        raise HTTPException(400, "數量必須大於 0")
    conn = get_db()
    row = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    if not row:
        raise HTTPException(404, "品項不存在")
    if req.qty > row["prepared_qty"]:
        raise HTTPException(400, f"準備中的數量只有 {row['prepared_qty']} {row['unit']}")

    new_prepared = row["prepared_qty"] - req.qty
    conn.execute("UPDATE items SET prepared_qty=?, updated_at=? WHERE id=?",
                 (new_prepared, datetime.datetime.now().isoformat(), item_id))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (item_id, 0, row["qty"], row["qty"], "退回準備", ""),
    )
    conn.commit()
    updated = conn.execute("SELECT * FROM items WHERE id=?", (item_id,)).fetchone()
    conn.close()
    return dict(updated)


@app.get("/api/prepared")
def list_prepared(site: Optional[str] = None):
    """準備中清單（已領出尚未出庫）"""
    conn = get_db()
    where = ""
    params = ()
    if site and site != "all":
        where = " AND site = ?"
        params = (site,)
    rows = conn.execute(f"""
        SELECT * FROM items WHERE prepared_qty > 0{where} ORDER BY brand COLLATE NOCASE, name
    """, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------- 整組（套件） ----------
class KitCreate(BaseModel):
    name: str
    items: list  # [{item_id, qty}]
    note: str = ""


class KitAssemble(BaseModel):
    qty: float = 1


@app.get("/api/kits")
def list_kits(site: Optional[str] = None):
    """套件清單（含組成材料）"""
    conn = get_db()
    where = ""
    params = ()
    if site and site != "all":
        where = " WHERE i.site = ?"
        params = (site,)
    kits = conn.execute(f"SELECT k.*, i.qty as stock_qty, i.unit, i.brand FROM kits k JOIN items i ON i.id = k.item_id{where} ORDER BY k.name", params).fetchall()
    result = []
    for k in kits:
        items = conn.execute("""
            SELECT ki.item_id, ki.qty as need_qty, i.name, i.brand, i.code, i.unit, i.qty as stock
            FROM kit_items ki JOIN items i ON i.id = ki.item_id
            WHERE ki.kit_id = ?
        """, (k["id"],)).fetchall()
        d = dict(k)
        d["components"] = [dict(x) for x in items]
        result.append(d)
    conn.close()
    return result


@app.post("/api/kits", status_code=201)
def create_kit(kit: KitCreate):
    """新增套件定義：建立套件品項 + 組成材料"""
    if not kit.name or not kit.items:
        raise HTTPException(400, "套件名稱與材料都不能空白")
    conn = get_db()
    # 建立套件品項
    cur = conn.execute(
        "INSERT INTO items (brand, name, qty, unit, note, is_kit) VALUES (?,?,?,?,?,1)",
        ("", kit.name, 0, "組", kit.note),
    )
    kit_item_id = cur.lastrowid
    # 建立套件定義
    cur2 = conn.execute(
        "INSERT INTO kits (item_id, name, note) VALUES (?,?,?)",
        (kit_item_id, kit.name, kit.note),
    )
    kit_id = cur2.lastrowid
    for comp in kit.items:
        conn.execute(
            "INSERT INTO kit_items (kit_id, item_id, qty) VALUES (?,?,?)",
            (kit_id, comp["item_id"], comp.get("qty", 1)),
        )
    conn.commit()
    conn.close()
    return {"id": kit_id, "item_id": kit_item_id, "name": kit.name}


@app.post("/api/kits/{kit_id}/assemble")
def assemble_kit(kit_id: int, req: KitAssemble):
    """組裝：從材料庫存扣掉所需數量，整組庫存增加"""
    if req.qty <= 0:
        raise HTTPException(400, "組裝數量必須大於 0")
    conn = get_db()
    kit = conn.execute("SELECT * FROM kits WHERE id=?", (kit_id,)).fetchone()
    if not kit:
        raise HTTPException(404, "套件不存在")
    comps = conn.execute("SELECT * FROM kit_items WHERE kit_id=?", (kit_id,)).fetchall()

    # 檢查材料庫存
    short = []
    for c in comps:
        mat = conn.execute("SELECT * FROM items WHERE id=?", (c["item_id"],)).fetchone()
        need = c["qty"] * req.qty
        if mat["qty"] < need:
            short.append(f"{mat['name']}（需要 {need}，剩 {mat['qty']}）")
    if short:
        conn.close()
        raise HTTPException(400, "材料不足：" + "、".join(short))

    # 扣材料
    for c in comps:
        mat = conn.execute("SELECT * FROM items WHERE id=?", (c["item_id"],)).fetchone()
        need = c["qty"] * req.qty
        new_qty = mat["qty"] - need
        conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                     (new_qty, datetime.datetime.now().isoformat(), c["item_id"]))
        conn.execute(
            "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
            (c["item_id"], -need, mat["qty"], new_qty, f"組裝套件:{kit['name']}", ""),
        )
    # 加整組庫存
    kit_item = conn.execute("SELECT * FROM items WHERE id=?", (kit["item_id"],)).fetchone()
    new_kit_qty = kit_item["qty"] + req.qty
    conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                 (new_kit_qty, datetime.datetime.now().isoformat(), kit["item_id"]))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (kit["item_id"], req.qty, kit_item["qty"], new_kit_qty, f"組裝完成:{kit['name']}", ""),
    )
    conn.commit()
    conn.close()
    return {"ok": True, "kit": kit["name"], "qty": req.qty}


@app.post("/api/kits/{kit_id}/disassemble")
def disassemble_kit(kit_id: int, req: KitAssemble):
    """拆解：整組扣掉，材料庫存加回"""
    if req.qty <= 0:
        raise HTTPException(400, "拆解數量必須大於 0")
    conn = get_db()
    kit = conn.execute("SELECT * FROM kits WHERE id=?", (kit_id,)).fetchone()
    if not kit:
        raise HTTPException(404, "套件不存在")
    kit_item = conn.execute("SELECT * FROM items WHERE id=?", (kit["item_id"],)).fetchone()
    if kit_item["qty"] < req.qty:
        conn.close()
        raise HTTPException(400, f"整組庫存不足！只剩 {kit_item['qty']} 組")

    # 扣整組
    new_kit_qty = kit_item["qty"] - req.qty
    conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                 (new_kit_qty, datetime.datetime.now().isoformat(), kit["item_id"]))
    conn.execute(
        "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
        (kit["item_id"], -req.qty, kit_item["qty"], new_kit_qty, f"拆解:{kit['name']}", ""),
    )
    # 加回材料
    comps = conn.execute("SELECT * FROM kit_items WHERE kit_id=?", (kit_id,)).fetchall()
    for c in comps:
        mat = conn.execute("SELECT * FROM items WHERE id=?", (c["item_id"],)).fetchone()
        add = c["qty"] * req.qty
        new_qty = mat["qty"] + add
        conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                     (new_qty, datetime.datetime.now().isoformat(), c["item_id"]))
        conn.execute(
            "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
            (c["item_id"], add, mat["qty"], new_qty, f"拆解套件:{kit['name']}", ""),
        )
    conn.commit()
    conn.close()
    return {"ok": True, "kit": kit["name"], "qty": req.qty}


# ---------- 盤點 ----------
class StocktakeSubmit(BaseModel):
    take_date: str = ""  # 預設今天
    items: list  # [{item_id, actual_qty, note}]


@app.post("/api/stocktake")
def submit_stocktake(req: StocktakeSubmit):
    """盤點：逐項輸入實際數量，計算盤盈/盤虧並更新庫存"""
    conn = get_db()
    take_date = req.take_date or datetime.date.today().isoformat()
    results = []

    for it in req.items:
        row = conn.execute("SELECT * FROM items WHERE id=?", (it["item_id"],)).fetchone()
        if not row:
            continue
        system_qty = row["qty"]
        actual_qty = float(it.get("actual_qty", system_qty))
        diff = round(actual_qty - system_qty, 3)
        note = it.get("note", "")

        conn.execute("UPDATE items SET qty=?, updated_at=? WHERE id=?",
                     (actual_qty, datetime.datetime.now().isoformat(), row["id"]))
        conn.execute(
            "INSERT INTO movements (item_id, delta, before_qty, after_qty, reason, destination) VALUES (?,?,?,?,?,?)",
            (row["id"], diff, system_qty, actual_qty, "盤點調整", ""),
        )
        conn.execute(
            "INSERT INTO stocktakes (take_date, item_id, system_qty, actual_qty, diff, note) VALUES (?,?,?,?,?,?)",
            (take_date, row["id"], system_qty, actual_qty, diff, note),
        )
        results.append({"item_id": row["id"], "name": row["name"], "system_qty": system_qty,
                        "actual_qty": actual_qty, "diff": diff})

    conn.commit()
    conn.close()
    return {"take_date": take_date, "count": len(results), "results": results}


@app.get("/api/stocktakes")
def list_stocktakes(limit: int = 200):
    """盤點紀錄（含差異）"""
    conn = get_db()
    rows = conn.execute("""
        SELECT s.*, i.name as item_name, i.brand, i.unit
        FROM stocktakes s JOIN items i ON i.id = s.item_id
        ORDER BY s.id DESC LIMIT ?
    """, (limit,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


@app.get("/api/stocktake/dates")
def stocktake_dates():
    """盤點日期清單（含每批差異統計）"""
    conn = get_db()
    rows = conn.execute("""
        SELECT take_date, COUNT(*) as item_count,
               SUM(CASE WHEN diff != 0 THEN 1 ELSE 0 END) as diff_count,
               ROUND(SUM(diff), 3) as total_diff
        FROM stocktakes GROUP BY take_date ORDER BY take_date DESC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------- 靜態檔案（前端） ----------
@app.get("/")
def index():
    idx = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(idx):
        return FileResponse(idx)
    return Response("<h1>庫存系統 API</h1><p>前端尚未建立，請先將 index.html 放到 static/</p>", media_type="text/html")


# 掛載靜態目錄（放在最後，避免吃掉 API 路由）
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
