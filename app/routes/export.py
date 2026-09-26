# -*- coding: utf-8 -*-
"""Build selectable single-inventory Excel reports with safe formulas."""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from copy import copy
from typing import Iterable

from fastapi import APIRouter, Depends, HTTPException
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.services import movement_time
from app.services.auth import require_perm
from app.services.safety import excel_safe, xlsx_download

router = APIRouter()
SITES = {"office": "公司", "warehouse": "倉庫", "van": "廂型車", "truck": "貨車"}
SITE_ORDER = tuple(SITES)
MAX_RANGE_DAYS = 366
DEFAULT_EXPORT_SECTIONS = ("inventory", "positions", "movements")
EXPORT_SECTION_ORDER = ("overview", "inventory", "positions", "alerts", "movements", "stats")
HEADER_FILL = "2E5C8A"
TITLE_FILL = "163B63"
STATUS_FILLS = {"資料異常": "FCA5A5", "缺貨": "FECACA", "低庫存": "FED7AA"}


def _parse_export_range(month: str | None, start_date: str | None, end_date: str | None, now: dt.datetime | None = None) -> tuple[dt.datetime, dt.datetime, str, str]:
    """解析月份或自訂日期範圍，回傳 SQL 邊界與 Excel 顯示期間。"""
    now = now or dt.datetime.strptime(movement_time.now_sql(), movement_time.SQL_DATETIME_FORMAT)
    if start_date is not None or end_date is not None:
        if not start_date or not end_date:
            raise HTTPException(400, "start_date 與 end_date 必須同時提供")
        try:
            start_day = dt.date.fromisoformat(start_date)
            end_day = dt.date.fromisoformat(end_date)
        except ValueError:
            raise HTTPException(400, "日期格式必須為 YYYY-MM-DD")
        inclusive_days = (end_day - start_day).days + 1
        if start_day > end_day or inclusive_days > MAX_RANGE_DAYS:
            raise HTTPException(400, "日期範圍無效或超過 366 天")
        start = dt.datetime.combine(start_day, dt.time.min)
        end = dt.datetime.combine(end_day + dt.timedelta(days=1), dt.time.min)
        label = f"{start_day.isoformat()} ～ {end_day.isoformat()}"
        return start, end, label, f"{start_day.strftime('%Y/%m/%d')} ～ {end_day.strftime('%Y/%m/%d')}"
    if month is not None:
        if not re.fullmatch(r"\d{4}-\d{2}", month):
            raise HTTPException(400, "month 格式必須為 YYYY-MM")
        try:
            year, mon = map(int, month.split("-"))
            start_day = dt.date(year, mon, 1)
        except ValueError:
            raise HTTPException(400, "month 不是合法月份")
        next_month = dt.date(year + (mon == 12), 1 if mon == 12 else mon + 1, 1)
        start = dt.datetime.combine(start_day, dt.time.min)
        end = dt.datetime.combine(next_month, dt.time.min)
        if start_day.year == now.year and start_day.month == now.month:
            end = now + dt.timedelta(seconds=1)
        display_end = now.strftime("%Y/%m/%d %H:%M") if start_day.year == now.year and start_day.month == now.month else (next_month - dt.timedelta(days=1)).strftime("%Y/%m/%d")
        return start, end, f"{year} 年 {mon:02d} 月", f"{start_day.strftime('%Y/%m/%d')} ～ {display_end}"
    return _parse_export_range(now.strftime("%Y-%m"), None, None, now)


def _safe(value):
    """將資料庫文字交給既有 Excel 公式注入防護 helper。"""
    return excel_safe("" if value is None else value)


def _site_label(site: str) -> str:
    """將內部庫存區代碼轉成報表顯示名稱。"""
    return SITES.get(site, site)


def _period_text(label: str, display_period: str) -> str:
    """組合報表期間與異動統計期間的標題文字。"""
    return f"報表期間：{label}　異動統計：{display_period}"


def _style_title(ws, title: str, period: str):
    """Write the workbook title and reporting period without a snapshot stamp."""
    ws.merge_cells("A1:D1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws["A2"] = period
    ws["A2"].font = Font(color="475569", italic=True, size=10)


def _style_header(ws, row: int):
    """套用資料表表頭樣式並設定凍結窗格。"""
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{row + 1}"


def _style_data(ws, header_row: int, qty_columns: Iterable[int] = (), note_columns: Iterable[int] = ()):
    """套用資料列對齊、備註換行與基礎數量格式。"""
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(horizontal="right" if cell.column in qty_columns else "left", vertical="top", wrap_text=cell.column in note_columns)
        for col in qty_columns:
            row[col - 1].number_format = "#,##0.###"


def _display_width(value) -> int:
    """Estimate Excel width, counting full-width East Asian glyphs twice."""
    lines = str(value).splitlines() or [""]
    return max(sum(2 if unicodedata.east_asian_width(char) in ("F", "W") else 1 for char in line) for line in lines)


def _autofit_columns(ws, body_only_columns: Iterable[int] = ()):
    """Size columns from headers and values, excluding long ID headings."""
    body_only = set(body_only_columns)
    for column in range(1, ws.max_column + 1):
        values = []
        for row in range(5, ws.max_row + 1):
            if column in body_only and row == 5:
                continue
            value = ws.cell(row, column).value
            if value is None or (isinstance(value, str) and value.startswith("=")):
                continue
            values.append(_display_width(value))
        ws.column_dimensions[get_column_letter(column)].width = min(max(max(values, default=0) + 2, 8), 255)


def _apply_workbook_styles(ws: Worksheet) -> None:
    """Apply the report font, text format, and centered body alignment.

    Rows 1–2 retain their title/period layout. Quantity formats set by each
    sheet builder remain numeric so Excel can still calculate with them.
    """
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            font = copy(cell.font)
            font.name = "Microsoft JhengHei"
            cell.font = font
            if cell.row >= 3:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=cell.alignment.wrap_text)
                if cell.number_format == "General":
                    cell.number_format = "@"


def _parse_export_sections(sections: str | None) -> set[str]:
    """Validate requested worksheets, using the standard three-sheet default."""
    if sections is None:
        return set(DEFAULT_EXPORT_SECTIONS)
    requested = [part.strip() for part in sections.split(",") if part.strip()]
    if not requested or any(part not in EXPORT_SECTION_ORDER for part in requested):
        raise HTTPException(400, "sections 含有不合法的匯出工作表")
    return set(requested)


def _add_table(ws, name: str, header_row: int, style: str = "TableStyleMedium2"):
    """在資料列存在時建立原生 Excel Table，並回傳是否建立成功。"""
    if ws.max_row <= header_row:
        return False
    table = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    return True


def _qty_format(qty_type: str) -> str:
    """依單位數量類型回傳 Excel number format。"""
    return {"integer": "#,##0", "fraction": "# ??/??", "decimal": "#,##0.###"}.get(qty_type, "#,##0.###")


def _write_empty(ws, row: int, text: str = "目前沒有資料"):
    """在空資料工作表寫入使用者可理解的空狀態訊息。"""
    ws.cell(row, 1).value = text
    ws.cell(row, 1).font = Font(color="64748B", italic=True)


def _write_headers(ws, headers, row: int = 5):
    """將欄位標題寫入指定表頭列。"""
    for column, value in enumerate(headers, 1):
        ws.cell(row, column).value = value


def _build_inventory_sheet(ws, items, positions, position_table_available: bool, qty_types):
    """Build the inventory summary, using formulas only when its source table exists.

    Args:
        ws: Inventory worksheet to populate.
        items: Selected item rows from the database.
        positions: Location rows used for values and structured table formulas.
        position_table_available: Whether the exported position table can be referenced.
        qty_types: Unit-name to numeric-format mapping.
    """
    headers = ["品項編號(系統編號)", "庫存區", "廠牌", "品項名稱", "型號", "單位", "低庫存門檻", "待領出", "總庫存", "可用庫存", "位置數", "庫存狀態", "警示序號"]
    _write_headers(ws, headers)
    _style_header(ws, 5)
    quantity_by_item = {}
    count_by_item = {}
    for position in positions:
        item_id = position["id"]
        quantity_by_item[item_id] = quantity_by_item.get(item_id, 0) + (position["qty"] or 0)
        count_by_item[item_id] = count_by_item.get(item_id, 0) + 1
    for item in items:
        row_idx = ws.max_row + 1
        if position_table_available:
            total_formula = f"=SUMIFS(tblPosition[位置數量],tblPosition[品項編號(系統編號)],A{row_idx})"
            position_count_formula = f"=COUNTIFS(tblPosition[品項編號(系統編號)],A{row_idx})"
        else:
            total_formula = quantity_by_item.get(item["id"], 0)
            position_count_formula = count_by_item.get(item["id"], 0)
        available_formula = f"=I{row_idx}-H{row_idx}"
        status_formula = f'=IF(J{row_idx}<0,"資料異常",IF(J{row_idx}=0,"缺貨",IF(AND(G{row_idx}>0,J{row_idx}<=G{row_idx}),"低庫存","正常")))'
        warning_index_formula = f'=IF(L{row_idx}<>"正常",COUNTIF($L$6:L{row_idx},"<>正常"),"")'
        ws.append([item["id"], _site_label(item["site"]), _safe(item["brand"] or "未設定廠牌"), _safe(item["name"]), _safe(item["code"]), _safe(item["unit"]), item["low_stock"] or 0, item["prepared_qty"] or 0, total_formula, available_formula, position_count_formula, status_formula, warning_index_formula])
    if not items:
        _write_empty(ws, 6)
    else:
        _add_table(ws, "tblInventory", 5)
        _style_data(ws, 5, qty_columns=(7, 8, 9, 10, 11))
        for row in range(6, ws.max_row + 1):
            ws.cell(row, 1).alignment = Alignment(horizontal="center")
            fmt = _qty_format(qty_types.get(str(ws.cell(row, 6).value).lstrip("'"), "decimal"))
            for col in (7, 8, 9, 10):
                ws.cell(row, col).number_format = fmt
            ws.cell(row, 11).number_format = "#,##0"
        ws.column_dimensions["M"].hidden = True
        status_range = f"L6:L{ws.max_row}"
        for status, color in STATUS_FILLS.items():
            ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$L6="{status}"'], fill=PatternFill("solid", fgColor=color)))


def _build_position_sheet(ws, positions, qty_types):
    """Build one row per location, omitting category data from the export."""
    headers = ["品項編號(系統編號)", "庫存區", "廠牌", "品項名稱", "型號", "單位", "位置", "位置數量", "位置備註"]
    _write_headers(ws, headers)
    _style_header(ws, 5)
    for row in positions:
        ws.append([row["id"], _site_label(row["site"]), _safe(row["brand"] or "未設定廠牌"), _safe(row["name"]), _safe(row["code"]), _safe(row["unit"]), _safe(row["location"]), row["qty"], _safe(row["note"])])
    if positions:
        _add_table(ws, "tblPosition", 5)
        _style_data(ws, 5, qty_columns=(8,), note_columns=(9,))
        for row in range(6, ws.max_row + 1):
            ws.cell(row, 8).number_format = _qty_format(qty_types.get(str(ws.cell(row, 6).value).lstrip("'"), "decimal"))
    else:
        _write_empty(ws, 6)


def _build_movement_sheet(ws, movements):
    """建立期間異動紀錄，保留原始數量並依值套用顯示格式。"""
    headers = ["時間", "異動類型", "品項編號(系統編號)", "廠牌", "品項名稱", "型號", "庫存區", "變動量", "異動前", "異動後", "去向", "原因"]
    _write_headers(ws, headers)
    _style_header(ws, 5)
    for row in movements:
        reason = row["reason"] or ""
        movement_type = _movement_type(reason, row["delta"])
        ws.append([_safe(row["created_at"]), movement_type, row["item_id"], _safe(row["brand"] or "未設定廠牌"), _safe(row["name"]), _safe(row["code"]), _site_label(row["site"]), row["delta"], row["before_qty"], row["after_qty"], _safe(row["destination"]), _safe(reason)])
    if movements:
        _add_table(ws, "tblMovement", 5)
        _style_data(ws, 5, qty_columns=(8, 9, 10), note_columns=(11, 12))
        for row in range(6, ws.max_row + 1):
            for col in (8, 9, 10):
                value = ws.cell(row, col).value
                if isinstance(value, (int, float)) and float(value).is_integer():
                    ws.cell(row, col).number_format = "#,##0"
    else:
        _write_empty(ws, 6, "選定期間沒有異動紀錄")


def _build_overview(ws: Worksheet, inventory_available: bool, period: str) -> None:
    """Build optional KPI and site summaries from the inventory table.

    Args:
        ws: Overview worksheet to populate.
        inventory_available: Whether a populated inventory table is exported.
        period: Display period shared with the other report sheets.
    """
    _style_title(ws, "庫存管理報表", period)
    _write_headers(ws, ["指標", "數值"])
    _style_header(ws, 5)
    kpis = [
        ("品項數", '=ROWS(tblInventory[品項編號(系統編號)])' if inventory_available else "=0"),
        ("總庫存", '=SUM(tblInventory[總庫存])' if inventory_available else "=SUM(0)"),
        ("待領出", '=SUM(tblInventory[待領出])' if inventory_available else "=SUM(0)"),
        ("可用庫存", '=SUM(tblInventory[可用庫存])' if inventory_available else "=SUM(0)"),
        ("低庫存", '=COUNTIF(tblInventory[庫存狀態],"低庫存")' if inventory_available else "=0"),
        ("缺貨", '=COUNTIF(tblInventory[庫存狀態],"缺貨")' if inventory_available else "=0"),
    ]
    for label, formula in kpis:
        ws.append([label, formula])
    ws["A14"] = "各庫存區摘要"
    ws["A14"].font = Font(bold=True, size=13, color=TITLE_FILL)
    _write_headers(ws, ["庫存區", "品項數", "總庫存", "待領出", "可用", "低庫存", "缺貨"], row=15)
    _style_header(ws, 15)
    for site in SITE_ORDER:
        row = ws.max_row + 1
        if inventory_available:
            values = [
                f'=COUNTIF(tblInventory[庫存區],A{row})',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[總庫存])',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[待領出])',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[可用庫存])',
                f'=COUNTIFS(tblInventory[庫存區],A{row},tblInventory[庫存狀態],"低庫存")',
                f'=COUNTIFS(tblInventory[庫存區],A{row},tblInventory[庫存狀態],"缺貨")',
            ]
        else:
            values = ["=0", "=SUM(0)", "=SUM(0)", "=SUM(0)", "=0", "=0"]
        ws.append([SITES[site], *values])


def _build_stats_sheet(ws: Worksheet, items: Iterable, positions: Iterable, inventory_available: bool, period: str) -> None:
    """Build optional site, category, and brand summaries.

    Category totals are calculated from the export snapshot because category
    is intentionally not a column in the inventory and position worksheets.

    Args:
        ws: Statistics worksheet to populate.
        items: Selected item rows, including category for aggregation.
        positions: Selected location rows used for on-hand totals.
        inventory_available: Whether formula-based table summaries are available.
        period: Display period shared with the other report sheets.
    """
    _style_title(ws, "庫存統計", period)
    ws["A4"] = "庫存區統計"
    ws["J4"] = "分類統計"
    ws["O4"] = "廠牌統計"
    for cell in (ws["A4"], ws["J4"], ws["O4"]):
        cell.font = Font(bold=True, size=13, color=TITLE_FILL)
    _write_headers(ws, ["庫存區", "品項數", "總庫存", "待領出", "可用庫存", "低庫存", "缺貨"], row=5)
    for column, header in enumerate(["分類", "品項數", "總庫存", "可用庫存"], 10):
        ws.cell(5, column).value = header
    for column, header in enumerate(["廠牌", "品項數", "總庫存", "可用庫存"], 15):
        ws.cell(5, column).value = header
    _style_header(ws, 5)

    quantity_by_item = {}
    for position in positions:
        item_id = position["id"]
        quantity_by_item[item_id] = quantity_by_item.get(item_id, 0) + (position["qty"] or 0)
    for site in SITE_ORDER:
        row = ws.max_row + 1
        if inventory_available:
            measures = [
                f'=COUNTIF(tblInventory[庫存區],A{row})',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[總庫存])',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[待領出])',
                f'=SUMIF(tblInventory[庫存區],A{row},tblInventory[可用庫存])',
                f'=COUNTIFS(tblInventory[庫存區],A{row},tblInventory[庫存狀態],"低庫存")',
                f'=COUNTIFS(tblInventory[庫存區],A{row},tblInventory[庫存狀態],"缺貨")',
            ]
        else:
            site_items = [item for item in items if item["site"] == site]
            measures = [
                len(site_items),
                sum(quantity_by_item.get(item["id"], 0) for item in site_items),
                sum(item["prepared_qty"] or 0 for item in site_items),
                sum(quantity_by_item.get(item["id"], 0) - (item["prepared_qty"] or 0) for item in site_items),
                sum(1 for item in site_items if 0 < quantity_by_item.get(item["id"], 0) - (item["prepared_qty"] or 0) <= (item["low_stock"] or 0) and (item["low_stock"] or 0) > 0),
                sum(1 for item in site_items if quantity_by_item.get(item["id"], 0) - (item["prepared_qty"] or 0) == 0),
            ]
        ws.append([SITES[site], *measures])

    category_totals = {}
    brand_totals = {}
    for item in items:
        total = quantity_by_item.get(item["id"], 0)
        available = total - (item["prepared_qty"] or 0)
        category = item["category"] or "未分類"
        brand = item["brand"] or "未設定廠牌"
        for table, label in ((category_totals, category), (brand_totals, brand)):
            values = table.setdefault(label, [0, 0, 0])
            values[0] += 1
            values[1] += total
            values[2] += available
    for row, (category, values) in enumerate(sorted(category_totals.items()), 6):
        ws.cell(row, 10).value = _safe(category)
        ws.cell(row, 11).value = values[0]
        ws.cell(row, 12).value = values[1]
        ws.cell(row, 13).value = values[2]
    for row, (brand, values) in enumerate(sorted(brand_totals.items()), 6):
        ws.cell(row, 15).value = _safe(brand)
        if inventory_available:
            ws.cell(row, 16).value = f'=COUNTIF(tblInventory[廠牌],O{row})'
            ws.cell(row, 17).value = f'=SUMIF(tblInventory[廠牌],O{row},tblInventory[總庫存])'
            ws.cell(row, 18).value = f'=SUMIF(tblInventory[廠牌],O{row},tblInventory[可用庫存])'
        else:
            ws.cell(row, 16).value, ws.cell(row, 17).value, ws.cell(row, 18).value = values
    if not items:
        ws["J6"] = "目前無資料"
        ws["O6"] = "目前無資料"


def _movement_type(reason: str, delta: float) -> str:
    """依 production movement reason contract 分類異動類型。

    先處理明確 reason，再以 delta 作為最後 fallback；避免「組裝套件」等
    負向整組異動被誤標成一般出庫。
    """
    reason = reason or ""
    if reason == "領出準備":
        return "待領出"
    if reason.startswith("解除"):
        return "解除待領出"
    if reason.startswith("出庫"):
        return "出庫"
    if reason == "退回已領出":
        return "退回"
    if reason == "庫存調撥":
        return "調撥"
    if reason.startswith("盤點"):
        return "盤點"
    if reason.startswith("組裝"):
        return "整組組裝"
    if reason.startswith("拆解"):
        return "整組拆解"
    if reason == "品項刪除清零":
        return "刪除清零"
    if "調整" in reason:
        return "庫存調整"
    if delta < 0:
        return "出庫"
    return "其他"


def _build_alert_sheet(ws, items, positions, inventory_available: bool):
    """Build alert rows from the inventory table or standalone snapshot values."""
    headers = ["庫存狀態", "品項編號(系統編號)", "庫存區", "廠牌", "品項名稱", "型號", "單位", "總庫存", "待領出", "可用庫存", "低庫存門檻"]
    _write_headers(ws, headers)
    _style_header(ws, 5)
    if inventory_available and items:
        for row in range(6, len(items) + 6):
            match_formula = "MATCH(ROW()-5,tblInventory[警示序號],0)"
            for column, header in enumerate(headers, 1):
                ws.cell(row, column).value = f'=IFERROR(INDEX(tblInventory[{header}],{match_formula}),"")'
    else:
        totals = {}
        for position in positions:
            totals[position["id"]] = totals.get(position["id"], 0) + (position["qty"] or 0)
        for item in items:
            total = totals.get(item["id"], 0)
            prepared = item["prepared_qty"] or 0
            available = total - prepared
            threshold = item["low_stock"] or 0
            status = "資料異常" if available < 0 else "缺貨" if available == 0 else "低庫存" if threshold > 0 and available <= threshold else "正常"
            if status != "正常":
                ws.append([status, item["id"], _site_label(item["site"]), _safe(item["brand"] or "未設定廠牌"), _safe(item["name"]), _safe(item["code"]), _safe(item["unit"]), total, prepared, available, threshold])
        if ws.max_row == 5:
            _write_empty(ws, 6)


@router.get("/api/export", dependencies=[Depends(require_perm("export"))])
def export_excel(month: str | None = None, start_date: str | None = None, end_date: str | None = None, days: int | None = None, sites: str | None = None, sections: str | None = None):
    """Validate request parameters and return the selected inventory workbook sheets.

    Args:
        month: Optional YYYY-MM period.
        start_date: Inclusive custom range start in YYYY-MM-DD form.
        end_date: Inclusive custom range end in YYYY-MM-DD form.
        days: Optional trailing-day range.
        sites: Comma-separated internal inventory-site identifiers.
        sections: Comma-separated requested workbook sections.

    Returns:
        An XLSX download response containing the requested sheets.

    Raises:
        HTTPException: If a date range, site, or section is invalid.
    """
    selected_sections = _parse_export_sections(sections)
    if days is not None and month is None and start_date is None and end_date is None:
        if days < 0 or days > MAX_RANGE_DAYS:
            raise HTTPException(400, "days 必須介於 0 到 366")
        now = dt.datetime.strptime(movement_time.now_sql(), movement_time.SQL_DATETIME_FORMAT); start = now - dt.timedelta(days=days); end = now; period = f"{start.strftime('%Y/%m/%d')} ～ {end.strftime('%Y/%m/%d')}"; display_period = period
    else:
        start, end, period, display_period = _parse_export_range(month, start_date, end_date)
    selected_sites = list(SITE_ORDER) if not sites else [s for s in sites.split(",") if s]
    if not selected_sites or any(s not in SITES for s in selected_sites):
        raise HTTPException(400, "sites 含有不合法的庫存區")
    conn = get_db()
    try:
        items = conn.execute("SELECT id, site, category, brand, name, code, unit, low_stock, prepared_qty FROM items WHERE is_deleted=0 AND site IN (%s) ORDER BY brand COLLATE NOCASE, name, id" % ",".join("?" * len(selected_sites)), selected_sites).fetchall()
        positions = conn.execute("SELECT i.id, i.site, i.brand, i.name, i.code, i.unit, s.location, s.qty, s.note FROM items i JOIN item_stocks s ON s.item_id=i.id WHERE i.is_deleted=0 AND i.site IN (%s) ORDER BY i.id, s.id" % ",".join("?" * len(selected_sites)), selected_sites).fetchall()
        movement_site = "COALESCE(NULLIF(m.return_site,''), NULLIF(m.source_site,''), NULLIF(i.site,''), '')"
        site_placeholders = ",".join("?" * len(selected_sites))
        movement_sql = (
            "SELECT m.created_at, m.item_id, m.delta, m.before_qty, m.after_qty, "
            "m.destination, m.reason, " + movement_site + " AS site, "
            "i.brand, i.name, i.code "
            "FROM movements m JOIN items i ON i.id=m.item_id "
            "WHERE (" + movement_site + f" IN ({site_placeholders}) OR "
            + movement_site + " = '') AND m.created_at >= ? AND m.created_at < ? "
            "ORDER BY m.created_at DESC, m.id DESC"
        )
        movements = conn.execute(
            movement_sql,
            [*selected_sites, movement_time.datetime_to_sql(start), movement_time.datetime_to_sql(end)],
        ).fetchall()
        qty_types = {row["name"]: row["qty_type"] for row in conn.execute("SELECT name, qty_type FROM units")}
    finally:
        conn.close()
    wb = Workbook(); wb.remove(wb.active); wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.calculation.calcMode = "auto"
    period_text = _period_text(period, display_period)
    if "overview" in selected_sections:
        overview = wb.create_sheet("01 總覽")
        _build_overview(overview, "inventory" in selected_sections and bool(items), period_text)
    if "inventory" in selected_sections:
        inventory = wb.create_sheet("庫存總表(單一庫存)")
        _style_title(inventory, "單一庫存總表", period_text)
        _build_inventory_sheet(inventory, items, positions, "positions" in selected_sections and bool(positions), qty_types)
    if "positions" in selected_sections:
        position = wb.create_sheet("位置明細(單一庫存)")
        _style_title(position, "單一庫存位置明細", period_text)
        _build_position_sheet(position, positions, qty_types)
    if "alerts" in selected_sections:
        alerts = wb.create_sheet("庫存警示(單一庫存)")
        _style_title(alerts, "單一庫存警示", period_text)
        _build_alert_sheet(alerts, items, positions, "inventory" in selected_sections and bool(items))
    if "movements" in selected_sections:
        movement = wb.create_sheet("異動紀錄(單一庫存)")
        _style_title(movement, "單一庫存異動紀錄", period_text)
        _build_movement_sheet(movement, movements)
    if "stats" in selected_sections:
        stats = wb.create_sheet("06 統計")
        _build_stats_sheet(stats, items, positions, "inventory" in selected_sections and bool(items), period_text)
    for sheet in wb.worksheets:
        _apply_workbook_styles(sheet)
        id_column = {"庫存總表(單一庫存)": 1, "位置明細(單一庫存)": 1, "庫存警示(單一庫存)": 2, "異動紀錄(單一庫存)": 3}.get(sheet.title)
        _autofit_columns(sheet, body_only_columns=(id_column,) if id_column else ())
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    stamp = movement_time.now_sql().replace("-", "").replace(":", "").replace(" ", "_")
    if month and not start_date and not end_date:
        filename = f"庫存報表_{month[:4]}年{month[5:]}月_{stamp}.xlsx"
    elif start_date and end_date:
        filename = f"庫存報表_{start_date.replace('-', '')}-{end_date.replace('-', '')}_{stamp}.xlsx"
    else:
        filename = f"庫存報表_{stamp}.xlsx"
    return xlsx_download(buf.getvalue(), filename)
