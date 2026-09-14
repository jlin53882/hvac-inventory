# -*- coding: utf-8 -*-
"""
報價單路由：CRUD、庫存帶入與 Excel/PDF 匯出。

報價單明細會快照品項名稱、規格、單位與單價；inventory_item_id 僅作來源追蹤，
即使日後庫存品項被刪除，歷史報價仍可讀取。
"""
import datetime
import io
import os
from urllib.parse import quote
from xml.sax.saxutils import escape as xml_escape

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from app.database import get_db
from app.models import QuotationIn
from app.services.auth import require_perm
from app.services.safety import excel_safe, parse_ymd, xlsx_download

router = APIRouter()

def _next_quote_number(conn) -> str:
    prefix = f"Q-{datetime.date.today():%Y%m}-"
    rows = conn.execute(
        "SELECT quote_number FROM quotations WHERE quote_number LIKE ?",
        (prefix + "%",),
    ).fetchall()
    used = set()
    for row in rows:
        suffix = row["quote_number"][len(prefix):]
        if suffix.isdigit():
            used.add(int(suffix))
    number = 1
    while number in used:
        number += 1
    return f"{prefix}{number:03d}"


def _quote_dict(conn, quote_id: int):
    row = conn.execute("SELECT * FROM quotations WHERE id=?", (quote_id,)).fetchone()
    if not row:
        raise HTTPException(404, "報價單不存在")
    items = conn.execute(
        "SELECT * FROM quotation_items WHERE quotation_id=? ORDER BY sort_order, id",
        (quote_id,),
    ).fetchall()
    line_total = sum(float(item["qty"]) * float(item["unit_price"]) for item in items)
    if row["tax_type"] == "included":
        tax = line_total * 5 / 105
        subtotal = line_total - tax
        total = line_total
    else:
        subtotal = line_total
        tax = line_total * 0.05
        total = subtotal + tax
    return {
        "id": row["id"],
        "quote_number": row["quote_number"],
        "quote_date": row["quote_date"],
        "customer_name": row["customer_name"],
        "contact": row["contact"] or "",
        "address": row["address"] or "",
        "valid_days": row["valid_days"],
        "tax_type": row["tax_type"],
        "note": row["note"] or "",
        "created_by": row["created_by"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "subtotal": round(subtotal, 2),
        "tax": round(tax, 2),
        "total": round(total, 2),
        "items": [{
            "id": item["id"],
            "inventory_item_id": item["inventory_item_id"],
            "item_name": item["item_name"],
            "specification": item["specification"] or "",
            "qty": item["qty"],
            "unit": item["unit"],
            "unit_price": item["unit_price"],
            "line_total": round(float(item["qty"]) * float(item["unit_price"]), 2),
        } for item in items],
    }


def _validate_inventory_ids(conn, body: QuotationIn) -> None:
    ids = {item.inventory_item_id for item in body.items if item.inventory_item_id is not None}
    if not ids:
        return
    placeholders = ",".join("?" * len(ids))
    rows = conn.execute(
        f"SELECT id FROM items WHERE id IN ({placeholders}) AND is_deleted=0",
        list(ids),
    ).fetchall()
    found = {row["id"] for row in rows}
    missing = sorted(ids - found)
    if missing:
        raise HTTPException(400, f"庫存品項不存在或已刪除: {missing}")


def _write_quote(conn, body: QuotationIn, user_id: int, quote_id: int | None = None):
    quote_date = parse_ymd(body.quote_date)
    customer_name = body.customer_name.strip()
    if not customer_name:
        raise HTTPException(400, "客戶名稱不可空白")
    auto_number = not (body.quote_number or "").strip()
    quote_number = (body.quote_number or "").strip() or _next_quote_number(conn)
    _validate_inventory_ids(conn, body)
    now = datetime.datetime.now().isoformat()
    fields = (
        quote_number, quote_date, customer_name, body.contact.strip(), body.address.strip(),
        body.valid_days, body.tax_type, body.note.strip(), user_id, now,
    )
    if quote_id is None:
        for attempt in range(3):
            try:
                cur = conn.execute(
                    """INSERT INTO quotations
                       (quote_number, quote_date, customer_name, contact, address,
                        valid_days, tax_type, note, created_by, updated_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?)""",
                    fields,
                )
                break
            except Exception as exc:
                if "UNIQUE" not in str(exc).upper() or not auto_number or attempt == 2:
                    if "UNIQUE" in str(exc).upper():
                        raise HTTPException(400, "報價單號已存在")
                    raise
                quote_number = _next_quote_number(conn)
                fields = (quote_number,) + fields[1:]
        else:
            raise HTTPException(409, "報價單號產生衝突，請重試")
        quote_id = cur.lastrowid
    else:
        existing = conn.execute("SELECT id FROM quotations WHERE id=?", (quote_id,)).fetchone()
        if not existing:
            raise HTTPException(404, "報價單不存在")
        try:
            conn.execute(
                """UPDATE quotations SET quote_number=?, quote_date=?, customer_name=?,
                   contact=?, address=?, valid_days=?, tax_type=?, note=?, updated_at=?
                   WHERE id=?""",
                fields[:-2] + (now, quote_id),
            )
        except Exception as exc:
            if "UNIQUE" in str(exc).upper():
                raise HTTPException(400, "報價單號已存在")
            raise
        conn.execute("DELETE FROM quotation_items WHERE quotation_id=?", (quote_id,))
    conn.executemany(
        """INSERT INTO quotation_items
           (quotation_id, inventory_item_id, item_name, specification, qty, unit, unit_price, sort_order)
           VALUES (?,?,?,?,?,?,?,?)""",
        [(quote_id, item.inventory_item_id, item.item_name.strip(), item.specification.strip(),
          item.qty, item.unit.strip(), item.unit_price, index)
         for index, item in enumerate(body.items)],
    )
    return quote_id


@router.get("/api/quotations/inventory-items", dependencies=[Depends(require_perm("view"))])
def quotation_inventory_items(q: str = Query("", max_length=100), limit: int = Query(30, ge=1, le=100)):
    conn = get_db()
    try:
        like = f"%{q.strip()}%"
        rows = conn.execute(
            """SELECT i.id, i.brand, i.code, i.name, i.unit, i.site,
                      COALESCE(SUM(s.qty), 0) AS total_qty
               FROM items i LEFT JOIN item_stocks s ON s.item_id=i.id
               WHERE i.is_deleted=0
                 AND (i.name LIKE ? OR i.brand LIKE ? OR i.code LIKE ?)
               GROUP BY i.id ORDER BY i.name COLLATE NOCASE LIMIT ?""",
            (like, like, like, limit),
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


@router.post("/api/quotations", status_code=201, dependencies=[Depends(require_perm("item-mgmt"))])
def create_quotation(body: QuotationIn, user: dict = Depends(require_perm("item-mgmt"))):
    conn = get_db()
    try:
        quote_id = _write_quote(conn, body, user["id"])
        conn.commit()
        return _quote_dict(conn, quote_id)
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.get("/api/quotations", dependencies=[Depends(require_perm("view"))])
def list_quotations(
    q: str = Query("", max_length=100),
    from_date: str | None = Query(None, max_length=10),
    to_date: str | None = Query(None, max_length=10),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    conn = get_db()
    try:
        clauses = ["1=1"]
        params = []
        if q.strip():
            like = f"%{q.strip()}%"
            clauses.append("(quote_number LIKE ? OR customer_name LIKE ? OR contact LIKE ? OR address LIKE ?)")
            params.extend([like] * 4)
        if from_date:
            clauses.append("quote_date >= ?")
            params.append(parse_ymd(from_date))
        if to_date:
            clauses.append("quote_date <= ?")
            params.append(parse_ymd(to_date))
        where = " AND ".join(clauses)
        total = conn.execute(f"SELECT COUNT(*) FROM quotations WHERE {where}", params).fetchone()[0]
        offset = (page - 1) * page_size
        rows = conn.execute(
            f"SELECT * FROM quotations WHERE {where} ORDER BY quote_date DESC, id DESC LIMIT ? OFFSET ?",
            params + [page_size, offset],
        ).fetchall()
        items = []
        for row in rows:
            quote = _quote_dict(conn, row["id"])
            items.append({k: quote[k] for k in ("id", "quote_number", "quote_date", "customer_name", "contact", "total", "tax_type", "updated_at")})
        return {"items": items, "total": total, "page": page, "page_size": page_size}
    finally:
        conn.close()


@router.get("/api/quotations/{quote_id}", dependencies=[Depends(require_perm("view"))])
def get_quotation(quote_id: int):
    conn = get_db()
    try:
        return _quote_dict(conn, quote_id)
    finally:
        conn.close()


@router.put("/api/quotations/{quote_id}", dependencies=[Depends(require_perm("item-mgmt"))])
def update_quotation(quote_id: int, body: QuotationIn, user: dict = Depends(require_perm("item-mgmt"))):
    conn = get_db()
    try:
        _write_quote(conn, body, user["id"], quote_id)
        conn.commit()
        return _quote_dict(conn, quote_id)
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.delete("/api/quotations/{quote_id}", dependencies=[Depends(require_perm("item-mgmt"))])
def delete_quotation(quote_id: int):
    conn = get_db()
    try:
        cur = conn.execute("DELETE FROM quotations WHERE id=?", (quote_id,))
        if cur.rowcount == 0:
            raise HTTPException(404, "報價單不存在")
        conn.commit()
        return {"ok": True, "deleted": quote_id}
    except HTTPException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _load_for_export(quote_id: int):
    conn = get_db()
    try:
        return _quote_dict(conn, quote_id)
    finally:
        conn.close()


@router.get("/api/quotations/{quote_id}/export.xlsx", dependencies=[Depends(require_perm("export"))])
def export_quotation_xlsx(quote_id: int):
    quote_data = _load_for_export(quote_id)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(500, "缺少 openpyxl，請先安裝")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "報價單"
    ws.append(["報價單號", excel_safe(quote_data["quote_number"])])
    ws.append(["報價日期", excel_safe(quote_data["quote_date"])])
    ws.append(["客戶名稱", excel_safe(quote_data["customer_name"])])
    ws.append(["聯絡人／電話", excel_safe(quote_data["contact"])])
    ws.append(["工程地址", excel_safe(quote_data["address"])])
    ws.append(["有效天數", quote_data["valid_days"]])
    ws.append(["稅別", excel_safe("含稅" if quote_data["tax_type"] == "included" else "未稅")])
    ws.append(["備註／付款條件", excel_safe(quote_data["note"])])
    ws.append([])
    ws.append(["品項名稱", "規格／說明", "數量", "單位", "單價", "小計"])
    header_row = ws.max_row
    for item in quote_data["items"]:
        ws.append([excel_safe(item["item_name"]), excel_safe(item["specification"]), item["qty"],
                   excel_safe(item["unit"]), item["unit_price"], item["line_total"]])
    ws.append([])
    ws.append(["未稅小計", quote_data["subtotal"]])
    ws.append(["稅額", quote_data["tax"]])
    ws.append(["報價總額", quote_data["total"]])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip("ABCDEF", [28, 38, 10, 10, 16, 16]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"報價單_{quote_data['quote_number']}.xlsx"
    return xlsx_download(buf.getvalue(), filename)


def _pdf_text(value) -> str:
    """Escape user text before passing it to ReportLab Paragraph markup."""
    return xml_escape(str(value or ""))


@router.get("/api/quotations/{quote_id}/export.pdf", dependencies=[Depends(require_perm("export"))])
def export_quotation_pdf(quote_id: int):
    quote_data = _load_for_export(quote_id)
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph
    except ImportError:
        raise HTTPException(500, "缺少 reportlab，請先安裝")
    font_name = "Helvetica"
    font_path = "C:/Windows/Fonts/msjh.ttc"
    try:
        if os.path.exists(font_path):
            pdfmetrics.registerFont(TTFont("MSJH", font_path, subfontIndex=0))
            font_name = "MSJH"
    except Exception:
        pass
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("QuoteTitle", parent=styles["Title"], fontName=font_name, fontSize=20, textColor=colors.HexColor("#1E3A5F"), alignment=1)
    body_style = ParagraphStyle("QuoteBody", parent=styles["BodyText"], fontName=font_name, fontSize=9, leading=13)
    right_style = ParagraphStyle("QuoteRight", parent=body_style, alignment=TA_RIGHT)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=15*mm, bottomMargin=15*mm)
    story = [Paragraph("振佳空調　報價單", title_style), Spacer(1, 8)]
    info = [[Paragraph(f"報價單號：{_pdf_text(quote_data['quote_number'])}", body_style), Paragraph(f"報價日期：{_pdf_text(quote_data['quote_date'])}", right_style)],
            [Paragraph(f"客戶名稱：{_pdf_text(quote_data['customer_name'])}", body_style), Paragraph(f"聯絡人：{_pdf_text(quote_data['contact'])}", right_style)],
            [Paragraph(f"工程地址：{_pdf_text(quote_data['address'])}", body_style), Paragraph(f"有效期限：{quote_data['valid_days']} 天", right_style)]]
    info_table = Table(info, colWidths=[105*mm, 65*mm])
    info_table.setStyle(TableStyle([("VALIGN", (0,0), (-1,-1), "TOP"), ("BOTTOMPADDING", (0,0), (-1,-1), 6)]))
    story += [info_table, Spacer(1, 8)]
    data = [[Paragraph("品項名稱", body_style), Paragraph("規格／說明", body_style), Paragraph("數量", body_style), Paragraph("單位", body_style), Paragraph("單價", body_style), Paragraph("小計", body_style)]]
    for item in quote_data["items"]:
        data.append([Paragraph(_pdf_text(item["item_name"]), body_style), Paragraph(_pdf_text(item["specification"]), body_style), str(item["qty"]), _pdf_text(item["unit"]), f"{item['unit_price']:,.2f}", f"{item['line_total']:,.2f}"])
    data += [["", "", "", "", "未稅小計", f"{quote_data['subtotal']:,.2f}"], ["", "", "", "", "稅額", f"{quote_data['tax']:,.2f}"], ["", "", "", "", "報價總額", f"{quote_data['total']:,.2f}"]]
    table = Table(data, colWidths=[35*mm, 60*mm, 15*mm, 15*mm, 25*mm, 30*mm], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1E3A5F")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,-1), font_name), ("FONTSIZE", (0,0), (-1,-1), 8), ("GRID", (0,0), (-1,-1), .5, colors.HexColor("#CBD5E1")), ("ALIGN", (2,1), (-1,-1), "RIGHT"), ("SPAN", (0,-3), (3,-3)), ("SPAN", (0,-2), (3,-2)), ("SPAN", (0,-1), (3,-1)), ("BACKGROUND", (4,-1), (-1,-1), colors.HexColor("#E8F0FE"))]))
    story += [table, Spacer(1, 12), Paragraph(f"備註／付款條件：{_pdf_text(quote_data['note'])}", body_style)]
    doc.build(story)
    filename = f"報價單_{quote_data['quote_number']}.pdf"
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
