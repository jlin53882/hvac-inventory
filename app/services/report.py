# -*- coding: utf-8 -*-
"""
工程日報表生成器（範本填值法）
================================
- 以使用者提供的 0811.xlsx（app/assets/工程日誌範本.xlsx）為範本：
  openpyxl 載入 → 填當天行程 → 另存 xlsx（BytesIO 回傳，不寫磁碟）
- 版面（框線/合併格/欄寬/字型）100% 來自範本，程式不重刻格式
- 填值規則見 docs/行事曆派工-設計文件.md §10.3
"""
import datetime
import io
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"
TEMPLATE_PATH = ASSETS_DIR / "工程日誌範本.xlsx"

# service_types 固定 id 1-4 → 日報表勾選 ✓ 欄位（欄位名 保養/維修/安裝/配管 由範本印在 D/F/H/J）
SVC_CHECK_COL = {1: "E", 2: "G", 3: "I", 4: "K"}
# 5 個預留區塊：資料列 + 備註區第 1 列（第 2 列 = 第 1 列 +1 列）
BLOCKS = [
    {"data": 4, "note": "B5"},
    {"data": 11, "note": "B12"},
    {"data": 18, "note": "B19"},
    {"data": 25, "note": "B26"},
    {"data": 32, "note": "B33"},
]
# 勾選 ✓ 所在欄（清空區塊時要清的資料格；欄位名 D/F/H/J 不動）
CHECK_COLS = ["E", "G", "I", "K"]


def _safe(value):
    """公式注入防護：= + - @ 開頭的字串加撇號，避免被 Excel 當公式執行（沿用 export.py）"""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def build_daily_report(date_str: str, day_events: list, engineers: list = None):
    """填值工程日誌範本 → 回傳 (BytesIO, mmdd)。

    day_events: 每筆含 client_name / address / service_type_id / service_name /
                start_time / end_time / note（依開始時間排序）
    engineers:  **已不使用**（2026-08-12 家豪指定：匯出 Excel 不帶工程師名稱，
                A2 只留「工程師：」由工程師手寫；API 層 assignees 資料仍保留）
    """
    d = datetime.date.fromisoformat(date_str)
    mmdd = d.strftime("%m%d")
    week = "一二三四五六日"[d.weekday()]  # weekday(): 0=週一
    date_label = f"{d.year}年 {d.month} 月 {d.day} 日 星期{week}"

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    ws.title = mmdd
    ws.column_dimensions["B"].width = 20  # B 欄寬固定 20（家豪指定：10.75 放不下時間）

    ws["A2"] = "工程師："  # 不帶名稱（家豪 2026-08-12 指定）
    ws["D2"] = "日期：" + date_label

    # 1. 清空所有區塊資料格（避免殘留範本佔位；欄位名 D/F/H/J 保留）
    for b in BLOCKS:
        r = b["data"]
        for col in ("A", "B", "C") + tuple(CHECK_COLS):
            ws.cell(row=r, column=column_index_from_string(col)).value = None
        ws[b["note"]].value = None
        ws[f"B{int(b['note'][1:]) + 1}"].value = None

    # 2. 填入當天行程（最多 5 筆）
    for i, e in enumerate(day_events[:5]):
        b = BLOCKS[i]
        r = b["data"]
        ws.cell(row=r, column=1).value = i + 1
        # 2026-08-12 補：時間欄也套 _safe()（H4 同類漏網——client/address/note 都有，時間獨漏）
        ws.cell(row=r, column=2).value = _safe(f"{e['start_time']}~{e['end_time']}")
        ws.cell(row=r, column=3).value = _safe(e["client_name"])
        check_col = SVC_CHECK_COL.get(e["service_type_id"])
        if check_col:
            ws[f"{check_col}{r}"] = "✓"
        elif e.get("service_name"):
            # 非固定四類（如自訂服務）→ 地點欄附註
            ws.cell(row=r, column=3).value = _safe(f"{e['client_name']}（{e['service_name']}）")
        # 備註區：第 1 列 地址（有填才顯示 + 前綴）、第 2 列 備註（不加前綴）
        if e.get("address"):
            ws[b["note"]].value = _safe("地址：" + e["address"])
        if e.get("note"):
            ws[f"B{int(b['note'][1:]) + 1}"].value = _safe(e["note"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, mmdd
