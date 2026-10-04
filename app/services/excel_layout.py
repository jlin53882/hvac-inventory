"""Shared content-based column sizing for generated Excel workbooks."""
from datetime import date, datetime, time
import math
import re
import unicodedata

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


def _display_width(value) -> int:
    """Estimate visible width, counting full-width East Asian characters twice."""
    return max(
        (
            sum(2 if unicodedata.east_asian_width(char) in ("F", "W") else 1 for char in line)
            for line in str(value).splitlines()
        ),
        default=0,
    )


def _date_display_value(value, number_format):
    """Format common Excel date/time patterns for a closer visible-width estimate."""
    number_format = (number_format or "").split(";")[0]
    if not re.search(r"[ymdhs]", number_format, re.IGNORECASE):
        return value.strftime("%Y-%m-%d %H:%M") if isinstance(value, datetime) else value.isoformat()

    has_time = bool(re.search(r"[hs]", number_format, re.IGNORECASE))
    token_pattern = re.compile(r"yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s", re.IGNORECASE)

    def replace_token(match):
        token = match.group(0)
        lowered = token.lower()
        if lowered.startswith("y"):
            return f"{value.year % 100:02d}" if lowered == "yy" else f"{value.year:04d}"
        if lowered.startswith("d"):
            if lowered in ("ddd", "dddd"):
                return value.strftime("%a" if lowered == "ddd" else "%A")
            return f"{value.day:02d}" if lowered == "dd" else str(value.day)
        if lowered.startswith("h"):
            return f"{value.hour:02d}" if lowered == "hh" else str(value.hour)
        if lowered.startswith("s"):
            return f"{value.second:02d}" if lowered == "ss" else str(value.second)
        if lowered.startswith("m") and has_time:
            return f"{value.minute:02d}" if lowered == "mm" else str(value.minute)
        if lowered == "mmmm":
            return value.strftime("%B")
        if lowered == "mmm":
            return value.strftime("%b")
        return f"{value.month:02d}" if lowered == "mm" else str(value.month)

    return token_pattern.sub(replace_token, number_format)


def _cell_display_value(cell):
    value = cell.value
    if value is None:
        return None
    if isinstance(value, str):
        # openpyxl does not calculate formula results; reserve a useful minimum
        # for their displayed values instead of sizing to the formula source.
        return "0" * 12 if value.startswith("=") else value
    if isinstance(value, datetime):
        return _date_display_value(value, cell.number_format)
    if isinstance(value, date):
        return _date_display_value(value, cell.number_format)
    if isinstance(value, time):
        return _date_display_value(datetime.combine(date.min, value), cell.number_format)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and math.isfinite(value):
        number_format = (cell.number_format or "").split(";")[0]
        if number_format in ("", "General"):
            return str(value)
        decimal_pattern = number_format.split(".", 1)[1] if "." in number_format else ""
        decimal_places = sum(char in "0#?" for char in decimal_pattern)
        if decimal_places:
            return f"{value:,.{decimal_places}f}".rstrip("0").rstrip(".")
        return f"{value:,.0f}"
    return str(value)


def auto_fit_columns(ws, min_width: float = 8, max_width: float = 255) -> None:
    """Size every worksheet column from all cell values, including headers.

    Widths use a server-side display estimate because openpyxl cannot ask Excel
    to calculate native AutoFit. Text in horizontal merged cells contributes
    to the combined width of that merged range instead of one anchor column.
    """
    widths = {column: min_width for column in range(1, ws.max_column + 1)}
    merged_anchors = {
        (merged.min_row, merged.min_col): merged
        for merged in ws.merged_cells.ranges
    }
    merged_requirements = []

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            display_value = _cell_display_value(cell)
            if display_value is None:
                continue
            required = _display_width(display_value) + 2
            merged = merged_anchors.get((cell.row, cell.column))
            if merged is not None and merged.max_col > merged.min_col:
                columns = tuple(range(merged.min_col, merged.max_col + 1))
                merged_requirements.append((columns, required))
            else:
                widths[cell.column] = max(widths[cell.column], required)

    # Each horizontal merged value constrains the total width of its span.
    # Iteration handles overlapping merged ranges without assigning a title's
    # full width to its first column.
    for _ in range(max(1, len(merged_requirements))):
        changed = False
        for columns, required in sorted(merged_requirements, key=lambda item: len(item[0])):
            available = sum(widths[column] for column in columns)
            if available < required:
                addition = (required - available) / len(columns)
                for column in columns:
                    widths[column] += addition
                changed = True
        if not changed:
            break

    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = min(max(width, min_width), max_width)
