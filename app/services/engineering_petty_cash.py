# -*- coding: utf-8 -*-
"""工程零用金：階層資料計算、檔名與範本匯出。"""
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import datetime
import io

from openpyxl import load_workbook

from app.services.petty_cash_excel import (
    configure_print_layout,
    copy_role_style,
    ensure_page_defaults,
    period_display,
    period_token,
)
from app.services.safety import excel_safe

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "工程零用金範本.xlsx"


def money(value) -> Decimal:
    try:
        return Decimal(str(value or 0)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0.00")


def calculate_engineering_totals(categories):
    total = Decimal("0.00")
    for category in categories:
        cat_total = Decimal("0.00")
        for group in category.get("groups", []):
            group_total = sum((money(r.get("amount")) for r in group.get("receipts", [])), Decimal("0.00"))
            group["subtotal"] = group_total
            cat_total += group_total
        category["subtotal"] = cat_total
        total += cat_total
    return {"categories": categories, "total_amount": total}


def engineering_summary_totals(conn, report_ids):
    """Return report totals in one aggregate query for mixed report lists."""
    ids = [int(report_id) for report_id in report_ids]
    if not ids:
        return {}
    placeholders = ",".join("?" for _ in ids)
    rows = conn.execute(
        f"""SELECT c.report_id, COALESCE(SUM(r.amount), 0) AS total_amount
            FROM engineering_expense_categories c
            JOIN engineering_expense_groups g ON g.category_id = c.id
            JOIN engineering_expense_receipts r ON r.group_id = g.id
            WHERE c.report_id IN ({placeholders})
            GROUP BY c.report_id""",
        ids,
    ).fetchall()
    return {row["report_id"]: money(row["total_amount"]) for row in rows}


def engineering_filename(start_date, end_date, owner, note=""):
    period = period_token(start_date, end_date)
    note = (note or "").strip()
    inside = f"{period} {note}" if note else period
    return excel_safe(f"({inside}){(owner or '').strip()} 工程零用金.xlsx")


def engineering_sheet_title(start_date, end_date):
    """工程零用金工作表名稱：使用不含 Excel 禁用字元的期間 token。"""
    return period_token(start_date, end_date)

def engineering_safe_filename(name):
    """工程報表保留規格所需括號/空格，同時阻擋路徑與危險字元。"""
    name = (name or "file").strip().replace("/", "_").replace("\\", "_")
    name = re.sub(r"[^0-9A-Za-z\u4e00-\u9fa5._()\- ]", "_", name)
    suffix = ".xlsx"
    stem = name[:-len(suffix)] if name.lower().endswith(suffix) else name
    if stem[:1] in (".", "-", "=", "+", "@"):
        stem = "_" + stem[1:]
    stem = stem[:120 - len(suffix)] or "file"
    return stem + suffix


def _merge(ws, first, last, col):
    if last > first:
        ws.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)


def _engineering_row(conn, report_id):
    row = conn.execute("SELECT * FROM petty_cash_reports WHERE id=? AND report_type='engineering'", (report_id,)).fetchone()
    if row is None:
        return None
    categories = []
    cat_rows = conn.execute("SELECT * FROM engineering_expense_categories WHERE report_id=? ORDER BY sort_order,id", (report_id,)).fetchall()
    for cat in cat_rows:
        groups = []
        for group in conn.execute("SELECT * FROM engineering_expense_groups WHERE category_id=? ORDER BY sort_order,id", (cat['id'],)).fetchall():
            receipts = []
            for receipt in conn.execute("SELECT * FROM engineering_expense_receipts WHERE group_id=? ORDER BY sort_order,id", (group['id'],)).fetchall():
                details = [r['description'] for r in conn.execute("SELECT description FROM engineering_expense_details WHERE receipt_id=? ORDER BY sort_order,id", (receipt['id'],)).fetchall()]
                receipts.append({"id": receipt['id'], "tax_id_mark": receipt['tax_id_mark'] or "", "receipt_number": receipt['receipt_number'] or "", "amount": money(receipt['amount']), "details": details or [""]})
            groups.append({"id": group['id'], "name": group['name'], "receipts": receipts})
        categories.append({"id": cat['id'], "name": cat['name'], "groups": groups})
    totals = calculate_engineering_totals(categories)
    return {"id": row['id'], "report_type": "engineering", "start_date": row['start_date'], "end_date": row['end_date'], "filename_text": row['filename_text'] or "", "upload_person": row['upload_person'], "prepared_by": row['prepared_by'], "status": row['status'], "created_at": row['created_at'], "updated_at": row['updated_at'], "categories": categories, "total_amount": totals['total_amount'], "filename": engineering_filename(row['start_date'], row['end_date'], row['upload_person'], row['filename_text'])}


def write_engineering(conn, body, user_id, report_id=None):
    start = datetime.date.fromisoformat(body.start_date); end = datetime.date.fromisoformat(body.end_date)
    if start > end: raise ValueError("開始日期不可晚於結束日期")
    if report_id is None:
        cur = conn.execute("INSERT INTO petty_cash_reports (report_type,start_date,end_date,filename_text,upload_person,uploader_user_id,prepared_by,status,created_by,updated_at) VALUES ('engineering',?,?,?,?,?,?,?, ?,datetime('now'))", (body.start_date, body.end_date, body.filename_text, body.upload_person, user_id, body.prepared_by, body.status, user_id))
        report_id = cur.lastrowid
    else:
        existing = conn.execute("SELECT id,report_type FROM petty_cash_reports WHERE id=?", (report_id,)).fetchone()
        if not existing: raise KeyError("零用金月報不存在")
        if existing['report_type'] != 'engineering': raise ValueError("報表類型不可變更")
        conn.execute("UPDATE petty_cash_reports SET start_date=?,end_date=?,filename_text=?,upload_person=?,prepared_by=?,status=?,updated_at=datetime('now') WHERE id=?", (body.start_date,body.end_date,body.filename_text,body.upload_person,body.prepared_by,body.status,report_id))
        conn.execute("DELETE FROM engineering_expense_categories WHERE report_id=?", (report_id,))
    for ci, cat in enumerate(body.categories):
        c = conn.execute("INSERT INTO engineering_expense_categories(report_id,name,sort_order) VALUES(?,?,?)", (report_id,cat.name,ci)).lastrowid
        for gi, group in enumerate(cat.groups):
            g = conn.execute("INSERT INTO engineering_expense_groups(category_id,name,sort_order) VALUES(?,?,?)", (c,group.name,gi)).lastrowid
            for ri, receipt in enumerate(group.receipts):
                r = conn.execute("INSERT INTO engineering_expense_receipts(group_id,tax_id_mark,receipt_number,amount,sort_order) VALUES(?,?,?,?,?)", (g,receipt.tax_id_mark,receipt.receipt_number,str(money(receipt.amount)),ri)).lastrowid
                for di, detail in enumerate(receipt.details):
                    conn.execute("INSERT INTO engineering_expense_details(receipt_id,description,sort_order) VALUES(?,?,?)", (r,detail,di))
    return report_id


def _xlsx_number(value):
    amount = money(value)
    return int(amount) if amount == amount.to_integral_value() else float(amount)


def _clear_visible_body(ws, first_body_row):
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    if ws.max_row >= first_body_row:
        ws.delete_rows(first_body_row, ws.max_row - first_body_row + 1)
    for key in list(ws.row_dimensions):
        try:
            row_number = int(key)
        except (TypeError, ValueError):
            continue
        if row_number >= first_body_row:
            del ws.row_dimensions[key]


def build_engineering_report(report):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    styles = wb["__styles__"]
    _clear_visible_body(ws, 3)
    ws.merge_cells("A1:F1")
    ws["A1"] = f"{period_display(report['start_date'], report['end_date'])}工程零用金明細表"
    for col, value in enumerate(["類別", "項目", "統編", "發票號碼", "細項", "金額"], 1):
        ws.cell(2, col).value = value

    categories = report.get("categories") or []
    calculated = calculate_engineering_totals(categories)
    categories = calculated["categories"]
    calculated_total = calculated["total_amount"]
    data_rows = []
    category_merges = []
    group_merges = []
    receipt_merges = []
    subtotal_rows = []
    current_row = 3

    def append_data(category_name, group_name, receipt, detail, category_start, group_start, receipt_start):
        data_rows.append({
            "category": category_name,
            "group": group_name,
            "tax_id_mark": (receipt or {}).get("tax_id_mark", "") if receipt else "",
            "receipt_number": (receipt or {}).get("receipt_number", "") if receipt else "",
            "detail": detail,
            "amount": (receipt or {}).get("amount", 0) if receipt else 0,
            "category_anchor": current_row == category_start,
            "group_anchor": current_row == group_start,
            "receipt_anchor": current_row == receipt_start,
        })

    for category in categories:
        category_start = current_row
        groups = category.get("groups") or []
        if not groups:
            append_data(category.get("name", ""), "", None, "", category_start, current_row, current_row)
            current_row += 1
        else:
            for group in groups:
                group_start = current_row
                receipts = group.get("receipts") or []
                if not receipts:
                    append_data(category.get("name", ""), group.get("name", ""), None, "", category_start, group_start, current_row)
                    current_row += 1
                else:
                    for receipt in receipts:
                        receipt_start = current_row
                        details = receipt.get("details") or [""]
                        for detail_index, detail in enumerate(details):
                            append_data(
                                category.get("name", ""),
                                group.get("name", ""),
                                receipt,
                                detail,
                                category_start,
                                group_start,
                                receipt_start,
                            )
                            current_row += 1
                        receipt_end = current_row - 1
                        if receipt_end > receipt_start:
                            receipt_merges.extend((receipt_start, receipt_end, col) for col in (3, 4, 6))
                group_end = current_row - 1
                if group_end >= group_start and group_end > group_start:
                    group_merges.append((group_start, group_end, 2))
        category_end = current_row - 1
        if category_end > category_start:
            category_merges.append((category_start, category_end, 1))
        subtotal_rows.append((current_row, category.get("subtotal", 0)))
        current_row += 1

    for index, record in enumerate(data_rows):
        row_number = 3 + index
        copy_role_style(styles, "engineering_data", ws, row_number)
        if record["category_anchor"]:
            ws.cell(row_number, 1).value = excel_safe(record["category"])
        if record["group_anchor"]:
            ws.cell(row_number, 2).value = excel_safe(record["group"])
        if record["receipt_anchor"]:
            tax = ws.cell(row_number, 3)
            tax.value = excel_safe(str(record["tax_id_mark"] or ""))
            tax.number_format = "@"
            number = ws.cell(row_number, 4)
            number.value = excel_safe(str(record["receipt_number"] or ""))
            number.number_format = "@"
            ws.cell(row_number, 6).value = _xlsx_number(record["amount"])
        ws.cell(row_number, 5).value = excel_safe(record["detail"])

    for first, last, col in receipt_merges + group_merges + category_merges:
        _merge(ws, first, last, col)

    for row_number, subtotal in subtotal_rows:
        copy_role_style(styles, "engineering_subtotal", ws, row_number)
        ws.cell(row_number, 1).value = "小計:"
        ws.cell(row_number, 6).value = _xlsx_number(subtotal)
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=5)

    spacer_row = current_row
    copy_role_style(styles, "engineering_spacer", ws, spacer_row)
    total_top = spacer_row + 1
    for role, row_number in (
        ("engineering_total_top", total_top),
        ("engineering_total_middle", total_top + 1),
        ("engineering_total_bottom", total_top + 2),
    ):
        copy_role_style(styles, role, ws, row_number)
    ws.cell(total_top, 1).value = "總計:"
    ws.cell(total_top, 6).value = _xlsx_number(report.get("total_amount", calculated_total))
    ws.merge_cells(start_row=total_top, start_column=1, end_row=total_top + 2, end_column=5)
    ws.merge_cells(start_row=total_top, start_column=6, end_row=total_top + 2, end_column=6)

    ws.title = engineering_sheet_title(report["start_date"], report["end_date"])
    ensure_page_defaults(ws)
    configure_print_layout(ws, total_top + 2, title_rows="$1:$2")
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
