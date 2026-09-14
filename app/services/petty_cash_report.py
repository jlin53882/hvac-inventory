# -*- coding: utf-8 -*-
"""
零用金月報 Excel 生成器（範本填值法）
=====================================
- 以 app/assets/零用金月報範本.xlsx（由真實主管報表清理而來）為範本：
  openpyxl 載入 → 填月報資料 → 另存 xlsx（BytesIO 回傳，不寫磁碟）
- 版面（框線/合併格/欄寬/字型/列印設定）100% 來自範本，程式不重刻格式
- 資料庫不用 merged-cell 思維：多項目支出在匯出時才展開成多列，
  屬於同一 Entry 的日期/收入/支出/科目欄做跨列合併並垂直置中
"""
import datetime
import io
from pathlib import Path

from openpyxl import load_workbook

from app.services.petty_cash_excel import (
    INTEGER_MONEY_FORMAT,
    configure_print_layout,
    copy_cell_style,
    copy_role_style,
    ensure_page_defaults,
    period_display,
    period_token,
)
from app.services.safety import excel_safe

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"
TEMPLATE_PATH = ASSETS_DIR / "零用金月報範本.xlsx"

def _num_g(value) -> str:
    """明細文字用數字格式：24.0 → '24'、1.5 → '1.5'（對齊原模板『24瓶 $200』寫法）。"""
    try:
        return "%g" % float(value)
    except (TypeError, ValueError):
        return str(value or "")


def _item_text(item: dict) -> str:
    name = (item.get("item_name") or "").strip()
    qty = _num_g(item.get("qty"))
    unit = (item.get("unit") or "").strip()
    amount = _num_g(item.get("amount"))
    return excel_safe(f"{name} {qty}{unit} ${amount}")


def _mmdd(value):
    return datetime.date.fromisoformat(value).strftime("%m%d")


def _xlsx_number(value):
    amount = round(float(value or 0), 2)
    return int(amount) if amount.is_integer() else amount


def _clear_visible_body(ws, first_body_row):
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    if ws.max_row >= first_body_row:
        ws.delete_rows(first_body_row, ws.max_row - first_body_row + 1)
    for key in list(ws.row_dimensions):
        try:
            row_number = int(key)
        except (TypeError, ValueError):
            continue
        if row_number >= first_body_row:
            del ws.row_dimensions[key]


def build_petty_cash_report(report: dict) -> io.BytesIO:
    """Render a general petty-cash report with an exact dynamic body/footer."""
    start = report["start_date"]
    end = report["end_date"]
    opening = round(float(report.get("opening_balance") or 0), 2)

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    styles = wb["__styles__"]
    _clear_visible_body(ws, 4)
    ws.merge_cells("A1:F1")
    ws["A1"] = f"{period_display(start, end)}零用金收支明細表"
    roc_year = datetime.date.fromisoformat(start).year - 1911
    ws["B2"] = f"{roc_year}年"
    ws["D2"] = "上期餘額"
    ws["E2"] = _xlsx_number(opening)
    ws["E2"].number_format = INTEGER_MONEY_FORMAT
    for col, value in enumerate(["項次", "日   期", "摘         要", "收   入", "支   出", "科 目"], 1):
        ws.cell(3, col).value = value

    entries = sorted(
        report.get("entries") or [],
        key=lambda entry: (entry.get("entry_date") or "", entry.get("sort_order") or 0, entry.get("id") or 0),
    )
    expanded = []
    for entry in entries:
        items = sorted(
            entry.get("items") or [],
            key=lambda item: (item.get("sort_order") or 0, item.get("id") or 0),
        )
        if entry.get("entry_type") == "expense" and items:
            expanded.extend((entry, item, index, len(items)) for index, item in enumerate(items))
        else:
            expanded.append((entry, None, 0, 1))

    body_count = len(expanded)
    closing_row = 4 + body_count
    label_row = closing_row + 2
    ws.insert_rows(4, body_count + 3)

    seq = 0
    for index, (entry, item, item_index, item_count) in enumerate(expanded):
        row_number = 4 + index
        if item_count == 1:
            role = "general_data"
        elif item_index == 0:
            role = "general_multi_first"
        elif item_index == item_count - 1:
            role = "general_multi_last"
        else:
            role = "general_multi_middle"
        copy_role_style(styles, role, ws, row_number)
        seq += 1
        ws.cell(row_number, 1).value = seq
        entry_date = datetime.date.fromisoformat(entry["entry_date"])
        if item is not None:
            ws.cell(row_number, 3).value = _item_text(item)
            if item_index == 0:
                ws.cell(row_number, 2).value = entry_date
                ws.cell(row_number, 5).value = _xlsx_number(entry["amount"])
                ws.cell(row_number, 5).number_format = INTEGER_MONEY_FORMAT
                ws.cell(row_number, 6).value = excel_safe((entry.get("category") or "").strip())
            if item_count > 1 and item_index == 0:
                for col in (2, 4, 5, 6):
                    ws.merge_cells(
                        start_row=row_number,
                        start_column=col,
                        end_row=row_number + item_count - 1,
                        end_column=col,
                    )
        else:
            ws.cell(row_number, 2).value = entry_date
            ws.cell(row_number, 3).value = excel_safe((entry.get("description") or "").strip())
            amount = _xlsx_number(entry["amount"])
            amount_column = 4 if entry.get("entry_type") == "income" else 5
            ws.cell(row_number, amount_column).value = amount
            ws.cell(row_number, amount_column).number_format = INTEGER_MONEY_FORMAT
            ws.cell(row_number, 6).value = excel_safe((entry.get("category") or "").strip())

    copy_role_style(styles, "general_closing", ws, closing_row)
    ws.cell(closing_row, 4).value = "本期餘額"
    if body_count:
        last_body_row = closing_row - 1
        closing_formula = f"=E2+SUM(D4:D{last_body_row})-SUM(E4:E{last_body_row})"
    else:
        closing_formula = "=E2"
    ws.cell(closing_row, 5).value = closing_formula
    ws.cell(closing_row, 5).number_format = INTEGER_MONEY_FORMAT
    copy_role_style(styles, "general_signature_label", ws, label_row)
    ws.cell(label_row, 2).value = "主管:"
    ws.cell(label_row, 4).value = "製表人:"
    copy_cell_style(styles.cell(10, 5), ws.cell(label_row, 5))
    ws.cell(label_row, 5).value = excel_safe((report.get("prepared_by") or "").strip())

    ws.title = period_token(start, end)
    ensure_page_defaults(ws)
    configure_print_layout(ws, label_row, title_rows="$1:$3")
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def download_filename(filename_text: str, start_date: str, end_date: str) -> str:
    """下載檔名（安全格式，無 '/'）：零用金-資材0826-0925.xlsx。"""
    return f"零用金-{filename_text}{_mmdd(start_date)}-{_mmdd(end_date)}.xlsx"
