# -*- coding: utf-8 -*-
"""Shared, presentation-only helpers for petty-cash Excel exports."""
from copy import copy
import datetime

from openpyxl.worksheet.page import PageMargins


XLSX_SUFFIX = ".xlsx"


def _as_date(value):
    return datetime.date.fromisoformat(str(value))


def period_display(start_date, end_date):
    """Return the user-facing report period used inside worksheet titles."""
    start = _as_date(start_date)
    end = _as_date(end_date)
    if start == end:
        return start.strftime("%m/%d")
    if start.year == end.year:
        return f"{start:%m/%d}~{end:%m/%d}"
    return f"{start:%Y/%m/%d}~{end:%Y/%m/%d}"


def period_token(start_date, end_date):
    """Return a sheet/filename-safe period token without slash characters."""
    start = _as_date(start_date)
    end = _as_date(end_date)
    if start == end:
        return start.strftime("%m%d")
    if start.year == end.year:
        return f"{start:%m%d}-{end:%m%d}"
    return f"{start:%Y%m%d}-{end:%Y%m%d}"


def copy_cell_style(source, target):
    """Copy the complete visual style without copying the source value."""
    if source.has_style:
        target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def copy_row_style(source_ws, source_row, target_ws, target_row, columns=6):
    """Copy a row prototype's height and cell styles."""
    target_ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height
    for col in range(1, columns + 1):
        copy_cell_style(source_ws.cell(source_row, col), target_ws.cell(target_row, col))


def copy_role_style(style_ws, role, target_ws, target_row, columns=6):
    """Copy a labelled row from the hidden ``__styles__`` sheet."""
    source_row = next(
        (row for row in range(1, style_ws.max_row + 1)
         if style_ws.cell(row, 7).value == role),
        None,
    )
    if source_row is None:
        raise KeyError(f"missing petty-cash Excel style role: {role}")
    copy_row_style(style_ws, source_row, target_ws, target_row, columns)


def configure_print_layout(ws, last_row, title_rows="$1:$3"):
    """Set output bounds so preformatted template tails cannot leak into print."""
    ws.print_area = f"$A$1:$F${last_row}"
    ws.print_title_rows = title_rows
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.scale = None
    ws.auto_filter.ref = None


def ensure_page_defaults(ws):
    """Keep the production portrait/A4 layout while allowing multiple pages."""
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(
        left=ws.page_margins.left or 0.25,
        right=ws.page_margins.right or 0.25,
        top=ws.page_margins.top or 0.75,
        bottom=ws.page_margins.bottom or 0.75,
        header=ws.page_margins.header or 0.3,
        footer=ws.page_margins.footer or 0.3,
    )
